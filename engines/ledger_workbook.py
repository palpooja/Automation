"""engines/ledger_workbook.py - one combined Ledgers workbook per client.

Built at the end of each client's Ledgers step from the files already in the
client's Ledgers folder (no portal work, so it adds no portal time):

    Ledgers_<TIN>_<from>_to_<to>.xlsx
      Contents                 every tax type and month of the period, with its status
      <tax type> - L2          the Level 2 ledger exactly as saved (merged headers kept)
      <tax type> - All L3      every Level 3 month stacked, with Month and Source file columns

The separate Level 2 / Level 3 files stay: they decide what is already downloaded,
and they are the portal's data file by file.  This workbook is rebuilt every run,
so it also contains months downloaded in earlier runs.

Accuracy rules: values, number formats and bold headers are copied cell by cell;
text stays text (never a formula); when a month's header differs from the first
month's, that month's own header rows are written above its rows, so no column
is ever misaligned.
"""
import datetime
import os
import re
from pathlib import Path

from engines.excel_files import excel_ok

MAX_SHEET = 31


def _sheet_name(text, used):
    name = re.sub(r"[\[\]:*?/\\]", "-", text)[:MAX_SHEET]
    base, n = name, 2
    while name in used:
        suffix = f" ({n})"
        name = base[:MAX_SHEET - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def _copy_cell(src, dst):
    from engines.sheet_safety import as_text_cell
    from openpyxl.styles import Font
    dst.value = src.value
    as_text_cell(dst)
    if isinstance(src.value, str):
        dst.data_type = "s"
    if src.has_style:
        dst.number_format = src.number_format
        if src.font is not None and src.font.b:
            dst.font = Font(bold=True)


def _rows(ws):
    return [list(r) for r in ws.iter_rows()]


def _header_count(rows):
    """Leading rows with no number in them are header rows."""
    h = 0
    for row in rows:
        if any(isinstance(c.value, (int, float)) and not isinstance(c.value, bool) for c in row):
            break
        h += 1
    return h


def _copy_sheet(src_ws, dst_ws, row0=1, col0=1, rows=None):
    """Copy cells (and merges inside the copied rows) to dst starting at (row0, col0)."""
    rows = rows if rows is not None else _rows(src_ws)
    first = rows[0][0].row if rows and rows[0] else 1
    last = rows[-1][0].row if rows and rows[-1] else 0
    for i, row in enumerate(rows):
        for j, c in enumerate(row):
            _copy_cell(c, dst_ws.cell(row=row0 + i, column=col0 + j))
    for m in src_ws.merged_cells.ranges:
        if m.min_row >= first and m.max_row <= last:
            dst_ws.merge_cells(start_row=row0 + m.min_row - first, start_column=col0 + m.min_col - 1,
                               end_row=row0 + m.max_row - first, end_column=col0 + m.max_col - 1)
    return len(rows)


def _texts(rows):
    return [tuple("" if c.value is None else str(c.value) for c in r) for r in rows]


def _stack_level3(dst_ws, sources):
    """sources: [(label, path)].  Returns {label: data_rows}."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    counts, r, current_header = {}, 1, None
    for label, path in sources:
        ws = load_workbook(path).active
        rows = _rows(ws)
        h = _header_count(rows)
        header, data = rows[:h], rows[h:]
        if header and _texts(header) != current_header:
            if current_header is not None:
                r += 1                                       # blank line before a different header
            dst_ws.cell(row=r, column=1, value="Month").font = Font(bold=True)
            dst_ws.cell(row=r, column=2, value="Source file").font = Font(bold=True)
            _copy_sheet(ws, dst_ws, row0=r, col0=3, rows=header)
            r += len(header)
            current_header = _texts(header)
        for row in data:
            dst_ws.cell(row=r, column=1, value=label).data_type = "s"
            dst_ws.cell(row=r, column=2, value=Path(path).name).data_type = "s"
            for j, c in enumerate(row):
                _copy_cell(c, dst_ws.cell(row=r, column=3 + j))
            r += 1
        counts[label] = len(data)
    return counts


def _status(results, section, item, progress):
    for rec in reversed(results):
        if rec.get("Section") == section and rec.get("Item") == item:
            return str(rec.get("Status", ""))
    value = str(progress.get(item, "") or "")
    return value.split(" @")[0] if value else "not downloaded"


def build(save_folder, tin, period_from, period_to, tax_types, month_labels, annual_types,
          results=(), progress=None, now=None):
    """Write the combined workbook; returns its path, or None when there is nothing to combine."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    save_folder = Path(save_folder)
    progress = progress or {}
    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    used = {"Contents"}
    lines = []                    # (tax type, month/item, level 3 status, rows, file)
    made_any = False
    for tax in tax_types:
        section = f"{tax} ledger"
        l2 = save_folder / f"{tax}_L2.xlsx"
        if excel_ok(l2):
            from openpyxl import load_workbook
            ws = wb.create_sheet(_sheet_name(f"{tax} - L2", used))
            _copy_sheet(load_workbook(l2).active, ws)
            made_any = True
            lines.append((tax, "Level 2", "included", "", l2.name))
        else:
            lines.append((tax, "Level 2", "not available - see the run summary", "", ""))
        if tax in annual_types:
            sources = [(p.stem[len(f"{tax}_L3_"):], p) for p in sorted(save_folder.glob(f"{tax}_L3_*.xlsx"))
                       if excel_ok(p)]
            wanted = []
        else:
            sources = [(m, save_folder / f"{tax}_L3_{m}.xlsx") for m in month_labels]
            wanted = list(month_labels)
            sources = [(m, p) for m, p in sources if excel_ok(p)]
        counts = {}
        if sources:
            ws3 = wb.create_sheet(_sheet_name(f"{tax} - All L3", used))
            counts = _stack_level3(ws3, sources)
            made_any = True
        tax_progress = progress.get(tax, {})
        for label in wanted or [s[0] for s in sources]:
            if label in counts:
                lines.append((tax, label, "included", counts[label], f"{tax}_L3_{label}.xlsx"))
            else:
                lines.append((tax, label, _status(results, section, label, tax_progress), "", ""))
    if not made_any:
        return None
    now = now or datetime.datetime.now()
    contents.append([f"Ledgers - client TIN {tin}"])
    contents["A1"].font = Font(bold=True, size=13)
    contents.append([f"Period {period_from} to {period_to}; built {now:%d/%m/%Y %H:%M} from the files in this "
                     "folder (they stay as they are; this workbook is rebuilt every run)"])
    contents.append([])
    contents.append(["Tax type", "Month / item", "Level 3", "Rows", "Source file"])
    for c in contents[4]:
        c.font = Font(bold=True)
    for line in lines:
        contents.append(list(line))
    for col, w in zip("ABCDE", (20, 24, 34, 8, 44)):
        contents.column_dimensions[col].width = w
    target = save_folder / f"Ledgers_{tin}_{period_from}_to_{period_to}.xlsx"
    tmp = save_folder / f"~Ledgers_{tin}.tmp.xlsx"
    wb.save(tmp)
    try:
        os.replace(tmp, target)
    except PermissionError:                               # open in Excel: save beside it
        target = target.with_name(f"{target.stem}_{now:%H%M%S}.xlsx")
        os.replace(tmp, target)
    return target
