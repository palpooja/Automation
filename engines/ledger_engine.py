import os
import re
import io
import sys
import base64
import json
import shutil
import time
import datetime
import traceback
import threading
import queue
import pandas as pd
from pathlib import Path
try:
    from dotenv import load_dotenv
except Exception:          # optional in the suite
    load_dotenv = lambda *a, **k: None
from playwright.sync_api import sync_playwright

# ---------- Stop button and on-screen questions (used by the window) ----------

class StopRequested(BaseException):
    """Raised when Stop is pressed. BaseException, so the many
    'except Exception' blocks below cannot swallow it."""

STOP_EVENT = threading.Event()
CREDS = {"tin": None, "password": None}   # the CURRENT client only; cleared after each client
CURRENT_TIN = ""                           # names error screenshots so clients never mix
NOTES = []                                 # (item, note) that a person must check; drained by callers
REFRESH_L2 = True                          # re-export a Level 2 ledger saved on an earlier day
PROMPT_HANDLER = None                      # set by the window; None = console questions
# Keeps the browser serviced while the program waits for the person (6 Oct 2026): with 'Load portal
# pages faster' ticked every request waits for the program's filter, which only answers while the
# program works with the browser - a pure wait held every request (manual mode: nothing clickable,
# no second frame).  Set by core.session.open_session; called about every 0.5 s by every wait.
BROWSER_KEEPALIVE = None

def check_stop():
    if STOP_EVENT.is_set():
        raise StopRequested()

def ask_user(message, kind="continue", defaults=("", ""), poll=None):
    """Ask the person something. In the window this shows a question panel with
    buttons; in console mode it falls back to input()."""
    if PROMPT_HANDLER is not None:
        if poll is None and BROWSER_KEEPALIVE is not None:
            def poll():                        # service the browser; never answers for the person
                BROWSER_KEEPALIVE()
                return False
        return PROMPT_HANDLER(message, kind, defaults, poll)
    return input(message + " ")

# ---------- Run summary (written to Excel at the end of every run) ----------

RESULTS = []            # one dict per item: section, item, status, detail, file
TIMINGS = []            # (tax type, month, s to open Level 3, s to save the file, s back to Level 2)
LEVEL3_SCREEN_CHECK = False         # option B check - OFF since 10.9: live run 6 Oct 2026, 0 of 27 identical
                                    # (the screen has 4 tables, the portal's file 1); with option C the
                                    # portal's file takes 1 s, so the ~1 s check per month is not worth it
SCREEN_CHECK_FILE = "_screen_check.json"
HIDDEN_L3 = True            # option C: Level 3 in a hidden frame, Level 2 stays (no Back per month)
HIDDEN_L3_VERIFY = True     # per tax type the first month is proved against the click route first
HIDDEN_FRAME = "bot_level3"
ROUTE_CHECK_FILE = "_route_check.json"
ROUTE_STATE = {}            # tax type -> {"hidden": bool, "verified": bool}
TIME_LINE_TEXT = ("Tax Position as at",)   # the download time inside every Level 3 - never compared
# 10.16 - the direct way (user, 9 Oct 2026: 'a blend of accuracy and speed ... build it, then we decide'):
L3_ROUTE = "auto"           # 'auto'   direct request, proved per tax type on its first month against the screen way
                            # 'direct' direct request first, no proof (falls back to the screen way on a failure)
                            # 'screen' the screen way only (click / background frame + Export, as in 10.15)
DIRECT_SWITCH_AFTER_S = 0   # > 0: give up the direct request after this many seconds and use the screen way
DIRECT_HEARTBEAT_S = 15     # 'URA still working on it' line after this long, then every 30 s
DIRECT_SAFETY_S = 15 * 60   # x SLOW_FACTOR: no answer at all for this long -> the screen way
DIRECT_STATE = {}           # tax type -> {"on": bool, "verified": bool}
SKIP_NIL = False        # skip Level 3 for months whose Level 2 payment is 0
KEEP_MONTH_FILES = False  # user, 8 Oct 2026: no Excel file per month - each month is a tab in <Tax head>.xlsx
ANNUAL_L3 = False       # Income Tax / Rental: also download Level 3 details
ANNUAL_ALL_YEARS = False  # ...for every year in the ledger, not only the period
PDF_QUIET_SECONDS = 4   # after a certificate PDF arrives, wait this long for more

def record(section, item, status, detail="", file=""):
    RESULTS.append({"Section": section, "Item": item, "Status": status,
                    "Detail": detail, "File": str(file) if file else ""})

def is_nil(amount) -> bool:
    """'0', '0.00', '-' -> True. Anything unreadable -> False (never skip on a guess)."""
    text = str(amount or "").replace(",", "").strip()
    if text in ("-",):
        return True
    try:
        return float(text) == 0
    except ValueError:
        return False

# 1.0 = the old, slow pauses. 0.35 = about a third of the waiting.
SPEED = float(os.getenv("URA_SPEED", "0.35"))
POLL_MS = 150          # how often wait_find looks again

def nap(page, ms):
    """A pause, scaled by SPEED. Never shorter than 120 ms."""
    check_stop()
    page.wait_for_timeout(max(120, int(ms * SPEED)))

def wait_find(page, selectors, timeout=10000, visible=True, frame=None):
    """Keep looking in every frame until the element turns up (150 ms poll),
    instead of looking once and sleeping a fixed time."""
    if isinstance(selectors, str):
        selectors = [selectors]
    end = time.monotonic() + timeout / 1000.0
    while True:
        for fr in ([frame] if frame else page.frames):
            for sel in selectors:
                try:
                    loc = fr.locator(sel).first
                    if loc.count() and (not visible or loc.is_visible()):
                        return fr, loc
                except Exception:
                    continue
        if time.monotonic() >= end:
            return None, None
        check_stop()
        page.wait_for_timeout(POLL_MS)

def js_click(locator, timeout=5000):
    """Click by script - works even when a menu is closed or something is
    on top of the element, and never waits for 'actionability'.  Short time
    limit: on the live portal a missing element used to cost 45 s here."""
    locator.first.evaluate("el => el.click()", timeout=timeout)


def js_click_later(locator, timeout=5000):
    """Script click that returns at once (the click runs a moment later), so a
    slow page function - e.g. the portal's exportTableToExcel - cannot hold up
    the program.  Use inside expect_download / a wait for the next screen."""
    locator.first.evaluate("el => { setTimeout(() => el.click(), 0); }", timeout=timeout)

def app_dir() -> Path:
    """Folder the program lives in - next to the .exe once it is built."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent

OUTPUT_ROOT = None     # set by the suite; None = <program folder>/downloads
SLOW_FACTOR = 1.0      # suite Settings > Internet speed: 1 normal, 2 slow, 4 very slow
ONLINE_WAIT = None     # suite sets a function that waits while the internet is down
BROWSER_CHANNEL = ""   # "" = Chrome, then Edge, then bundled; "chrome" / "msedge" to force
PROXY = ""             # e.g. http://proxy.company.com:8080
NAV_TIMEOUT_MS = 45000 # default wait for page elements
WORK_DIR = None        # where error screenshots go; None = program folder
load_dotenv(app_dir() / ".env")


# ---------- Generic helpers ----------

def find_frame_containing(page, role, name, timeout=60):
    timeout = timeout * SLOW_FACTOR
    deadline = time.time() + timeout
    while time.time() < deadline:
        for frame in page.frames:
            try:
                if frame.get_by_role(role, name=name).count() > 0:
                    return frame
            except Exception:
                continue
        nap(page, 500)
    raise TimeoutError(f"Could not find a frame containing {role} '{name}' within {timeout}s")

def find_frame_containing_selector(page, selector, timeout=30):
    timeout = timeout * SLOW_FACTOR
    deadline = time.time() + timeout
    while time.time() < deadline:
        for frame in page.frames:
            try:
                if frame.locator(selector).count() > 0:
                    return frame
            except Exception:
                continue
        nap(page, 500)
    raise TimeoutError(f"Could not find a frame containing '{selector}' within {timeout}s")

def find_frame_containing_text(page, text, timeout=20):
    timeout = timeout * SLOW_FACTOR
    deadline = time.time() + timeout
    while time.time() < deadline:
        for frame in page.frames:
            try:
                if frame.get_by_text(text, exact=True).count() > 0:
                    return frame
            except Exception:
                continue
        nap(page, 500)
    raise TimeoutError(f"Could not find text '{text}' in any frame within {timeout}s")

DEBUG_DUMPS = os.getenv("URA_DEBUG", "0") == "1"

def work_file(name: str) -> Path:
    """Absolute path in the troubleshooting folder, prefixed with the client TIN."""
    folder = Path(WORK_DIR) if WORK_DIR else app_dir()
    prefix = f"{safe_part(CURRENT_TIN)}_" if CURRENT_TIN else ""
    return folder / f"{prefix}{name}"


def safe_part(text) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(text or "")).strip("._") or "x"


def safe_screenshot(page, name: str, full_page: bool = False):
    """Screenshot for troubleshooting.  Never raises: an error handler must not
    fail because the page it wanted to photograph has already closed (M2)."""
    try:
        path = work_file(name)
        path.parent.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(path), full_page=full_page)
        return path
    except Exception as e:
        print(f"(no screenshot: {type(e).__name__})")
        return None


def debug_dump(page, label):
    """Screenshot + frame list. Off by default - each one costs a second or
    two, and they are only useful when something is broken."""
    if not DEBUG_DUMPS:
        return
    shot = safe_screenshot(page, f"debug_{label}.png")
    print(f"--- Frames after '{label}' ---")
    for i, frame in enumerate(page.frames):
        print(f"  Frame {i}: {frame.url}")
    print(f"Saved {shot}")

def dismiss_cookie_banner(page):
    try:
        btn = page.get_by_role("button", name="Accept All")
        if btn.is_visible(timeout=1500):
            btn.click()
            print("Dismissed cookie banner.")
    except Exception:
        pass

def dismiss_pre_login_popup(page):
    """Close the portal's notice box if it is showing.  UNIVERSAL RULE (6 Oct 2026): a pop-up that
    may not appear is checked at once and closed only if it is showing - never waited for.  The old
    click(timeout=5000) waited 5 s whenever the box was absent (after every Level 3 month); in the
    5 Oct logs the box was closed 10 times and was already showing every time.  A box that appears
    later is still closed by clear_overlays (it is in OVERLAY_SELECTORS) before the next clicks."""
    try:
        btn = page.locator("#process-static-modal").get_by_role("button", name="Close")
        if btn.count() and btn.first.is_visible():
            btn.first.click(timeout=2000)      # it is showing: a normal click, as before
            print("Dismissed pre-login popup.")
    except Exception:
        pass  # quieter now, since this gets called many times per run

OVERLAY_SELECTORS = ["#NINBRNMsg", "#TINBRNMsg", "#parishMsgA", "#process-static-modal",
                     ".ModalPopUpBackgd", ".modal-backdrop"]

def any_overlay(page):
    """Cheap check: is any blocking pop-up actually on screen?"""
    for frame in page.frames:
        for sel in OVERLAY_SELECTORS:
            try:
                loc = frame.locator(sel)
                if loc.count() > 0 and loc.first.is_visible():
                    return True
            except Exception:
                continue
    return False

def clear_overlays(page, quiet=True):
    """Portal modals (NIN/BRN, parish message, cookie/static popups) sit on top
    and swallow every click. Press their Skip/Close button if there is one;
    otherwise hide the element so clicks get through."""
    for frame in page.frames:
        for sel in OVERLAY_SELECTORS:
            try:
                loc = frame.locator(sel)
                if loc.count() == 0 or not loc.first.is_visible():
                    continue
                clicked = False
                for label in ("Skip", "OK", "Close", "Cancel", "\u00d7"):
                    btn = loc.first.get_by_text(label, exact=False)
                    if btn.count() > 0:
                        try:
                            btn.first.click(timeout=2000)
                            print(f"Dismissed overlay {sel} using '{label}'.")
                            clicked = True
                            break
                        except Exception:
                            continue
                if not clicked:
                    loc.first.evaluate(
                        "el => { el.style.display='none'; el.style.visibility='hidden'; }")
                    if not quiet:
                        print(f"Hid blocking overlay {sel}.")
            except Exception:
                continue

def dismiss_popup_if_present(page):
    clear_overlays(page, quiet=False)

def safe_click(page, locator, what="element", timeout=2000):
    """Click that survives overlays. One short real click, then a script
    click - waiting 8 s for 'actionability' on every blocked element was
    costing minutes per run."""
    def gone():
        # Live portal, 3 Oct 2026: a click often WORKS (the next screen opens)
        # but is reported as timed out. If the element has disappeared, the
        # click went through - do not click again (a 2nd click undid ticks and
        # cost 45 s per payment link).
        try:
            return locator.first.count() == 0
        except Exception:
            return False
    try:
        locator.first.click(timeout=timeout)
        return
    except Exception:
        if gone():
            return
    clear_overlays(page)
    try:
        locator.first.click(timeout=timeout)
        return
    except Exception:
        if gone():
            return
    try:
        js_click(locator)
    except Exception as e:
        if not gone():
            print(f"Could not click {what} ({type(e).__name__})")

# ---------- Section 3: Login ----------

LOGIN_URL = "https://ura.go.ug/en/etax-login/"

def open_portal(page):
    """Go to the e-Tax login page. Returns False when we are in fact still
    logged in (the login URL then redirects into the portal)."""
    if not login_form_present(page) and session_alive(page):
        print("Still logged in - no need for the login form.")
        return False
    for attempt in range(1, 4):
        try:
            page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=60000)
            clear_overlays(page)
            if login_form_present(page):
                print("Reached the login page.")
                return True
            if session_alive(page):
                print("Still logged in - no need for the login form.")
                return False
            find_login_frame(page, timeout=90)
            print("Reached the login page.")
            return True
        except Exception as e:
            print(f"Login page not ready (attempt {attempt}/3): {type(e).__name__}")
            if attempt == 2:      # try the long way round once
                try:
                    page.goto("https://ura.go.ug/en/", wait_until="domcontentloaded")
                    clear_overlays(page)
                    page.get_by_role("link", name="Login alert").click()
                    page.get_by_role("link", name="Portal").click()
                    clear_overlays(page)
                except Exception:
                    pass
            nap(page, 3000)
    raise RuntimeError("Could not open the URA login page")

def find_login_frame(page, timeout=90):
    """The frame holding the login form. Looks for the field's id first - the
    portal is slow, and the accessible name is not always there yet."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        for f in page.frames:
            try:
                if f.locator("#userName").count() > 0:
                    return f
                if f.get_by_role("textbox", name="Please enter Login Id.").count() > 0:
                    return f
            except Exception:
                continue
        nap(page, 1000)
    raise TimeoutError(f"The login form did not appear within {timeout}s")

def logged_in(page) -> bool:
    """True once the portal home (menus) has loaded after Log In."""
    for f in page.frames:
        try:
            if "loadhomepagecontent" in (f.url or "").lower():
                return True
            if f.locator("#mailbox").count() > 0:
                return True
        except Exception:
            continue
    return menu_available(page)

LOGIN_ERROR_RE = re.compile(r"invalid|incorrect|wrong|does not match|not valid|try again", re.I)


def login_rejected(page) -> bool:
    """True when the login frame shows an error message or the 'Back' button
    the portal uses after a failed attempt (#loginRespSec)."""
    for f in page.frames:
        try:
            back = f.locator("#loginRespSec input, #loginRespSec > input")
            if back.count() and back.first.is_visible():
                try:
                    back.first.click(timeout=3000)      # return to the form
                except Exception:
                    pass
                return True
            if f.locator("#userName").count():
                txt = f.locator("body").inner_text(timeout=1500)
                if LOGIN_ERROR_RE.search(txt or ""):
                    return True
        except Exception:
            continue
    return False


def fill_login_form(page, tin, password):
    frame = find_login_frame(page)
    user_box = frame.locator("#userName")
    if user_box.count() == 0:
        user_box = frame.get_by_role("textbox", name="Please enter Login Id.")
    pass_box = frame.locator("input[type='password']")
    if pass_box.count() == 0:
        pass_box = frame.get_by_role("textbox", name="Please enter Password.")
    try:
        user_box.first.fill(tin)
    except Exception as e:
        raise RuntimeError(f"Could not type the Login ID into the login form "
                           f"({type(e).__name__} - the box was not ready)") from None
    try:
        pass_box.first.fill(password)
    except Exception as e:
        # Playwright's message repeats the typed text: never let it out (finding H6)
        raise RuntimeError(f"Could not type the password into the login form "
                           f"({type(e).__name__} - the box was not ready)") from None
    print("TIN and password entered.")

def login(page, tin, password, attempts=3):
    # Fall back to the saved credentials ONLY for the same client (a re-login in
    # the middle of that client's job).  Never pair one client's TIN with another
    # client's password (finding H2).
    if not tin:
        tin, password = CREDS["tin"], password or CREDS["password"]
    elif not password and tin == CREDS["tin"]:
        password = CREDS["password"]
    if not tin or not password:
        raise RuntimeError(f"No password for Login ID {tin or '(none)'} - open Clients > Edit "
                           "and type it again")
    if not open_portal(page):
        return                      # session is still good

    for attempt in range(1, attempts + 1):
        fill_login_form(page, tin, password)

        # Carry on as soon as EITHER the button in the window is pressed, OR the
        # person clicks Log In in the browser themselves (login form disappears).
        gone = {"polls": 0}
        def form_gone():
            gone["polls"] = gone["polls"] + 1 if not login_form_present(page) else 0
            return gone["polls"] >= 2          # gone for about a second
        answer = ask_user(
            "In the browser: read the security code and type it in the box. Then press "
            "'Log In' in this window (or click Log In in the browser yourself).",
            kind="login", poll=form_gone)
        if answer == "q":
            raise StopRequested()

        if answer != "done_in_browser":
            login_button_frame = page.frame(name="tcs-resource")
            if login_button_frame is None:
                try:
                    login_button_frame = find_frame_containing(page, "button", "Log In", timeout=20)
                except TimeoutError:
                    login_button_frame = None
            if login_button_frame is not None:
                clear_overlays(page)
                safe_click(page, login_button_frame.get_by_role("button", name="Log In"), "Log In")
                print("Clicked Log In.")
        else:
            print("Log In was clicked in the browser.")

        # Did it work? Wait for the portal menus rather than assuming.
        deadline = time.time() + 60
        while time.time() < deadline:
            if logged_in(page):
                print("Logged in.")
                return
            if login_rejected(page):            # wrong code: re-ask at once, not after 60 s
                print("The portal rejected the login (usually a mistyped security code).")
                break
            nap(page, 1000)
        if login_form_present(page):
            print(f"Login did not go through - the security code was probably wrong "
                  f"(attempt {attempt}/{attempts}).")
            continue
        print("Portal menus not seen yet, but the login form has gone - carrying on.")
        return
    raise RuntimeError("Could not log in after several attempts - check the TIN, password "
                       "and security code")

# ---------- Sections 7 & 8: Tax Statement (Level 1) ----------

def open_tax_statement(page):
    dismiss_cookie_banner(page)
    dismiss_pre_login_popup(page)

    frame = find_frame_containing(page, "heading", "Reports", timeout=90)
    safe_click(page, frame.get_by_role("heading", name="Reports"), "Reports menu")
    dismiss_cookie_banner(page)
    clear_overlays(page)
    safe_click(page, frame.get_by_role("link", name="Tax Statement"), "Tax Statement")

    if any_overlay(page):
        clear_overlays(page)
    tcs_frame = find_frame_containing_selector(page, "#btnSubmit", timeout=40)
    try:
        radio = tcs_frame.locator("#ViewLedger")
        if radio.count() > 0:
            radio.first.evaluate("el => el.click()")   # Ledger Position = Tax Position
    except Exception:
        pass
    submit_btn = tcs_frame.locator("#btnSubmit")
    try:
        submit_btn.evaluate("el => el.click()")
        print("Click executed.")
    except Exception as e:
        print(f"Click reported an issue ({e}), but the page may have already moved on - checking...")

    wait_find(page, ["#firstPageOfAR", "#tblData"], timeout=30000, visible=False)
    debug_dump(page, "after_submit")

def read_level2_table(frame, max_attempts=3):
    """Find the real Level 2 ledger table. Retries a few times in case the page is still rendering."""
    for attempt in range(1, max_attempts + 1):
        frame.wait_for_load_state("domcontentloaded")
        wait_find(frame.page, "table", timeout=8000, frame=frame)

        table_count = frame.locator("table").count()

        for i in range(table_count):
            try:
                locator = frame.locator("table").nth(i)
                locator.wait_for(state="attached", timeout=8000)
                html = locator.evaluate("el => el.outerHTML")
                raw = pd.read_html(io.StringIO(html))[0]
            except Exception:
                continue

            raw = raw.dropna(axis=1, how="all")
            header_row_idx = None
            for ridx in range(len(raw)):
                row_vals = raw.iloc[ridx].astype(str).tolist()
                if "Payment" in row_vals and "Period" in row_vals:
                    header_row_idx = ridx
                    break

            if header_row_idx is None:
                continue

            new_columns = raw.iloc[header_row_idx].tolist()
            clean = raw.iloc[header_row_idx + 1:].copy()
            clean.columns = new_columns
            clean = clean.reset_index(drop=True)
            clean = clean.loc[:, ~clean.columns.duplicated()]
            clean = clean[clean["Period"].astype(str) != "Period"]
            clean = clean.dropna(subset=["Period"])
            return clean

        print(f"Ledger table not readable yet (attempt {attempt}/{max_attempts}) - retrying...")
        nap(frame.page, 2000)

    raise RuntimeError("Could not find the Level 2 ledger table after several attempts")

# ---------- Level 2 / Level 3 download logic ----------

def progress_path(save_folder: Path) -> Path:
    return save_folder / "_progress.json"

def load_progress(save_folder: Path) -> dict:
    f = progress_path(save_folder)
    if f.exists():
        try:
            return json.loads(f.read_text())
        except Exception:
            pass
    return {}

# values that mean "it did NOT succeed"; they are always retried
FAILURE_VALUES = ("generate failed", "download failed", "skipped by user", "failed", "error")
# values that can change on the portal later (the month gets posted, a payment
# arrives): they count as done only on the day they were recorded, so a
# same-day Resume stays fast and a later run looks again (finding H8a)
RECHECK_VALUES = ("not in ledger", "no payment link")


def _today() -> str:
    return datetime.date.today().isoformat()


def mark_progress(save_folder: Path, tax_type: str, key: str, value="done"):
    if str(value).strip().lower() in RECHECK_VALUES:
        value = f"{value} @{_today()}"
    data = load_progress(save_folder)
    data.setdefault(tax_type, {})[key] = value
    save_folder.mkdir(parents=True, exist_ok=True)
    path = progress_path(save_folder)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(data, indent=1), encoding="utf-8")
    os.replace(tmp, path)                  # never a half-written progress file


def already_done(save_folder: Path, tax_type: str, key: str) -> bool:
    value = load_progress(save_folder).get(tax_type, {}).get(key)
    if value is None:
        return False
    v = str(value).strip().lower()
    if v in FAILURE_VALUES:
        return False          # it was recorded, but it did NOT succeed
    for r in RECHECK_VALUES:
        if v == r or v.startswith(r + " @"):
            return v == f"{r} @{_today()}"
    return True


def level2_is_current(path: Path) -> bool:
    """A Level 2 ledger file saved today.  A ledger is a running statement, so a
    copy from an earlier day may be missing later postings (finding H8b)."""
    try:
        saved = datetime.date.fromtimestamp(Path(path).stat().st_mtime)
    except OSError:
        return False
    return saved == datetime.date.today()


def level2_usable(path: Path) -> bool:
    """Can the Level 2 file on disk be used as it is?  (a real workbook, saved today
    or refreshing switched off)"""
    from engines.excel_files import excel_ok
    return excel_ok(path) and (level2_is_current(path) or not REFRESH_L2)


def month_file_done(save_folder: Path, tax_type: str, key: str, target: Path) -> bool:
    """Level 3 already done?  A real workbook on disk, the month inside the tax-head workbook
    (<tax>.xlsx - since 10.12 the monthly file is removed once it is in there), or a progress
    mark that is not 'saved' (nil, not in ledger today...).  'saved' with neither the file nor
    the month in the workbook means: fetch it again."""
    from engines.excel_files import excel_ok
    from engines import ledger_taxhead
    if excel_ok(target):
        return True
    if ledger_taxhead.is_period(key) and ledger_taxhead.canonical(key) in \
            ledger_taxhead.months_in_workbook(save_folder, tax_type):
        return True
    value = str(load_progress(save_folder).get(tax_type, {}).get(key, ""))
    return already_done(save_folder, tax_type, key) and not value.startswith("saved")


def keep_previous_level2(path: Path):
    """Move an older Level 2 file to previous/<name>_saved-<date>.xlsx before re-exporting."""
    path = Path(path)
    if not path.exists():
        return None
    day = datetime.date.fromtimestamp(path.stat().st_mtime).isoformat()
    dest = path.parent / "previous" / f"{path.stem}_saved-{day}{path.suffix}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    n = 2
    while dest.exists():
        dest = dest.with_name(f"{path.stem}_saved-{day}_{n}{path.suffix}")
        n += 1
    os.replace(path, dest)
    print(f"Level 2 from {day} moved to {dest.name} - fetching today's copy.")
    return dest


def drain_notes(section):
    """Turn the notes collected during a step into 'check' items for the Review page."""
    while NOTES:
        item, note = NOTES.pop(0)
        record(section, item, "check", note)

def menu_available(page) -> bool:
    """Is the Reports menu (home-page frame) currently loaded?"""
    for f in page.frames:
        try:
            if f.get_by_role("heading", name="Reports").count() > 0:
                return True
        except Exception:
            continue
    return False

def login_form_present(page) -> bool:
    """Only a visible login box means we actually have to log in again."""
    for f in page.frames:
        try:
            for sel in ("#userName", "input[type='password']"):
                box = f.locator(sel)
                if box.count() > 0 and box.first.is_visible():
                    return True
        except Exception:
            continue
    return False

def save_download(download, folder: Path, filename: str):
    """Save a portal download.  A requested .xlsx is checked by its contents: the portal's
    'Export To Excel' sends an HTML table named .xls, which Excel refuses under a .xlsx
    name (live run 3 Oct 2026) - it is converted to a real workbook with the same text,
    and the portal's own file is kept in portal_original/."""
    from engines import excel_files
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / filename
    if path.suffix.lower() != ".xlsx":
        download.save_as(str(path))
        print(f"Saved: {path}")
        return
    tmp = folder / (filename + ".download")
    download.save_as(str(tmp))
    try:
        how = excel_files.save_portal_excel(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()
    print(f"Saved: {path}" + (" (portal table converted to a real Excel workbook)"
                              if how == "html converted" else ""))

LEVEL_CONTAINERS = [("Level 3", "#levelThreeTwoOfAR"), ("Level 2", "#levelTwoOfAR"),
                    ("Level 1", "#firstPageOfAR")]

def visible_level(page):
    """Which portal section is showing right now (ids taken from the recording)."""
    for f in page.frames:
        for name, selector in LEVEL_CONTAINERS:
            try:
                loc = f.locator(selector)
                if loc.count() > 0 and loc.first.is_visible():
                    return name, f
            except Exception:
                continue

    # Fallback: the container can report as not visible while the page is
    # perfectly usable (this is what caused the false "session was lost").
    for f in page.frames:
        url = f.url or ""
        try:
            if "loadLevelThree" in url:
                return "Level 3", f
            if "loadLevelTwo" in url:
                return "Level 2", f
            if "loadLevelOne" in url and f.locator("#tblData").count() > 0:
                return "Level 1", f
            for name, selector in LEVEL_CONTAINERS:
                if f.locator(selector).count() > 0 and f.locator("#tblData:visible").count() > 0:
                    return name, f
        except Exception:
            continue
    return None, None

def wait_for_level(page, wanted, timeout_seconds=25):
    """Wait until the wanted section (Level 1/2/3) is actually visible. Returns its frame, or None."""
    deadline = time.time() + timeout_seconds * SLOW_FACTOR
    while time.time() < deadline:
        name, f = visible_level(page)
        if name == wanted:
            return f
        check_stop()
        page.wait_for_timeout(POLL_MS)      # 150 ms, not a scaled sleep
    return None

BACK_SELECTORS = ["input#backbtn", "#backbtn", "input[value='Back']",
                  "button:has-text('Back')", "a:has-text('Back')"]

def click_back(page, frame, container=None):
    """Click the Back button of the panel we are actually on.

    The Level 3-2 (Details Of Payment) screen has its own Back next to
    'Export To Excel'. Calling the page's goBack() instead jumps straight out
    to Level 1, which is why the following months could not find their row.
    So: click the real button inside the panel first, JS-click it second, and
    only use goBack() as a last resort.
    """
    clear_overlays(page)
    scopes = []
    if container:
        scopes.append(frame.locator(container))
    scopes.append(frame)

    for scope in scopes:
        for sel in BACK_SELECTORS:
            try:
                loc = scope.locator(sel)
                n = loc.count()
            except Exception:
                continue
            for i in range(n):
                btn = loc.nth(i)
                try:
                    if not btn.is_visible():
                        continue
                except Exception:
                    continue
                try:
                    btn.click(timeout=6000)
                    nap(page, 1200)
                    return True
                except Exception:
                    try:
                        btn.evaluate("el => el.click()")
                        nap(page, 1200)
                        return True
                    except Exception:
                        continue
    try:
        frame.evaluate("() => goBack()")        # last resort
        nap(page, 1500)
        return True
    except Exception:
        return False

MONTH_ROW_XPATH = "xpath=.//tr[td[string-length(normalize-space())=8 and contains(normalize-space(),'-20')]]"

def level2_has_months(frame, timeout_ms=15000):
    """True once the Level 2 ledger is showing its period rows again."""
    return ledger_table(frame, timeout_ms=timeout_ms) is not None

def session_alive(page):
    """True if we are still inside the portal (any level, the menu, or a
    portal frame that is loaded)."""
    if visible_level(page)[0] is not None:
        return True
    for f in page.frames:
        try:
            if f.locator("#mailbox").count() > 0:
                return True
            url = (f.url or "").lower()
            # the login page itself lives under /Portal/ too - don't count it
            if "/portal/" in url and "login" not in url:
                return True
        except Exception:
            continue
    return False

LEVEL1_PATIENCE_S = 90   # x SLOW_FACTOR.  The live Level 1 can take ~100 s (CLAUDE_CONTEXT.md, Version 10
                         # notes).  5 Oct 2026: after 10 s the program left for a route that never worked
                         # and failed 2 minutes later - each time Level 1 was showing 1 s after the failure.


def wait_for_level1_patiently(page, timeout_seconds=LEVEL1_PATIENCE_S):
    """Wait for Level 1 with a heartbeat every 30 s.  Stops early when the login box shows
    (then waiting longer cannot help).  True once Level 1 is showing."""
    t0 = time.time()
    deadline = t0 + timeout_seconds * SLOW_FACTOR
    next_note = 30
    from core.runner import browser_gone
    while time.time() < deadline:
        if browser_gone(page):                 # Chrome closed by hand: waiting cannot help
            return False
        if visible_level(page)[0] == "Level 1":
            return True
        if login_form_present(page):
            return False
        waited = time.time() - t0
        if waited >= next_note:
            print(f"Still waiting for Level 1 to show ({int(waited)} s)...")
            next_note += 30
        check_stop()
        page.wait_for_timeout(500)
    return visible_level(page)[0] == "Level 1"


def recover_to_level1(page, tin, password, timeout_seconds=40):
    """Get back to Level 1 no matter what broke: Back if possible, otherwise
    re-open Tax Statement, logging in again if the portal timed us out."""
    for _ in range(4):                       # L3 -> L2 -> L1, whatever we are on
        level, frame = visible_level(page)
        if level == "Level 1":
            print("Back on Level 1.")
            return
        if level in ("Level 2", "Level 3") and frame is not None:
            container = "#levelThreeTwoOfAR" if level == "Level 3" else "#levelTwoOfAR"
            click_back(page, frame, container=container)
            nap(page, 1500)
            continue
        break
    if wait_for_level1_patiently(page):
        print("Back on Level 1.")
        return

    if login_form_present(page):
        print("Portal has logged us out - logging in again.")
        login(page, tin, password)
        dismiss_cookie_banner(page)
        clear_overlays(page)
    elif not session_alive(page):
        print("Portal page lost - reloading it.")
        page.reload(wait_until="domcontentloaded")
        clear_overlays(page)
        if login_form_present(page):
            login(page, tin, password)

    for attempt in range(1, 4):
        if attempt > 1 or not menu_available(page):
            # The Reports menu lives on the portal home page.  Send the portal frame to its home
            # address (ensure_portal_home - the route that worked every time in the 5 Oct logs);
            # it falls back to Login alert > Portal and logs in again only if the session ended.
            # Never reload the outer page: it is the etax-login address and reloading logs out.
            print("Getting the menu back (portal home page, staying logged in).")
            try:
                ensure_portal_home(page, tin, password)
            except Exception as exc:
                print(f"  The portal home page did not open (attempt {attempt}/3): "
                      f"{(str(exc).splitlines() or [''])[0][:150]}")
                if attempt == 3:
                    raise
                continue
        try:
            open_tax_statement(page)
        except Exception as exc:
            print(f"  The Tax Statement did not open (attempt {attempt}/3): "
                  f"{(str(exc).splitlines() or [''])[0][:150]}")
            if attempt == 3:
                raise
            continue
        if wait_for_level(page, "Level 1", timeout_seconds=timeout_seconds) is not None:
            print("Back on Level 1 (re-opened Tax Statement).")
            return
        print(f"Level 1 still not showing (attempt {attempt}/3) - retrying")
        clear_overlays(page, quiet=False)
    raise RuntimeError(f"Could not reach Level 1 (showing: {visible_level(page)[0]})")

LEVEL1_READY_JS = r"""(names) => {
  const norm = t => (t || '').replace(/[\u00a0\s]+/g, ' ').trim().toLowerCase();
  for (const t of document.querySelectorAll('table')) {
    const txt = norm(t.innerText);
    if (!txt.includes('tax type')) continue;
    if (names.filter(n => txt.includes(n.toLowerCase())).length >= 2) return true;
  }
  return false;
}"""
TAX_LINK_JS = r"""(label) => {
  const norm = t => (t || '').replace(/[\u00a0\s]+/g, ' ').trim().toLowerCase();
  document.querySelectorAll('[data-bot-tax]').forEach(e => e.removeAttribute('data-bot-tax'));
  const want = norm(label);
  for (const el of document.querySelectorAll('a, td')) {
    if (norm(el.innerText) !== want) continue;
    const a = el.matches('a') ? el : (el.querySelector('a') || el);
    a.setAttribute('data-bot-tax', '1');
    return true;
  }
  return false;
}"""
LEVEL1_NAMES = ["Value Added Tax", "IT-WHT", "PAYE", "Income Tax", "Rental", "Local Excise Duty",
                "VAT-WHT", "Stamp Duty"]


def level1_table_frame(page):
    """The frame whose Level 1 'Tax Type' table has finished filling in, or None."""
    for f in page.frames:
        try:
            if f.evaluate(LEVEL1_READY_JS, LEVEL1_NAMES):
                return f
        except Exception:
            continue
    return None


def open_level2(page, tax_type_label, tin=None, password=None):
    """Click a tax type on Level 1 and wait until Level 2 is really showing.
    Live run 3 Oct 2026: after Back, the Level 1 box appeared at once but its
    table of tax types took up to ~100 s to fill in, so a 20 s search for the
    text failed. Now: wait for the TABLE (heartbeat every 30 s), click by
    script, and if it never fills in, reopen the Tax Statement once."""
    dismiss_cookie_banner(page)
    clear_overlays(page)
    if visible_level(page)[0] != "Level 1":     # still inside the previous tax type
        recover_to_level1(page, tin, password)
    for round_ in (1, 2):
        limit = 120 * SLOW_FACTOR
        t0 = time.time()
        next_note = 30
        frame = None
        while time.time() - t0 < limit:
            frame = level1_table_frame(page)
            if frame is not None:
                break
            waited = time.time() - t0
            if waited >= next_note:
                print(f"Still waiting for the Level 1 list of tax types to load ({int(waited)} s)...")
                next_note += 30
            clear_overlays(page)
            nap(page, 1000)
        if frame is not None:
            if not frame.evaluate(TAX_LINK_JS, tax_type_label):
                raise RuntimeError(f"'{tax_type_label}' is not in this taxpayer's Level 1 list")
            frame.locator("[data-bot-tax='1']").first.evaluate("el => { setTimeout(() => el.click(), 0); }",
                                                              timeout=5000)
            l2 = wait_for_level(page, "Level 2", timeout_seconds=45)
            if l2 is not None:
                return l2
            raise RuntimeError(f"Level 2 did not appear for {tax_type_label} "
                               f"(showing: {visible_level(page)[0]})")
        if round_ == 1:
            print("The Level 1 list did not load - opening the Tax Statement again.")
            if ONLINE_WAIT:
                ONLINE_WAIT()
            ensure_portal_home(page, tin, password)
            open_tax_statement(page)
    raise RuntimeError(f"the Level 1 list of tax types did not load (waited {int(limit)} s twice)")


def export_table_to_excel(page, scope, save_folder: Path, filename, table_locator=None,
                          wait_s=30):
    """Press the portal's 'Export To Excel' by script and save the download.
    If no file arrives within wait_s, build the workbook from the table that is
    on screen instead (same rows and columns), so a slow or stuck export never
    costs the ledger.  Returns 'portal export' or 'built from the screen'."""
    clear_overlays(page)
    btn = scope.get_by_role("button", name="Export To Excel")
    try:
        with page.expect_download(timeout=wait_s * 1000) as dl_info:
            js_click_later(btn)
        save_download(dl_info.value, save_folder, filename)
        return "portal export"
    except Exception as e:
        print(f"Export To Excel gave no file within {wait_s} s ({type(e).__name__}) - "
              "building the Excel file from the table on screen.")
    import io
    import pandas as pd
    tbl = table_locator if table_locator is not None else scope.locator("#tblData")
    html = tbl.first.evaluate("t => t.outerHTML", timeout=10000)
    frames = pd.read_html(io.StringIO(html))
    if not frames:
        raise RuntimeError("no table on screen to build the Excel file from")
    target = Path(save_folder) / filename
    target.parent.mkdir(parents=True, exist_ok=True)
    from engines.sheet_safety import neutralise_sheet
    with pd.ExcelWriter(target, engine="openpyxl") as xw:
        frames[0].to_excel(xw, index=False)
        neutralise_sheet(next(iter(xw.sheets.values())))
    print(f"Saved (built from the screen): {target}")
    return "built from the screen"


def export_level2_to_excel(page, frame, save_folder: Path, tax_type_label):
    keep_previous_level2(save_folder / f"{tax_type_label}_L2.xlsx")
    how = export_table_to_excel(page, frame, save_folder, f"{tax_type_label}_L2.xlsx",
                                table_locator=ledger_table(frame) or frame.locator("#tblData"))
    if how == "built from the screen":
        NOTES.append(("Level 2", "built from the on-screen table because the portal's Export "
                                 "gave no file - compare it with the portal"))
    return how

# Column position of "Payment" in the Level 2 table:
# S/N, Period, Return, Assessment, Penalty, Interest, Payment, Net Position
PAYMENT_COLUMN_INDEX = 6

def ledger_table(frame, timeout_ms=20000):
    """The Level 2 ledger table, found by its content (Period + Payment + a
    month row) instead of by id - the id is not reliable across screens."""
    deadline = time.time() + timeout_ms / 1000
    while time.time() < deadline:
        tables = frame.locator("table")
        for i in range(tables.count()):
            t = tables.nth(i)
            try:
                if not t.is_visible():
                    continue
                txt = t.inner_text(timeout=3000)
            except Exception:
                continue
            if "Period" in txt and "Payment" in txt and "-20" in txt:
                return t
        nap(frame.page, 1000)
    return None

MARK_LINK_JS = r"""
(table, args) => {
  const norm = s => (s || '').replace(/\s+/g, ' ').trim();
  table.querySelectorAll('[data-bot-target]').forEach(e => e.removeAttribute('data-bot-target'));
  const rows = [...table.querySelectorAll('tr')];
  const head = rows.find(r => r.querySelectorAll('th').length > 2);
  const heads = head ? [...head.children].map(c => norm(c.innerText).toLowerCase()) : [];
  const col = heads.findIndex(h => h.includes('payment'));
  const target = rows.find(r => [...r.children]
      .some(c => norm(c.innerText).toUpperCase() === args.month));
  if (!target) return {row: false};
  let link = null;
  if (col >= 0 && target.children[col]) link = target.children[col].querySelector('a');
  if (!link && args.amount) {
    link = [...target.querySelectorAll('a')].find(a => norm(a.innerText) === args.amount) || null;
  }
  if (!link) {
    const links = [...target.querySelectorAll('a')];
    link = links.length ? links[links.length - 1] : null;
  }
  if (!link) return {row: true, link: false};
  link.setAttribute('data-bot-target', '1');
  return {row: true, link: true, amount: norm(link.innerText)};
}"""

def mark_payment_link(frame, month_label, amount=None):
    """Tag this month's Payment link in the DOM so it can be clicked directly."""
    table = ledger_table(frame)
    if table is None:
        return {"row": False, "table": False}
    try:
        result = table.evaluate(MARK_LINK_JS, {"month": month_label, "amount": amount})
    except Exception as e:
        print(f"{month_label}: could not read the ledger table ({type(e).__name__})")
        return {"row": False, "table": True}
    result["table"] = True
    return result

ERROR_PAGE_GRACE_S = 8     # an answer without the Level 3-2 box: give it this long, then stop


def find_level3_box(page, month_label, timeout_seconds=None):
    """(frame, text) of the Level 3-2 box, searched in EVERY frame until it shows.
    Live 5 Oct 2026: the program waited 45 s for the box in one frame while the screen showed it
    (screenshots <TIN>_error_PAYE_1.png and <TIN>_error_Value_Added_Tax_2.png) - the frame chosen by
    its address did not hold the box.  A box naming the month wins at once; a visible box that does
    not name it is taken after 5 s (the caller then warns, as before).  Same 45 s patience as before."""
    limit = timeout_seconds if timeout_seconds is not None else NAV_TIMEOUT_MS / 1000.0
    deadline = time.time() + limit
    other_since = None
    error_since, said = None, ""
    while True:
        best = other = None
        error_page = False
        for f in page.frames:
            try:
                box = f.locator("#levelThreeTwoOfAR").first
                if box.count() == 0:
                    # 10.15: the portal answered the Level 3-2 address with a page that has no Level 3-2 box
                    # (an error page) - stop after a short grace instead of waiting the full 45 s
                    if "loadlevelthreetwo" in (f.url or "").lower() and \
                            f.evaluate("() => document.readyState") == "complete":
                        error_page = True
                        said = " ".join(f.evaluate("() => (document.title + ' ' + (document.body ? "
                                                   "document.body.innerText : ''))").split())[:100]
                    continue
                shown = box.is_visible()
                text = box.inner_text(timeout=3000)
            except Exception:
                continue
            if month_label in text:
                if shown:
                    return f, text
                best = best or (f, text)
            elif shown:
                other = other or (f, text)
        if best:
            return best
        if other:
            other_since = other_since or time.time()
            if time.time() - other_since >= 5:
                return other
        if error_page and not other:
            error_since = error_since or time.time()
            if time.time() - error_since >= ERROR_PAGE_GRACE_S:
                raise RuntimeError(f"the portal answered {month_label} without the Level 3-2 page ('{said}')")
        else:
            error_since = None
        if time.time() >= deadline:
            break
        check_stop()
        page.wait_for_timeout(POLL_MS)
    print("Frames right now:")
    for i, f in enumerate(page.frames):
        try:
            print(f"  {i}: {f.url[:120]}")
        except Exception:
            pass
    raise RuntimeError(f"the Level 3-2 box for {month_label} did not appear in any frame "
                       f"within {int(limit)} s")


def check_screen_copy(l3_box, saved_path: Path):
    """Option B (6 Oct 2026): would reading Level 3 from the screen give the same file as the
    portal's Export?  The on-screen box goes through the SAME converter as the portal's file
    (excel_files.html_to_xlsx) into a scratch file outside the client folder, and every cell is
    compared.  The saved file is never changed.  Returns (result, detail)."""
    import tempfile
    from engines import excel_files
    if not saved_path.exists():
        return "not checked", "no portal file"
    first_diff = ""
    for variant, get in (("whole box", lambda: l3_box.first.evaluate("el => el.outerHTML")),
                         ("last table", lambda: l3_box.locator("table").last.evaluate("el => el.outerHTML"))):
        try:
            html = get()
        except Exception as e:
            first_diff = first_diff or f"{variant}: could not read the screen ({type(e).__name__})"
            continue
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td) / "screen.xlsx"
            try:
                excel_files.html_to_xlsx(html, tmp)
            except ValueError as e:
                first_diff = first_diff or f"{variant}: {e}"
                continue
            same, diffs = excel_files.compare_cells(saved_path, tmp, limit=3)
        if same:
            return "identical", variant
        first_diff = first_diff or f"{variant}: " + "; ".join(diffs)
    return "different", first_diff


def _remember_screen_check(save_folder: Path, filename, result, detail):
    path = Path(save_folder) / SCREEN_CHECK_FILE
    try:
        data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    except Exception:
        data = {}
    data[filename] = {"result": result, "detail": detail,
                      "checked": datetime.datetime.now().strftime("%d/%m/%Y %H:%M")}
    try:
        path.write_text(json.dumps(data, indent=1), encoding="utf-8")
    except Exception as e:
        print(f"Could not record the screen-copy check ({type(e).__name__})")


def summarise_level3_checks(save_folder: Path):
    """End-of-step lines: median time per part of a Level 3 month, and the screen-copy results."""
    import statistics
    click = [t for t in TIMINGS if (t[5] if len(t) > 5 else "click") == "click"]
    hidden = [t for t in TIMINGS if len(t) > 5 and t[5] == "background"]
    if click:
        med = [int(round(statistics.median(t[i] for t in click))) for i in (2, 3, 4)]
        print(f"Level 3 months timed: {len(click)} - median: Level 3 opened after {med[0]} s, "
              f"file saved {med[1]} s later, back on Level 2 {med[2]} s later.")
    direct = [t for t in TIMINGS if len(t) > 5 and t[5] == "direct"]
    if direct:
        med = int(round(statistics.median(t[2] for t in direct)))
        print(f"Level 3 pages asked for directly: {len(direct)} - median: URA answered after {med} s.")
    if hidden:
        med = [int(round(statistics.median(t[i] for t in hidden))) for i in (2, 3)]
        print(f"Level 3 months in the background: {len(hidden)} - median: Level 3 opened after {med[0]} s, "
              f"file saved {med[1]} s later (no Back).")
    data = {}
    for path in sorted(Path(save_folder).rglob(SCREEN_CHECK_FILE)):
        try:
            data.update(json.loads(path.read_text(encoding="utf-8")))
        except Exception:
            continue
    path = Path(SCREEN_CHECK_FILE)
    if data:
        same = sum(1 for v in data.values() if v.get("result") == "identical")
        diff = sum(1 for v in data.values() if v.get("result") == "different")
        print(f"Screen-copy check: {same} of {len(data)} Level 3 file(s) identical, {diff} different "
              f"(details: {path.name}).")


# ---------- Option C (6 Oct 2026): Level 3 in a hidden frame, Level 2 stays ----------
# The user's capture of the live Level 2: the Payment link is
#   <a onclick="openLevelThreeTwo('AUG-2026')" style="cursor:pointer;">1,329,127</a>
# and openLevelThreeTwo() points the POST form 'levelTwoOfAR' (hidden fields _csrf, taxType, prm_datefrom,
# prm_dateto, lastEventTS ...) at loadLevelThreeTwo?period=<month> and submits it.  The form normally
# answers into the same frame, replacing Level 2 - hence Back, which makes the portal rebuild Level 2
# (46-68 s per month, 5 Oct 2026 logs).  Here the link's OWN onclick runs with the form's target set to
# a hidden frame: the portal receives exactly what a click sends; only where the answer shows changes.
HIDDEN_SUBMIT_JS = r"""(link, name) => {
  const doc = link.ownerDocument;
  const form = doc.forms['levelTwoOfAR'];
  if (!form || typeof link.onclick !== 'function') return 'the Level 2 form or the link script is not there';
  const old = doc.querySelector('iframe[name="' + name + '"]');
  if (old) old.remove();                                   // a fresh frame each month: nothing stale
  const fr = doc.createElement('iframe');
  fr.name = name;
  fr.style.cssText = 'position:absolute;left:-12000px;top:0;width:1400px;height:900px;border:0';
  doc.body.appendChild(fr);
  const oldTarget = form.getAttribute('target'), oldAction = form.getAttribute('action');
  form.setAttribute('target', name);
  try { link.onclick.call(link); }
  finally {
    if (oldTarget === null) form.removeAttribute('target'); else form.setAttribute('target', oldTarget);
    if (oldAction !== null) form.setAttribute('action', oldAction);
  }
  return 'sent';
}"""


def _hidden_frame(l2_frame):
    """The newest LIVE hidden frame.  Playwright still lists a removed frame (detached) among the
    child frames for a while - found on the test portal: the 2nd month waited on the 1st's dead frame."""
    found = None
    for f in l2_frame.child_frames:
        try:
            if f.name == HIDDEN_FRAME and not f.is_detached():
                found = f
        except Exception:
            continue
    return found


def _remove_hidden(l2_frame):
    try:
        l2_frame.evaluate("n => { const f = document.querySelector('iframe[name=\"' + n + '\"]'); "
                          "if (f) f.remove(); }", HIDDEN_FRAME)
    except Exception:
        pass


def _wait_hidden_level3(page, l2_frame, month_label):
    """(frame, text, why): the hidden frame once it shows this month's Level 3-2.  An answer from the
    portal without the Level 3 box (an error page) ends the wait at once."""
    limit = NAV_TIMEOUT_MS / 1000.0
    deadline = time.time() + limit
    while time.time() < deadline:
        check_stop()
        f = _hidden_frame(l2_frame)
        if f is not None:
            try:
                box = f.locator("#levelThreeTwoOfAR").first
                if box.count():
                    text = box.inner_text(timeout=3000)
                    if month_label in text:
                        return f, text, ""
                elif "loadlevelthreetwo" in (f.url or "").lower() and \
                        f.evaluate("() => document.readyState") == "complete":
                    said = " ".join(f.evaluate("() => (document.title + ' ' + (document.body ? "
                                               "document.body.innerText : ''))").split())[:100]
                    return None, "", f"the portal answered without Level 3 ('{said}')"
            except Exception:
                pass
        page.wait_for_timeout(POLL_MS)
    return None, "", f"no Level 3 for {month_label} within {int(limit)} s"


def download_level3_hidden(page, frame, save_folder: Path, tax_type_label, month_label, payment_value=None,
                           proof=False):
    """Option C for one month.  Returns (level-2 frame, outcome, why): outcome 'saved', 'no link' or
    'fallback' (why says what went wrong; the caller then uses the click route)."""
    clear_overlays(page)
    l2_frame = wait_for_level(page, "Level 2", timeout_seconds=25)
    if l2_frame is None:
        return frame, "fallback", "Level 2 not showing"
    marked = mark_payment_link(l2_frame, month_label, payment_value)
    if not marked.get("row"):
        return l2_frame, "fallback", "row not on the current Level 2 view"
    if not marked.get("link"):
        print(f"{month_label}: no payment link in this row - nothing to open.")
        return l2_frame, "no link", ""
    print(f"{month_label}: opening payment {marked.get('amount')} "
          + ("(proof copy, in the background)" if proof else "(in the background)"))
    t_click = time.monotonic()
    try:
        sent = l2_frame.locator("[data-bot-target='1']").first.evaluate(HIDDEN_SUBMIT_JS, HIDDEN_FRAME)
    except Exception as e:
        sent = f"could not send the form ({type(e).__name__})"
    if sent != "sent":
        return l2_frame, "fallback", sent
    l3_frame, l3_text, why = _wait_hidden_level3(page, l2_frame, month_label)
    if l3_frame is None:
        _remove_hidden(l2_frame)
        return l2_frame, "fallback", why
    t_open = time.monotonic()
    how, t_saved, _ = _save_level3(page, l3_frame.locator("#levelThreeTwoOfAR"), save_folder,
                                   tax_type_label, month_label)
    _remove_hidden(l2_frame)
    opened, saved = t_open - t_click, t_saved - t_open
    if proof:
        print(f"{month_label}: proof copy - Level 3 opened after {opened:.0f} s, saved {saved:.0f} s later")
    else:
        TIMINGS.append((tax_type_label, month_label, opened, saved, 0.0, "background"))
        print(f"{month_label}: times - Level 3 opened after {opened:.0f} s, file saved {saved:.0f} s later "
              f"({how}), Level 2 kept (no Back)")
    return l2_frame, "saved", ""


def refresh_level2(page, frame, tax_type_label):
    """A freshly loaded Level 2 (new security token) - the route the program already uses when a
    month's row is not on the current Level 2 view."""
    recover_to_level1(page, None, None)
    return open_level2(page, tax_type_label)


# ---------- 10.16: the direct way ----------
# The payment link only asks URA for one page: openLevelThreeTwo('<period>') posts the Level 2 form 'levelTwoOfAR'
# to loadLevelThreeTwo?period=<period> (user's captures).  The direct way sends that same request from the Level 2
# page itself (fetch, same session and cookies, same form fields) and keeps the answer off the screen: nothing is
# drawn, no frame to search, no Back, so Level 2 is never rebuilt.  The portal's Export To Excel only copies the
# table #tblData of that page into a file (exportTableToExcel, captured); the direct way takes the same table
# from URA's answer.  Offline test 9 Oct 2026: the user's 4 captured pages identical either way; on the test
# portal identical files; with answers after 40 / 75 / 150 s the direct way took 40 / 75 / 150 s, the screen way
# 117 / 177 / 252 s for the first month.  The wait follows URA's answer, not a clock: it ends when URA answers;
# a clear 'no' (error, refused, login page, no table, another period) switches to the screen way at once.
DIRECT_START_JS = r"""(args) => {
  const form = document.forms['levelTwoOfAR'];
  if (!form) return 'the Level 2 form is not on this page';
  window.__botL3 = window.__botL3 || {};
  const ctl = new AbortController();
  const slot = {done: false, status: 0, text: '', error: '', ctl};
  window.__botL3[args.key] = slot;
  const url = new URL('/Portal/accountReportDetailedController/loadLevelThreeTwo?period=' + args.period,
                      location.href);
  fetch(url.href, {method: 'POST', body: new URLSearchParams(new FormData(form)), credentials: 'same-origin',
                   signal: ctl.signal})
    .then(r => r.text().then(t => { slot.status = r.status; slot.text = t; slot.done = true; }))
    .catch(e => { slot.error = String((e && e.name) || e); slot.done = true; });
  return 'sent';
}"""
# the answer read like the portal's own Export: #tblData of the Level 3-2 page (else the table with 'Sr No.')
DIRECT_POLL_JS = r"""(key) => {
  const s = (window.__botL3 || {})[key];
  if (!s) return null;
  if (!s.done) return {done: false};
  delete window.__botL3[key];
  if (s.error || s.status !== 200) return {done: true, status: s.status, error: s.error};
  const doc = new DOMParser().parseFromString(s.text, 'text/html');
  const box = doc.getElementById('levelThreeTwoOfAR');
  let t = doc.querySelector('#levelThreeTwoOfAR #tblData') || doc.getElementById('tblData');
  if (!t) t = [...doc.querySelectorAll('table')].find(x => !x.querySelector('table') &&
                [...x.querySelectorAll('th')].some(h => h.textContent.trim() === 'Sr No.'));
  const said = ((doc.title || '') + ' ' + (doc.body ? doc.body.innerText || doc.body.textContent : ''))
                 .replace(/\s+/g, ' ').trim().slice(0, 100);
  return {done: true, status: s.status, box: !!box, table: t ? t.outerHTML : '',
          text: box ? box.textContent.replace(/\s+/g, ' ') : '', said,
          login: !!doc.querySelector('input[type=password]')};
}"""
DIRECT_ABORT_JS = r"""(key) => { const s = (window.__botL3 || {})[key];
  if (s && !s.done) s.ctl.abort(); if (window.__botL3) delete window.__botL3[key]; }"""
LINK_PERIOD_RE = re.compile(r"openLevelThreeTwo\(\s*['\"]([^'\"]+)['\"]")


def download_level3_direct(page, l2_frame, save_folder: Path, tax_type_label, label, period_arg):
    """One Level 3-2 by the direct way.  Returns (outcome, why): 'saved' or 'fallback' (why says what went wrong;
    the caller then uses the screen way).  Saves <tax>_L3_<label>.xlsx exactly like a portal export: the table
    is converted by excel_files.save_portal_excel and its HTML kept in portal_original."""
    from engines import excel_files
    key = f"{tax_type_label}|{period_arg}|{time.monotonic()}"
    try:
        sent = l2_frame.evaluate(DIRECT_START_JS, {"key": key, "period": period_arg})
    except Exception as e:
        return "fallback", f"could not send the request ({type(e).__name__})"
    if sent != "sent":
        return "fallback", sent
    t0 = time.monotonic()
    next_note = DIRECT_HEARTBEAT_S
    safety = DIRECT_SAFETY_S * SLOW_FACTOR
    answer = None
    try:
        while True:
            check_stop()
            try:
                answer = l2_frame.evaluate(DIRECT_POLL_JS, key)
            except Exception as e:
                return "fallback", f"the Level 2 page changed while waiting ({type(e).__name__})"
            if answer is None:
                return "fallback", "the request was lost (the Level 2 page changed)"
            if answer.get("done"):
                break
            waited = time.monotonic() - t0
            if DIRECT_SWITCH_AFTER_S and waited >= DIRECT_SWITCH_AFTER_S:
                return "fallback", f"no answer within {DIRECT_SWITCH_AFTER_S} s (switch time set on the Ledgers page)"
            if waited >= safety:
                return "fallback", f"no answer within {int(safety // 60)} min"
            if waited >= next_note:
                print(f"{label}: URA still working on it ({waited:.0f} s)...")
                next_note += 30
            page.wait_for_timeout(300)
    finally:
        if answer is None or not answer.get("done"):
            try:
                l2_frame.evaluate(DIRECT_ABORT_JS, key)
            except Exception:
                pass
    took = time.monotonic() - t0
    if answer.get("error"):
        return "fallback", f"the request did not complete ({answer['error']})"
    if answer.get("status") != 200:
        return "fallback", f"URA answered {answer.get('status')}"
    if answer.get("login"):
        return "fallback", "URA answered with the login page"
    if not answer.get("box") or not answer.get("table"):
        return "fallback", f"the answer has no Level 3-2 table ('{answer.get('said', '')}')"
    if label.upper() not in (answer.get("text") or "").upper():
        return "fallback", f"the answer is not for {label}"
    name = f"{tax_type_label}_L3_{label}.xlsx"
    folder = Path(save_folder)
    folder.mkdir(parents=True, exist_ok=True)
    tmp = folder / (name + ".download")
    tmp.write_bytes(answer["table"].encode("utf-8"))
    try:
        excel_files.save_portal_excel(tmp, folder / name)
    finally:
        tmp.unlink(missing_ok=True)
    TIMINGS.append((tax_type_label, label, took, 0.0, 0.0, "direct"))
    print(f"Saved: {folder / name} (direct way, URA answered after {took:.0f} s)")
    return "saved", ""


SAME_LINES = "same lines, different order"


def compare_level3(a, b):
    """(result, detail): 'identical' (every cell; the download time line ignored), SAME_LINES (the same headings
    and the same lines, listed in another order - Sr No. follows the order), or 'different'.
    Live 9 Oct 2026, VAT FEB-2026: one payment split over 'Original Return' and 'Penalty Assessment' with the same
    date came in the opposite order in two downloads; every value was the same."""
    from openpyxl import load_workbook
    from engines import excel_files
    same, diffs = excel_files.compare_cells(a, b, limit=3, ignore_text=TIME_LINE_TEXT)
    if same:
        return "identical", ""

    def lines(path):
        wb = load_workbook(path, read_only=True)
        try:
            rows = [tuple(r) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)
                and not any(str(v or "").startswith(TIME_LINE_TEXT) for v in r)]
        head = next((i for i, r in enumerate(rows) if any(str(v or "").strip() == "Sr No." for v in r)), None)
        if head is None:
            return None, None
        return rows[:head + 1], sorted(tuple("" if v is None else str(v) for v in r[1:]) for r in rows[head + 1:])
    try:
        (ha, la), (hb, lb) = lines(a), lines(b)
    except Exception:
        ha = hb = la = lb = None
    if ha is not None and ha == hb and la == lb:
        return SAME_LINES, "the same lines and amounts, listed in another order"
    return "different", "; ".join(diffs)


def _with_direct(page, frame, save_folder, tax_type_label, label, find_link, screen_way):
    """The direct way first when it is on and proved for this tax type; otherwise the screen way, and - the first
    time for this tax type - the direct way into a scratch folder, compared cell by cell (download time line
    ignored).  find_link(frame) -> ('link', period) / ('no link', '') / ('no row', '').
    screen_way(frame) -> (frame, outcome).  Returns (frame, outcome)."""
    import tempfile
    from engines import excel_files
    st = DIRECT_STATE.setdefault(tax_type_label, {"on": L3_ROUTE in ("auto", "direct"),
                                                  "verified": L3_ROUTE == "direct"})
    if st["on"] and st["verified"]:
        l2 = wait_for_level(page, "Level 2", timeout_seconds=25) or frame
        found, period_arg = find_link(l2)
        if found == "no link":
            print(f"{label}: no payment link - nothing to open.")
            return l2, "no link"
        if found == "link":
            print(f"{label}: asking URA for the Level 3-2 directly")
            outcome, why = download_level3_direct(page, l2, save_folder, tax_type_label, label, period_arg)
            if outcome == "saved":
                return l2, "saved"
            st["on"] = False
            print(f"{label}: the direct way did not work ({why}) - the screen way for the rest of {tax_type_label}.")
            _remember_route_check(save_folder, f"{tax_type_label} - direct", label, "switched off", why)
            frame = refresh_level2(page, l2, tax_type_label)       # a fresh Level 2 (security token) first
        else:
            frame = l2
    frame, outcome = screen_way(frame)
    if st["on"] and not st["verified"] and outcome == "saved":
        st["verified"] = True                                       # proved once per tax type, whatever the result
        name = f"{tax_type_label}_L3_{label}.xlsx"
        l2 = wait_for_level(page, "Level 2", timeout_seconds=25) or frame
        found, period_arg = find_link(l2)
        with tempfile.TemporaryDirectory() as td:
            if found != "link":
                check, why = "fallback", "the payment link was not on Level 2 again"
            else:
                check, why = download_level3_direct(page, l2, Path(td), tax_type_label, label, period_arg)
            if check == "saved":
                result, detail = compare_level3(Path(save_folder) / name, Path(td) / name)
            else:
                result, detail = "not possible", why
        _remember_route_check(save_folder, f"{tax_type_label} - direct", label, result, detail)
        if result in ("identical", SAME_LINES):
            print(f"{label}: direct way {result} to the screen way's file - the other {tax_type_label} "
                  f"Level 3 pages are asked for directly.")
        else:
            st["on"] = False
            print(f"{label}: direct way {result} ({detail}) - the screen way for the rest of {tax_type_label}.")
            if result == "different":
                NOTES.append((f"{label} - direct way proof", f"direct way different from the portal's Export: {detail}"))
        frame = l2
    return frame, outcome


def _route_check_path(save_folder) -> Path:
    """Ledgers\\<TIN>\\_route_check.json.  level3_month is given the _months\\<tax> folder; until 10.12
    the file was written there (and kept that folder from being tidied away)."""
    folder = Path(save_folder)
    if folder.parent.name == "_months":
        folder = folder.parent.parent
    return folder / ROUTE_CHECK_FILE


def move_route_check_up(month_dir):
    """A _route_check.json left inside _months\\<tax> by 10.9-10.11 is merged into the client folder's."""
    old = Path(month_dir) / ROUTE_CHECK_FILE
    if not old.exists():
        return
    new = _route_check_path(month_dir)
    try:
        data = json.loads(new.read_text(encoding="utf-8")) if new.exists() else {}
        for tax, entry in json.loads(old.read_text(encoding="utf-8")).items():
            data.setdefault(tax, entry)
        new.write_text(json.dumps(data, indent=1), encoding="utf-8")
        old.unlink()
    except Exception as e:
        print(f"Could not move {old.name} to {new.parent} ({type(e).__name__})")


def _remember_route_check(save_folder, tax_type_label, month_label, result, detail):
    path = _route_check_path(save_folder)
    try:
        data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    except Exception:
        data = {}
    data[tax_type_label] = {"month": month_label, "result": result, "detail": detail,
                            "checked": datetime.datetime.now().strftime("%d/%m/%Y %H:%M")}
    try:
        path.write_text(json.dumps(data, indent=1), encoding="utf-8")
    except Exception as e:
        print(f"Could not record the route check ({type(e).__name__})")


def level3_month(page, frame, save_folder: Path, tax_type_label, month_label, payment_value=None):
    """10.16: the direct way first (proved per tax type on its first month), else the screen way below."""
    def find_link(l2):
        marked = mark_payment_link(l2, month_label, payment_value)
        if not marked.get("row"):
            return "no row", ""
        if not marked.get("link"):
            return "no link", ""
        try:
            onclick = l2.locator("[data-bot-target='1']").first.get_attribute("onclick", timeout=3000) or ""
        except Exception:
            onclick = ""
        m = LINK_PERIOD_RE.search(onclick)
        return "link", (m.group(1) if m else month_label)
    return _with_direct(page, frame, save_folder, tax_type_label, month_label, find_link,
                        lambda fr: level3_month_screen(page, fr, save_folder, tax_type_label, month_label,
                                                       payment_value))


def level3_month_screen(page, frame, save_folder: Path, tax_type_label, month_label, payment_value=None):
    """One month's Level 3 by the best proven screen route.  Per tax type: the first month by the click
    route (today's file), then the same month in the background into a scratch folder, compared
    cell by cell except the download-time line.  Identical -> the other months in the background;
    otherwise, or if the portal ever refuses the background request, the rest of that tax type by
    the click route (Level 2 reloaded first, so its security token is fresh).  No month is lost."""
    import tempfile
    from engines import excel_files
    st = ROUTE_STATE.setdefault(tax_type_label, {"hidden": HIDDEN_L3, "verified": not HIDDEN_L3_VERIFY})
    if st["hidden"] and st["verified"]:
        l2_frame, outcome, why = download_level3_hidden(page, frame, save_folder, tax_type_label,
                                                        month_label, payment_value)
        if outcome != "fallback":
            return l2_frame, outcome
        st["hidden"] = False
        print(f"{month_label}: background Level 3 did not work ({why}) - the usual click-and-Back way "
              f"for the rest of {tax_type_label}.")
        frame = refresh_level2(page, l2_frame, tax_type_label)
    frame, outcome = download_level3_for_month(page, frame, save_folder, tax_type_label, month_label,
                                               payment_value=payment_value)
    if st["hidden"] and not st["verified"] and outcome == "saved":
        st["verified"] = True
        name = f"{tax_type_label}_L3_{month_label}.xlsx"
        with tempfile.TemporaryDirectory() as td:
            frame, check, why = download_level3_hidden(page, frame, Path(td), tax_type_label,
                                                       month_label, payment_value, proof=True)
            if check == "saved":
                # 10.16: the same lines in another order also count (live VAT FEB-2026, 9 Oct 2026)
                result, detail = compare_level3(Path(save_folder) / name, Path(td) / name)
            else:
                result, detail = "not possible", why
        _remember_route_check(save_folder, tax_type_label, month_label, result, detail)
        if result in ("identical", SAME_LINES):
            print(f"{month_label}: background copy {result} to the click route's file - the other "
                  f"{tax_type_label} months go in the background (no Back).")
        else:
            st["hidden"] = False
            print(f"{month_label}: background copy {result} ({detail}) - the usual click-and-Back way "
                  f"for the rest of {tax_type_label}.")
            if check == "fallback":
                frame = refresh_level2(page, frame, tax_type_label)
    return frame, outcome


def download_level3_for_month(page, frame, save_folder: Path, tax_type_label, month_label,
                              payment_value=None):
    """Level 2 -> this month's Payment link -> Level 3 -> Export To Excel -> Back.
    Returns (frame, outcome): outcome is 'saved', 'no link' (nothing to open),
    or 'failed' (retried on the next run)."""
    clear_overlays(page)
    dismiss_cookie_banner(page)

    l2_frame = wait_for_level(page, "Level 2", timeout_seconds=25)
    if l2_frame is None:
        raise RuntimeError(f"Level 2 not showing before {month_label} (showing: {visible_level(page)[0]})")

    marked = mark_payment_link(l2_frame, month_label, payment_value)
    if not marked.get("row"):
        print(f"{month_label}: row not on the current Level 2 view - reloading the ledger")
        recover_to_level1(page, None, None)
        l2_frame = open_level2(page, tax_type_label)
        marked = mark_payment_link(l2_frame, month_label, payment_value)
        if not marked.get("row"):
            print(f"{month_label}: still no row - will retry on the next run.")
            return l2_frame, "failed"
    if not marked.get("link"):
        print(f"{month_label}: no payment link in this row - nothing to open.")
        return l2_frame, "no link"
    print(f"{month_label}: opening payment {marked.get('amount')}")
    link = l2_frame.locator("[data-bot-target='1']")
    t_click = time.monotonic()

    l3_frame = None
    for click_attempt in range(1, 4):
        try:
            js_click_later(link)          # one click, never repeated by a fallback
        except Exception:
            if link.count() == 0:         # already gone = Level 3 is opening
                pass
            else:
                safe_click(page, link, f"payment link for {month_label}")
        l3_frame = wait_for_level(page, "Level 3", timeout_seconds=20)
        if l3_frame is not None:
            break
        print(f"{month_label}: Level 3 did not open (attempt {click_attempt}/3) - retrying the click")
        nap(page, 2000)
        clear_overlays(page)
        l2_frame = wait_for_level(page, "Level 2", timeout_seconds=15) or l2_frame
        again = mark_payment_link(l2_frame, month_label, payment_value)
        if not again.get("link"):
            break
        link = l2_frame.locator("[data-bot-target='1']")

    if l3_frame is None:
        print(f"{month_label}: Level 3 would not open - will retry on the next run "
              f"(showing: {visible_level(page)[0]}).")
        return (wait_for_level(page, "Level 2", timeout_seconds=15) or l2_frame), "failed"

    l3_frame, l3_text = find_level3_box(page, month_label)      # the frame that really shows it
    t_open = time.monotonic()
    if month_label not in l3_text:
        print(f"WARNING: Level 3 page does not mention {month_label} - check this file manually.")
        NOTES.append((month_label, "the Level 3 page did not mention this month - check the file"))

    l3_box = l3_frame.locator("#levelThreeTwoOfAR")
    how, t_saved, t_checked = _save_level3(page, l3_box, save_folder, tax_type_label, month_label)

    nap(page, 1500)
    dismiss_cookie_banner(page)
    dismiss_pre_login_popup(page)

    new_frame = None
    for back_try in range(1, 4):
        click_back(page, l3_frame, container="#levelThreeTwoOfAR")
        new_frame = wait_for_level(page, "Level 2", timeout_seconds=10)
        if new_frame is not None:
            break
        print(f"{month_label}: still on Level 3 - pressing Back again ({back_try}/3)")
        l3_frame = wait_for_level(page, "Level 3", timeout_seconds=5) or l3_frame
    if new_frame is None:
        raise RuntimeError(f"Did not return to Level 2 after {month_label} (showing: {visible_level(page)[0]})")
    for _ in range(10):                       # the ledger redraws slowly
        if level2_has_months(new_frame, timeout_ms=3000):
            break
        nap(page, 1000)
    t_back = time.monotonic()
    opened, saved, back = t_open - t_click, t_saved - t_open, t_back - t_checked
    TIMINGS.append((tax_type_label, month_label, opened, saved, back, "click"))
    print(f"{month_label}: times - Level 3 opened after {opened:.0f} s, file saved {saved:.0f} s later "
          f"({how}), back on Level 2 {back:.0f} s later")
    return new_frame, "saved"


def _save_level3(page, l3_box, save_folder, tax_type_label, month_label):
    """Save one Level 3 with the portal's Export (unchanged) and run the screen-copy check.
    Returns (how, time saved, time checked)."""
    l3_name = f"{tax_type_label}_L3_{month_label}.xlsx"
    how = export_table_to_excel(page, l3_box, save_folder, l3_name,
                                table_locator=l3_box.locator("table").last)
    t_saved = time.monotonic()
    if how == "built from the screen":
        NOTES.append((month_label, "Level 3 built from the on-screen table because the portal's "
                                   "Export gave no file - compare it with the portal"))
    elif LEVEL3_SCREEN_CHECK:
        try:
            result, detail = check_screen_copy(l3_box, Path(save_folder) / l3_name)
        except Exception as e:
            result, detail = "not checked", f"{type(e).__name__}"
        _remember_screen_check(save_folder, l3_name, result, detail)
        if result == "identical":
            print(f"{month_label}: screen copy identical to the portal's file ({detail}).")
        else:
            print(f"{month_label}: screen copy {result.upper()} from the portal's file - {detail}")
    t_checked = time.monotonic()
    return how, t_saved, t_checked

# ---------- Yearly tax heads (Income Tax, Rental) - 10.15
# User's captures of 8 Oct 2026 (Income Tax and Rental Level 2, Income Tax Level 3-2):
#   Level 2 has TWO heading rows: Return [Provisional, Final], Payment [PRN, Tax Credit], ...  Every amount is a
#   link: Return/Assessment/Penalty -> openLevelThreeOne('JUL-2026 to JUN-2027'), Interest -> openLevelThreeFive,
#   PRN AND Tax Credit -> openLevelThreeTwo (the SAME Level 3-2 page for both), Amount Allocated ->
#   openLevelThreeThree, Tax waiver -> openLevelThreeFour.  Most links read "0".
#   Level 3-2 of a year is the monthly page: form#levelThreeTwoOfAR, table#tblData ('Period: JUL-2026 to
#   JUN-2027'), lines 'Payment (PRN)', 'WHT Credit', 'ASYCUDA Credit'; <button> Export To Excel; Back.
# Until 10.14 the program picked links by heading position; with two heading rows the 'Payment' heading sits
# over the Interest cell, so it opened Level 3-5 (Interest) and waited for an Export button that page does not
# have - live 8 Oct 2026, 22:33 and 22:35: 'Locator.click: Timeout 30000ms exceeded', 3 tries each, the whole
# of Income Tax and Rental failed.  Now: the openLevelThreeTwo links only, found by what they call; one Level
# 3-2 per year; a year whose PRN and Tax Credit are both 0 is not opened; one year's failure is recorded for
# that year and the next year goes on.

ANNUAL_LINKS_JS = r"""
() => {
  const norm = s => (s || '').replace(/\s+/g, ' ').trim();
  document.querySelectorAll('[data-bot-annual]').forEach(e => e.removeAttribute('data-bot-annual'));
  const years = [];
  document.querySelectorAll('a[onclick*="openLevelThreeTwo"]').forEach(a => {
    const m = /openLevelThreeTwo\(\s*['"]([^'"]+)['"]/.exec(a.getAttribute('onclick') || '');
    if (!m) return;
    const period = norm(m[1]);
    let y = years.find(x => x.period === period);
    if (!y) { y = {period, amounts: [], n: -1}; years.push(y); }
    const amount = norm(a.innerText);
    y.amounts.push(amount);
    if (y.n < 0 && /[1-9]/.test(amount)) {        // not "0": there is something to open
      y.n = years.indexOf(y);
      a.setAttribute('data-bot-annual', String(y.n));
    }
  });
  return years;
}"""


def annual_links(frame, wait_s=15):
    """[{period, amounts: [PRN, Tax Credit], n}] - one entry per year (n = -1: every amount is 0).
    Waits for the links: the page may still be redrawing (the Export button is a <button> inside the
    form, so pressing it also reloads Level 2 - same in the test copy of the captured page)."""
    deadline = time.time() + wait_s
    error = None
    while True:
        try:
            found = frame.evaluate(ANNUAL_LINKS_JS) or []
            if found:
                return found
            error = None
        except Exception as e:
            error = e
        if time.time() >= deadline:
            break
        check_stop()
        try:
            frame.page.wait_for_timeout(500)
        except Exception:
            time.sleep(0.5)
    if error is not None:
        print(f"Could not read the yearly ledger ({type(error).__name__})")
    return []


def annual_detail(item):
    """'Payment 350,000; Tax credit 460,000' - the PRN link comes before the Tax Credit link in every
    row of both captures (headings PRN | Tax Credit)."""
    a = item.get("amounts") or []
    parts = [f"Payment {a[0]}"] if a else []
    if len(a) > 1:
        parts.append(f"Tax credit {a[1]}")
    return "; ".join(parts)


def year_in_period(label, period_from, period_to) -> bool:
    """True when the year overlaps the period.  'JUL-2025 to JUN-2026' against 2025-07..2026-06 is in,
    'JUL-2024 to JUN-2025' is out (until 10.14 any shared calendar year counted, so 3 years were fetched
    for a 1-year period).  A label that is not a year period is kept (never dropped on a guess)."""
    from engines import ledger_taxhead
    text = ledger_taxhead.canonical(label)
    if " to " not in text:
        years = [int(y) for y in re.findall(r"(?:19|20)\d{2}", label or "")]
        if not years:
            return True
        lo, hi = int(period_from[:4]), int(period_to[:4])
        return any(lo <= y <= hi for y in years)
    a, b = (datetime.datetime.strptime(x, "%b-%Y") for x in text.split(" to "))
    p_from = datetime.datetime.strptime(period_from, "%Y-%m")
    p_to = datetime.datetime.strptime(period_to, "%Y-%m")
    return a <= p_to and b >= p_from


def level3_year(page, frame, month_dir: Path, tax_type_label, period):
    """10.16: the direct way first (proved per tax type on its first year), else download_level3_year."""
    def find_link(l2):
        item = next((x for x in annual_links(l2) if x["period"] == period), None)
        if item is None:
            return "no row", ""
        return ("link", period) if item["n"] >= 0 else ("no link", "")
    return _with_direct(page, frame, month_dir, tax_type_label, period, find_link,
                        lambda fr: download_level3_year(page, fr, month_dir, tax_type_label, period))


def save_failure_copy(page, stem):
    """A screenshot and a copy of every frame, so a failure can be looked at later (troubleshooting folder)."""
    saved = []
    shot = safe_screenshot(page, f"{stem}.png")
    if shot:
        saved.append(shot.name)
    for i, f in enumerate(page.frames):
        try:
            html = f.content()
        except Exception:
            continue
        if not html.strip() or (f.url or "about:blank") == "about:blank":
            continue
        path = work_file(f"{stem}_frame{i}.html")
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(html, encoding="utf-8")
            saved.append(path.name)
        except Exception:
            pass
    return saved


def download_level3_year(page, frame, month_dir: Path, tax_type_label, period):
    """Level 2 -> this year's PRN / Tax Credit link -> Level 3-2 -> Export To Excel (screen copy if the
    portal gives no file) -> Back.  Returns (level-2 frame, outcome): 'saved', 'no link' (every amount 0)
    or 'failed'.  May raise; the caller records the year as failed and goes on."""
    clear_overlays(page)
    dismiss_cookie_banner(page)
    l2_frame = wait_for_level(page, "Level 2", timeout_seconds=25) or frame
    item = next((x for x in annual_links(l2_frame) if x["period"] == period), None)
    if item is None:
        print(f"{period}: not on the current Level 2 view - reloading the ledger")
        recover_to_level1(page, None, None)
        l2_frame = open_level2(page, tax_type_label)
        item = next((x for x in annual_links(l2_frame) if x["period"] == period), None)
        if item is None:
            print(f"{period}: still not on Level 2 - will retry.")
            return l2_frame, "failed"
    if item["n"] < 0:
        print(f"{period}: payment and tax credit are 0 - nothing to open.")
        return l2_frame, "no link"
    print(f"{period}: opening the payments and tax credits ({annual_detail(item)})")
    link = l2_frame.locator(f"[data-bot-annual='{item['n']}']")
    t_click = time.monotonic()
    l3_frame = None
    for click_attempt in range(1, 4):
        try:
            js_click_later(link)                       # one click, never repeated by a fallback
        except Exception:
            if link.count():
                safe_click(page, link, f"payment link for {period}")
        l3_frame = wait_for_level(page, "Level 3", timeout_seconds=20)
        if l3_frame is not None:
            break
        print(f"{period}: Level 3 did not open (attempt {click_attempt}/3) - retrying the click")
        nap(page, 2000)
        clear_overlays(page)
        l2_frame = wait_for_level(page, "Level 2", timeout_seconds=15) or l2_frame
        item = next((x for x in annual_links(l2_frame) if x["period"] == period), None)
        if not item or item["n"] < 0:
            break
        link = l2_frame.locator(f"[data-bot-annual='{item['n']}']")
    if l3_frame is None:
        print(f"{period}: Level 3 would not open (showing: {visible_level(page)[0]}).")
        return (wait_for_level(page, "Level 2", timeout_seconds=15) or l2_frame), "failed"
    l3_frame, l3_text = find_level3_box(page, period)       # the frame that really shows it
    t_open = time.monotonic()
    if period not in l3_text:
        print(f"WARNING: Level 3 page does not mention {period} - check this file manually.")
        NOTES.append((period, "the Level 3 page did not mention this year - check the file"))
    how, t_saved, t_checked = _save_level3(page, l3_frame.locator("#levelThreeTwoOfAR"), month_dir,
                                           tax_type_label, period)
    nap(page, 1500)
    new_frame = None
    for back_try in range(1, 4):
        current = wait_for_level(page, "Level 3", timeout_seconds=5) or l3_frame
        click_back(page, current, container="#levelThreeTwoOfAR")
        new_frame = wait_for_level(page, "Level 2", timeout_seconds=15)
        if new_frame is not None:
            break
        print(f"{period}: still on Level 3 - pressing Back again ({back_try}/3)")
    if new_frame is None:
        raise RuntimeError(f"Did not return to Level 2 after {period} (showing: {visible_level(page)[0]})")
    TIMINGS.append((tax_type_label, period, t_open - t_click, t_saved - t_open, time.monotonic() - t_checked,
                    "click"))
    print(f"{period}: Level 3 opened after {t_open - t_click:.0f} s, file saved {t_saved - t_open:.0f} s later "
          f"({how})")
    return new_frame, "saved"


def _back_to_level1_after(page, tax_type_label, tin, password):
    """After a tax type is done: back to Level 1.  If that fails the tax type is still done
    (5 Oct 2026, run 17:11: VAT was finished, the way back failed, and all of VAT was attempted
    again); the next tax type opens the Tax Statement itself."""
    try:
        recover_to_level1(page, tin, password)
    except Exception as exc:
        print(f"Could not return to Level 1 after {tax_type_label} "
              f"({(str(exc).splitlines() or [''])[0][:150]}) - the next tax type opens the Tax Statement again.")


def process_annual_tax_type(page, tax_type_label, save_folder: Path, tin=None, password=None,
                            period_from=None, period_to=None):
    """Income Tax / Rental: Level 2 export; optionally Level 3-2 of each year (one page per year).
    Every year goes into <Tax head>.xlsx (Summary, Level 3-2, a tab per year) like the monthly tax heads.
    Years already saved are skipped, so a second run (or the second try at the end of a client) fetches
    only what is missing."""
    from engines import ledger_taxhead
    section = f"{tax_type_label} ledger"
    l2_file = save_folder / f"{tax_type_label}_L2.xlsx"
    month_dir = ledger_taxhead.months_folder(save_folder, tax_type_label)
    workbook = save_folder / f"{tax_type_label}.xlsx"
    if level2_usable(l2_file) and not ANNUAL_L3:
        print(f"{tax_type_label}: already downloaded today - skipping this tax type.")
        record(section, "Level 2", "already downloaded", file=l2_file)
        return

    frame = open_level2(page, tax_type_label, tin, password)
    if level2_usable(l2_file):
        record(section, "Level 2", "already downloaded", file=l2_file)
    else:
        how = export_level2_to_excel(page, frame, save_folder, tax_type_label)
        mark_progress(save_folder, tax_type_label, "L2", "saved")
        record(section, "Level 2", "saved", detail=how, file=l2_file)
        drain_notes(section)

    if ANNUAL_L3:
        frame = wait_for_level(page, "Level 2", timeout_seconds=25) or frame
        items = annual_links(frame)
        if not ANNUAL_ALL_YEARS and period_from:
            kept = [x for x in items if year_in_period(x["period"], period_from, period_to)]
            print(f"{tax_type_label}: {len(items)} year(s) in the ledger, {len(kept)} inside the period.")
            items = kept
        else:
            print(f"{tax_type_label}: {len(items)} year(s) in the ledger (all years).")
        for item in items:
            period, detail = item["period"], annual_detail(item)
            target = month_dir / f"{tax_type_label}_L3_{period}.xlsx"
            if month_file_done(save_folder, tax_type_label, period, target):
                print(f"{period}: already downloaded - skipping.")
                record(section, period, "already downloaded", detail=detail, file=workbook)
                continue
            try:
                frame, outcome = level3_year(page, frame, month_dir, tax_type_label, period)
                why = ""
            except StopRequested:
                raise
            except Exception as e:
                from core.runner import browser_gone
                if browser_gone(page, e):
                    raise
                why = (str(e).splitlines() or [""])[0][:150]
                outcome = "failed"
                copies = save_failure_copy(page, f"error_{safe_part(tax_type_label)}_{safe_part(period)}")
                print(f"{period}: failed - {why}" + (f" (page copy: {', '.join(copies)})" if copies else ""))
                try:                                   # back to a fresh Level 2 for the next year
                    recover_to_level1(page, tin, password)
                    frame = open_level2(page, tax_type_label, tin, password)
                except StopRequested:
                    raise
                except Exception as rec:
                    raise RuntimeError(f"{period} failed ({why}) and Level 2 could not be opened again "
                                       f"({(str(rec).splitlines() or [''])[0][:100]})") from rec
            status = {"saved": "saved", "no link": "no payment link"}.get(outcome, "failed")
            mark_progress(save_folder, tax_type_label, period, status)
            record(section, period, status, detail=f"{detail}; {why}" if why else detail,
                   file=(month_dir / target.name if KEEP_MONTH_FILES else workbook) if status == "saved" else "")
            drain_notes(section)
    print(f"=== {tax_type_label}: finished ===")
    _back_to_level1_after(page, tax_type_label, tin, password)

def generate_month_labels(period_from: str, period_to: str):
    start = datetime.datetime.strptime(period_from, "%Y-%m")
    end = datetime.datetime.strptime(period_to, "%Y-%m")
    labels = []
    current = start
    while current <= end:
        labels.append(current.strftime("%b-%Y").upper())
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1)
        else:
            current = current.replace(month=current.month + 1)
    return labels

def get_current_period():
    """period_to = current month. period_from = 17 months earlier -> 18 months inclusive."""
    today = datetime.date.today()
    total_months = today.year * 12 + (today.month - 1) - 17
    from_year = total_months // 12
    from_month = total_months % 12 + 1
    period_from = f"{from_year:04d}-{from_month:02d}"
    period_to = f"{today.year:04d}-{today.month:02d}"
    return period_from, period_to

def process_tax_type(page, tax_type_label, period_from, period_to, save_folder: Path,
                     tin=None, password=None):
    """Level 2 export, then Level 3 for every month of the period that the ledger
    lists (clicking Payment in that row), then back to Level 1."""
    section = f"{tax_type_label} ledger"
    wanted_months = generate_month_labels(period_from, period_to)
    l2_file = save_folder / f"{tax_type_label}_L2.xlsx"
    from engines import ledger_taxhead
    month_dir = ledger_taxhead.months_folder(save_folder, tax_type_label)   # the portal's monthly files
    if ledger_taxhead.move_old_month_files(save_folder, tax_type_label):
        print(f"{tax_type_label}: monthly Level 3 files moved to {month_dir}")
    move_route_check_up(month_dir)
    if level2_usable(l2_file) and all(
            month_file_done(save_folder, tax_type_label, m, month_dir / f"{tax_type_label}_L3_{m}.xlsx")
            for m in wanted_months):
        print(f"{tax_type_label}: everything already downloaded - skipping this tax type.")
        record(section, "Level 2 and all months", "already downloaded")
        return

    frame = open_level2(page, tax_type_label, tin, password)
    df = read_level2_table(frame)
    listed = [str(p) for p in df["Period"] if re.fullmatch(r"[A-Z]{3}-\d{4}", str(p))]
    print(f"{tax_type_label}: ledger lists {len(listed)} period(s)"
          + (f", {listed[0]} to {listed[-1]}" if listed else " (no data)") + ".")

    if level2_usable(l2_file):
        record(section, "Level 2", "already downloaded", file=l2_file)
    else:
        how = export_level2_to_excel(page, frame, save_folder, tax_type_label)
        record(section, "Level 2", "saved", detail=how, file=l2_file)
        drain_notes(section)

    months_to_process = [m for m in listed if m in wanted_months]
    payments = {str(r["Period"]): str(r.get("Payment", "")).strip() for _, r in df.iterrows()}
    print(f"Months in the period: {', '.join(months_to_process) or 'none'}")

    for month in wanted_months:
        if month not in months_to_process:
            record(section, month, "not in ledger")
            if not already_done(save_folder, tax_type_label, month):
                mark_progress(save_folder, tax_type_label, month, "not in ledger")

    current_frame = frame
    workbook = save_folder / f"{tax_type_label}.xlsx"     # where every month ends up (tab + table rows)
    for period_text in months_to_process:
        target = month_dir / f"{tax_type_label}_L3_{period_text}.xlsx"
        shown = target if KEEP_MONTH_FILES else workbook   # the run summary links to what will exist
        pay = payments.get(period_text, "")
        if month_file_done(save_folder, tax_type_label, period_text, target):
            print(f"{period_text}: already downloaded - skipping.")
            record(section, period_text, "already downloaded", detail=f"Payment {pay}",
                   file=shown if (target.exists() or workbook.exists()) else "")
            continue
        if SKIP_NIL and is_nil(pay):
            print(f"{period_text}: payment is nil - skipped (option ticked).")
            record(section, period_text, "nil payment - skipped", detail=f"Payment {pay}")
            continue          # not written to progress, so an unticked run still fetches it
        current_frame, outcome = level3_month(
            page, current_frame, month_dir, tax_type_label, period_text, payment_value=pay)
        mark_progress(save_folder, tax_type_label, period_text,
                      {"saved": "saved", "no link": "no payment link"}.get(outcome, "failed"))
        record(section, period_text, {"saved": "saved", "no link": "no payment link"}.get(outcome, "failed"),
               detail=f"Payment {pay}", file=shown if outcome == "saved" else "")
        drain_notes(section)

    print(f"=== {tax_type_label}: finished ===")
    _back_to_level1_after(page, tax_type_label, tin, password)

# ---------- Main ----------

# ---------------- speed settings ----------------
HEADLESS = os.getenv("URA_HEADLESS", "0") == "1"   # set URA_HEADLESS=1 in .env
SPEED_UPS = os.getenv("URA_SPEEDUPS", "0") == "1"  # set URA_SPEEDUPS=1 in .env to enable
DEBUG_BLOCKING = os.getenv("URA_DEBUG", "0") == "1"   # log blocked requests
BLOCK_IMAGES = True        # skip images/fonts/media from OTHER sites only (never ura.go.ug)
BLOCK_CHAT = True          # skip the chat + chatbot widgets, they never finish loading

BLOCKED_HOSTS = ["hclb.ura.go.ug", "chatbot", "googletagmanager", "google-analytics",
                 "doubleclick", "facebook", "hotjar", "clarity.ms", "youtube"]
BLOCKED_TYPES = {"image", "font", "media"}

def speed_up(context):
    """Stop the browser downloading things the bot never looks at. The portal
    itself is slow server-side, but this removes most of the waiting.
    The security-code image is always let through - you have to read it."""
    def blocked(url, why):
        if DEBUG_BLOCKING:
            print(f"  [blocked: {why}] {url[:110]}")

    def handler(route, request):
        url = request.url.lower()
        if BLOCK_CHAT and any(h in url for h in BLOCKED_HOSTS):
            blocked(url, "chat/analytics")
            return route.abort()
        # Never block anything served by the portal itself - the security code
        # and the Refresh icon are images, and blocking them breaks login.
        if "ura.go.ug" in url:
            return route.continue_()
        if BLOCK_IMAGES and request.resource_type in BLOCKED_TYPES:
            blocked(url, request.resource_type)
            return route.abort()
        return route.continue_()
    if not SPEED_UPS:
        print("Faster page loading is off (tick 'Load portal pages faster' on the Ledgers page to turn it on).")
        return
    try:
        context.route("**/*", handler)
        print("Faster page loading is on: the chat window and pictures from other websites "
              "are not loaded (everything from ura.go.ug still loads).")
    except Exception as e:
        print(f"Could not enable speed-ups ({type(e).__name__})")

RUN_SECTIONS_4_AND_10 = False   # turn on once those screens are recorded

TAX_TYPES = ["Value Added Tax", "IT-WHT", "PAYE", "Local Excise Duty"]
ANNUAL_TAX_TYPES = ["Income Tax", "Rental"]


# ============================================================================
# Section 9 of the SOP: tax credit certificates
#   9.1-9.2  Withholding tax credit certificate - one per month
#   9.3-9.4  PAYE tax credit certificates for the period, saved in one zip
# The selectors below are best guesses from the SOP screenshots. Each one is a
# list of candidates and the script prints what it could not find, so fixing
# them means editing these lists only.
# ============================================================================

def open_menu_path(page, labels, timeout=60):
    """Click through a chain of menu labels, e.g. e-Services > Return > ..."""
    frame = None
    for raw_label in labels:
        optional = raw_label.startswith("?")
        label = raw_label.lstrip("?")
        clear_overlays(page)
        deadline = time.time() + (8 if optional else timeout)
        clicked = False
        while time.time() < deadline and not clicked:
            for f in page.frames:
                for loc in (f.get_by_role("link", name=label, exact=True),
                            f.get_by_role("heading", name=label, exact=True),
                            f.get_by_text(label, exact=True)):
                    try:
                        if loc.count() > 0 and loc.first.is_visible():
                            safe_click(page, loc, f"menu '{label}'")
                            frame, clicked = f, True
                            break
                    except Exception:
                        continue
                if clicked:
                    break
            if not clicked:
                nap(page, 1000)
        if not clicked:
            if optional:
                print(f"Menu level '{label}' not on this portal - skipping it.")
                continue
            raise RuntimeError(f"Menu item '{label}' not found - check the menu path lists")
        nap(page, 2000)
    return frame

def _locator_for(frame, candidate):
    if candidate.startswith(("#", ".", "[", "input", "select", "button", "a")):
        return frame.locator(candidate)
    return frame.get_by_label(candidate)

def set_field(page, frame, candidates, value, what):
    """Fill an input or choose a dropdown option, trying each candidate."""
    for candidate in candidates:
        try:
            loc = _locator_for(frame, candidate)
            if loc.count() == 0 or not loc.first.is_visible():
                continue
            tag = loc.first.evaluate("el => el.tagName.toLowerCase()")
            if tag == "select":
                try:
                    loc.first.select_option(label=value)
                except Exception:
                    options = loc.first.evaluate(
                        "el => [...el.options].map(o => o.text.trim())")
                    match = next((o for o in options if value.lower() in o.lower()), None)
                    if match is None:
                        print(f"{what}: no option like '{value}' in {options}")
                        return False
                    loc.first.select_option(label=match)
            else:
                loc.first.fill(value)
            return True
        except Exception:
            continue
    print(f"Could not set {what} (tried {candidates}) - fix the candidate list")
    return False

def click_button(page, frame, names, what):
    for name in names:
        for loc in (frame.get_by_role("button", name=name),
                    frame.locator(f"input[value='{name}']"),
                    frame.get_by_text(name, exact=True)):
            try:
                if loc.count() > 0 and loc.first.is_visible():
                    safe_click(page, loc, f"{what} ('{name}')")
                    return name
            except Exception:
                continue
    print(f"Could not find a {what} button (tried {names})")
    return None

def frame_shows_no_records(frame):
    for text in NO_RECORD_TEXTS:
        try:
            if frame.get_by_text(text, exact=False).count() > 0:
                return True
        except Exception:
            continue
    return False

def click_and_save(page, locator, folder: Path, filename: str):
    """Save a download; if the portal opens a PDF in a new tab instead,
    fetch that tab's URL through the same logged-in session."""
    before = len(page.context.pages)
    try:
        with page.expect_download(timeout=45000) as info:
            locator.first.click()
        save_download(info.value, folder, filename)
        return True
    except Exception:
        nap(page, 3000)
        if len(page.context.pages) > before:
            tab = page.context.pages[-1]
            try:
                tab.wait_for_load_state("domcontentloaded")
                body = page.context.request.get(tab.url).body()
                folder.mkdir(parents=True, exist_ok=True)
                (folder / filename).write_bytes(body)
                print(f"Saved: {folder / filename} (from a new tab)")
                tab.close()
                return True
            except Exception as e:
                print(f"Could not save the new tab ({type(e).__name__})")
                try:
                    tab.close()
                except Exception:
                    pass
        return False

# ---------- Section 9: tax credit certificates (from the recording) ----------
# Selectors below come from Tax_credit_certificates_(withholding_tax_and_PAYE).json
WHT_FORM = "#GenTaxCreditCert"
WHT_MENU_LINK = "Generate Tax Credit Certificate for Withholding"
WHT_GENERATE_FOR = "#generateFor"          # D = Payee ( Deductee )
WHT_TO_DATE = "#toDt"
WHT_SEARCH = [f"{WHT_FORM} tr.tdbar input[value='Search']", f"{WHT_FORM} input[value='Search']"]

PAYE_FORM = "#TaxCreditCertForm"
PAYE_MENU_LINK = "Generate Tax Credit Certificate for Pay As You Earn"
PAYE_BRANCH = "#retBrnch"                  # Headquarter
PAYE_PERIOD = "#period"                    # e.g. 2025-2026
PAYE_SEARCH = [f"{PAYE_FORM} button:has-text('Search')", f"{PAYE_FORM} input[value='Search']"]

GENERATE_TCC = ["Generate TCC", "GenerateTCC", "Generate Tax Credit Certificate", "Generate"]
NO_RECORD_TEXTS = ["No record", "No Data Found", "No records found"]

PORTAL_ENTRY = [
    "#main-header li:nth-of-type(5) li:nth-of-type(1) > a",   # Login alert > Portal
    "#logoutFormISSP a",                                      # portal landing page links
]
LOGIN_ALERT = "#middle40Drop > span"
SKIP_BUTTON = "#parishMsgA button"

def portal_menu_frame(page):
    """The frame holding the e-Services / Reports menu, if it is loaded.

    #mailbox is present but NOT 'visible' until a menu is opened, so presence
    (or the home-page URL) is what counts here, not visibility.
    """
    for f in page.frames:
        try:
            if "loadhomepagecontent" in (f.url or "").lower():
                return f
        except Exception:
            continue
    for f in page.frames:
        try:
            if f.locator("#mailbox").count() > 0:
                return f
            if f.get_by_text("e-Services", exact=True).count() > 0:
                return f
        except Exception:
            continue
    return None

def click_first_visible(page, selectors_or_texts, what):
    """Click the first of these that is actually on screen, in any frame.  Every match is
    looked at, not only the first: a hidden menu entry with the same text can come earlier in
    the page (live 5 Oct 2026: 'Getting the menu back (Portal link)' clicked nothing, 5 times)."""
    for item in selectors_or_texts:
        for f in page.frames:
            try:
                loc = f.locator(item) if item.startswith(("#", ".", "[", "div", "a", "input", "button")) \
                    else f.get_by_text(item, exact=True)
                n = loc.count()
            except Exception:
                continue
            for i in range(min(n, 20)):
                try:
                    one = loc.nth(i)
                    if not one.is_visible():
                        continue
                    print(f"{what}: clicking {item!r}")
                    safe_click(page, one, what)
                    return True
                except Exception:
                    continue
    return False

PORTAL_HOME_URL = "https://ura.go.ug/Portal/loginController/loadHomePageContent"

def portal_session_frame(page):
    """A live portal frame (not the login page). Its cookies are the session."""
    for f in page.frames:
        url = (f.url or "").lower()
        if "/portal/" in url and "login" not in url.replace("logincontroller/loadhomepage", ""):
            return f
        if "/portal/" in url and "isspportallogin" not in url:
            return f
    return None

def go_home_inside_portal(page):
    """Send the portal frame itself back to its home page. This keeps the
    session alive - going out through the public site's Portal link is what
    can drop you on the login screen."""
    frame = portal_session_frame(page)
    if frame is None:
        return False
    try:
        frame.evaluate("url => { window.location.href = url; }", PORTAL_HOME_URL)
        print("Returning to the portal home page (staying logged in).")
        wait_find(page, ["#mailbox", "#mailbox a"], timeout=30000, visible=False)
        return portal_menu_frame(page) is not None
    except Exception as e:
        print(f"Could not navigate inside the portal ({type(e).__name__})")
        return False

def ensure_portal_home(page, tin=None, password=None):
    """Always start a screen from the real portal home page. After a failure the
    portal frame can be on some other screen (e.g. Change Password) that still
    looks like 'in the portal'."""
    frame = portal_session_frame(page)
    if frame is not None and "loadhomepagecontent" not in (frame.url or "").lower():
        go_home_inside_portal(page)
    return go_to_portal_home(page, tin, password)

def on_login_page(page) -> bool:
    """The portal login page, even before its fields have rendered."""
    for f in page.frames:
        url = (f.url or "").lower()
        if "issportallogin" in url or "logincontroller/login" in url or "etax-login" in url:
            if "loadhomepagecontent" not in url:
                return True
    return False

def go_to_portal_home(page, tin=None, password=None):
    """Back to the portal menus WITHOUT logging out.

    Home only reaches the public site, so from there we go through
    'Login alert' > 'Portal' (as in the recording) or the links on the
    portal landing page, then dismiss the Skip popup.
    """
    for attempt in range(1, 5):
        clear_overlays(page)
        frame = portal_menu_frame(page)
        if frame is not None:
            return frame

        # the Skip popup hides the menus behind it
        for f in page.frames:
            try:
                btn = f.locator(SKIP_BUTTON)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click()
                    nap(page, 2000)
            except Exception:
                continue
        if portal_menu_frame(page) is not None:
            return portal_menu_frame(page)

        # The Portal link lands on the login page when the session has ended.
        # Wait a moment for the form to draw before deciding - checking too
        # early made the bot click Portal over and over instead.
        if login_form_present(page) or (on_login_page(page) and portal_menu_frame(page) is None):
            try:
                find_login_frame(page, timeout=20)
            except Exception:
                pass
            if login_form_present(page):
                print("Session has ended - logging in again.")
                login(page, tin, password)
                wait_find(page, ["#mailbox", "#mailbox a"], timeout=30000, visible=False)
                continue

        print(f"Getting back into the portal (attempt {attempt}/4)")

        # First choice: stay inside the portal and go to its home page.
        if go_home_inside_portal(page):
            frame = portal_menu_frame(page)
            if frame is not None:
                return frame

        # Fallback: out through the public site's Login alert > Portal link.
        click_first_visible(page, [LOGIN_ALERT], "Login alert")     # opens the menu
        nap(page, 1000)
        if not click_first_visible(page, PORTAL_ENTRY + ["Portal"], "Portal link"):
            print("Portal link not on screen yet - waiting and trying again")
        nap(page, 6000)
        clear_overlays(page)

    print("Frames right now:")
    for i, f in enumerate(page.frames):
        print(f"  {i}: {f.url}")
    shot = safe_screenshot(page, "error_portal_home.png")
    raise RuntimeError("Could not get back to the portal menus "
                       f"- see {shot} and the frame list above")

# ---------- Section 9: tax credit certificates (from the recording) ----------
# Selectors below come from Tax_credit_certificates_(withholding_tax_and_PAYE).json

def open_cert_form(page, menu_link, form_selector, tin=None, password=None):
    """e-Services > Return > Other Return Services > <certificate screen>."""
    ensure_portal_home(page, tin, password)
    frame = find_frame_containing_text(page, "e-Services", timeout=60)
    safe_click(page, frame.locator("#mailbox header h3"), "e-Services menu")
    nap(page, 1200)

    for label in ("Return", "Other Return Services"):
        for loc in (frame.get_by_role("link", name=label, exact=True),
                    frame.get_by_text(label, exact=True)):
            try:
                if loc.count() > 0 and loc.first.is_visible():
                    safe_click(page, loc, f"menu '{label}'")
                    nap(page, 1000)
                    break
            except Exception:
                continue

    safe_click(page, frame.get_by_text(menu_link, exact=False), f"menu '{menu_link}'")
    nap(page, 4000)
    clear_overlays(page)

    form_frame = None
    deadline = time.time() + 60
    while time.time() < deadline and form_frame is None:
        for f in page.frames:
            try:
                if f.locator(form_selector).count() > 0:
                    form_frame = f
                    break
            except Exception:
                continue
        if form_frame is None:
            nap(page, 1000)
    if form_frame is None:
        raise RuntimeError(f"The certificate form {form_selector} did not open")
    return form_frame

CLEAN_DATES_JS = r"""
(form) => {
  const bad = [];
  form.querySelectorAll("input[type='text'], input:not([type])").forEach(el => {
    const v = (el.value || '').trim();
    if (v.startsWith('<') || v.toLowerCase() === 'undefined' || v.toLowerCase() === 'nan') {
      bad.push(el.id || el.name || '(no id)');
      el.value = '';
    }
  });
  return bad;
}"""

SET_DATE_JS = r"""
(form, args) => {
  const norm = s => (s || '').replace(/\s+/g, ' ').trim().toLowerCase();
  const rows = [...form.querySelectorAll('tr')];
  for (const row of rows) {
    const cells = [...row.children];
    for (let i = 0; i < cells.length; i++) {
      if (norm(cells[i].innerText).startsWith(norm(args.label))) {
        for (let j = i; j < cells.length; j++) {
          const input = cells[j].querySelector("input[type='text'], input:not([type])");
          if (input) {
            input.value = args.value;
            input.dispatchEvent(new Event('input', {bubbles: true}));
            input.dispatchEvent(new Event('change', {bubbles: true}));
            if (typeof input.onchange === 'function') input.onchange();
            return input.id || input.name || 'ok';
          }
        }
      }
    }
  }
  return null;
}"""

def clean_date_fields(page, frame, form_selector):
    """The portal sometimes leaves page HTML or 'undefined' in its date boxes,
    which breaks its own calendar. Empty those before we touch them."""
    try:
        bad = frame.locator(form_selector).first.evaluate(CLEAN_DATES_JS)
        if bad:
            print(f"Cleared junk out of: {', '.join(bad)}")
    except Exception as e:
        print(f"Could not tidy the date fields ({type(e).__name__})")

def set_form_date(frame, form_selector, label, value):
    """Type a date into the box next to a label, e.g. 'Return Date From'."""
    try:
        where = frame.locator(form_selector).first.evaluate(
            SET_DATE_JS, {"label": label, "value": value})
    except Exception as e:
        print(f"Could not set '{label}' ({type(e).__name__})")
        return False
    if where:
        print(f"Set {label} = {value}  (field {where})")
        return True
    print(f"Could not find the field for '{label}'")
    return False

def period_dates(period_from, period_to):
    """'2025-01','2026-06' -> ('01/01/2025', '30/06/2026')."""
    start = datetime.datetime.strptime(period_from, "%Y-%m")
    end = datetime.datetime.strptime(period_to, "%Y-%m")
    nxt = end.replace(year=end.year + 1, month=1) if end.month == 12 else end.replace(month=end.month + 1)
    last = nxt - datetime.timedelta(days=1)
    return start.strftime("%d/%m/%Y"), last.strftime("%d/%m/%Y")

def ask_dates(what, form_hint, defaults=("", "")):
    """Either the user types the dates here, or picks them in the browser."""
    if PROMPT_HANDLER is not None:
        return ask_user(f"{what}  ({form_hint})", kind="dates", defaults=defaults)
    print(f"\n>>> {what}")
    print("    Type the dates here as dd/mm/yyyy, e.g. 01/07/2025 30/06/2026")
    print(f"    - or set them in the browser yourself ({form_hint}) and just press Enter")
    print("    - 's' to skip this one, 'q' to stop this section.")
    answer = input("    > ").strip()
    low = answer.lower()
    if low in ("s", "q", ""):
        return low, None, None
    parts = answer.replace(",", " ").split()
    if len(parts) == 2:
        return "dates", parts[0], parts[1]
    print("    Did not understand that - carrying on with whatever is in the browser.")
    return "", None, None

def hand_over(what):
    """Pause so the user can set something the bot should not guess."""
    if PROMPT_HANDLER is not None:
        return (ask_user(f"In the browser: {what}", kind="handover") or "").strip().lower()
    print(f"\n>>> IN THE BROWSER: {what}")
    print("    Then press Enter here to let the bot carry on"
          " (type 's' to skip this one, 'q' to stop this section).")
    return input("    > ").strip().lower()

CONTROLS_JS = r"""
() => {
  const norm = s => (s || '').replace(/\s+/g, ' ').trim();
  const out = {buttons: [], checkboxes: 0, checkboxIds: []};
  document.querySelectorAll("button, input[type='button'], input[type='submit'], a")
    .forEach(el => {
      const label = norm(el.value || el.innerText || el.getAttribute('title') || '');
      if (label) out.buttons.push(`${el.tagName.toLowerCase()}` +
        (el.id ? '#' + el.id : '') + ` "${label}"` +
        (el.offsetParent === null ? ' (hidden)' : ''));
    });
  document.querySelectorAll("input[type='checkbox']").forEach(el => {
    out.checkboxes++;
    if (el.id || el.name) out.checkboxIds.push(el.id || el.name);
  });
  return out;
}"""

def dump_controls(page, tag="after_search"):
    """Print the buttons / checkboxes the portal is actually showing, so the
    right names can be put into GENERATE_TCC."""
    print(f"--- controls on screen ({tag}) ---")
    for i, f in enumerate(page.frames):
        try:
            info = f.evaluate(CONTROLS_JS)
        except Exception:
            continue
        if not info["buttons"] and not info["checkboxes"]:
            continue
        print(f"  frame {i}: {f.url[:90]}")
        print(f"    checkboxes: {info['checkboxes']}  {info['checkboxIds'][:8]}")
        for b in info["buttons"][:25]:
            print(f"    {b}")
    try:
        html = "\n\n".join(f"<!-- frame {i}: {f.url} -->\n" + f.content()
                            for i, f in enumerate(page.frames))
        target = work_file(f"debug_{tag}.html")
        target.write_text(html, encoding="utf-8", errors="ignore")
        print(f"  (full HTML written to {target})")
    except Exception:
        pass

SELECT_ALL = "#selectAll"
GENERATE_TCC_BTN = ["input[value='GenerateTCC']", "input[value*='generate' i]",
                    "#GenerateTCC", "button:has-text('GenerateTCC')"]
RESULT_MARKERS = [SELECT_ALL, "input[value='GenerateTCC']",
                  "input[type='checkbox'][name='listPaye']", "input[value*='generate' i]"]

def portal_frames(page):
    """Only the frames belonging to the portal application. Frame 0 is the
    public web site - its newsletter checkbox is NOT a result row."""
    out = []
    for f in page.frames:
        url = (f.url or "").lower()
        if "/portal/" in url and "chatbot" not in url:
            out.append(f)
    return out or list(page.frames)

def wait_for_results(page, timeout=90):
    """After Search, wait until the result list is really on screen. The
    portal takes several seconds, and looking too early is why the select-all
    box and the GenerateTCC button appeared to be missing."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        for f in portal_frames(page):
            for sel in RESULT_MARKERS:
                try:
                    if f.locator(sel).count() > 0:
                        print("Results are on screen.")
                        return f
                except Exception:
                    continue
            try:                       # a "no records" answer is also a result
                if f.locator("text=No record").count() > 0:
                    return f
            except Exception:
                pass
        page.wait_for_timeout(POLL_MS)
    print("The result list never appeared.")
    return None

CERT_ROWS_JS = r"""
() => {
  const norm = s => (s || '').replace(/\s+/g, ' ').trim();
  let best = null, bestN = 0;
  for (const t of document.querySelectorAll('table')) {
    const n = [...t.querySelectorAll('tr')].filter(r =>
      r.querySelector("input[type='checkbox']") && !r.querySelector('th')).length;
    if (n > bestN) { best = t; bestN = n; }
  }
  if (!best) return null;
  const rows = [...best.querySelectorAll('tr')];
  const head = rows.find(r => r.querySelectorAll('th').length > 1) || null;
  const heads = head ? [...head.children].map(c => norm(c.innerText).toLowerCase()) : [];
  const tccCol = heads.findIndex(h => /\btcc\b|certificate\s*(no|num)|cert\.?\s*no/.test(h));
  const out = [];
  rows.forEach(r => {
    if (r === head) return;
    const box = r.querySelector("input[type='checkbox']");
    if (!box || /selectall/i.test(box.id || '')) return;
    box.setAttribute('data-bot-cert', String(out.length));
    const cells = [...r.children];
    const tcc = tccCol >= 0 && cells[tccCol] ? norm(cells[tccCol].innerText) : '';
    out.push({n: out.length, tcc, disabled: box.disabled, checked: box.checked,
              text: norm(r.innerText).slice(0, 150)});
  });
  return {heads, tccCol, rows: out};
}"""

EMPTY_TCC = {"", "-", "--", "n/a", "na", "null", "none", "0"}

def read_cert_rows(page):
    """The certificate result rows (in whichever portal frame has the most),
    each with its TCC number if the portal already generated one."""
    best_frame, best = None, None
    for f in portal_frames(page):
        try:
            info = f.evaluate(CERT_ROWS_JS)
        except Exception:
            continue
        if info and (best is None or len(info["rows"]) > len(best["rows"])):
            best_frame, best = f, info
    return best_frame, best

def tick_cert_rows(page, frame, rows):
    """Tick exactly these rows; untick every other row."""
    wanted = {r["n"] for r in rows}
    for n_box in range(frame.locator("[data-bot-cert]").count()):
        box = frame.locator(f"[data-bot-cert='{n_box}']")
        try:
            if box.count() == 0 or box.first.is_disabled():
                continue
            should = n_box in wanted
            # one script click only when the state differs, then read back -
            # a mouse click that 'timed out' but worked, followed by a second
            # click, used to untick every row (live portal, 3 Oct 2026)
            box.first.evaluate("(el, want) => { if (el.checked !== want) el.click(); "
                               "return el.checked; }", should, timeout=5000)
        except Exception:
            continue
    ticked = 0
    for n in wanted:
        try:
            if frame.locator(f"[data-bot-cert='{n}']").first.is_checked():
                ticked += 1
        except Exception:
            pass
    return ticked

def find_generate_button(page, frame):
    frames = [frame] + [f for f in portal_frames(page) if f is not frame]
    for f in frames:
        for sel in ("input[value='GenerateTCC']", "input[value*='generate' i]",
                    "button:has-text('Generate')"):
            try:
                loc = f.locator(sel)
                for i in range(loc.count()):
                    if loc.nth(i).is_visible():
                        return loc.nth(i)
            except Exception:
                continue
    return None

def click_and_capture_pdfs(page, click, folder: Path, stem: str, first_wait=90,
                           manual_message=None):
    """Catch every PDF the portal sends after `click()` - as a download, as a
    PDF opened in a new tab, or as a PDF loaded in the page - and accept any
    'Are you sure?' box on the way. With manual_message, nothing is clicked:
    the person clicks in the browser and the same catching applies."""
    ctx = page.context
    pdf_responses, downloads, new_pages, dialogs = [], [], [], []

    def on_response(resp):
        try:
            if "pdf" in (resp.headers.get("content-type") or "").lower():
                pdf_responses.append(resp)
        except Exception:
            pass
    def on_download(dl):
        downloads.append(dl)
    def on_dialog(dialog):
        dialogs.append(dialog.message)
        try:
            dialog.accept()
        except Exception:
            pass
    def on_page(p):
        new_pages.append(p)
        p.on("download", on_download)
        p.on("dialog", on_dialog)

    ctx.on("response", on_response)
    ctx.on("page", on_page)
    page.on("download", on_download)
    page.on("dialog", on_dialog)
    try:
        if manual_message:
            answer = ask_user(manual_message, kind="handover",
                              poll=lambda: (page.wait_for_timeout(50) or bool(pdf_responses or downloads)))
            if answer in ("s", "q") and not (pdf_responses or downloads):
                return []
        else:
            click()
        deadline = time.time() + first_wait
        last_n, quiet_since = 0, time.time()
        while time.time() < deadline:
            nap(page, 300)
            n = len(pdf_responses) + len(downloads)
            if n != last_n:
                last_n, quiet_since = n, time.time()
            elif n and time.time() - quiet_since >= PDF_QUIET_SECONDS:
                break
            elif not n and new_pages and time.time() - quiet_since > 15:
                break              # a tab opened but no PDF came through it
        for msg in dialogs:
            print(f"Portal message accepted: {msg}")

        folder.mkdir(parents=True, exist_ok=True)
        saved, seen = [], set()
        def target(i, total, suggested=None):
            return folder / (f"{stem}.pdf" if total == 1 else f"{stem}_{i:02d}.pdf")
        total = len(downloads) + len(pdf_responses)
        i = 0
        for dl in downloads:
            i += 1
            dest = target(i, total, dl.suggested_filename)
            try:
                dl.save_as(str(dest)); saved.append(dest); seen.add(dl.url)
                print(f"Saved: {dest}")
            except Exception as e:
                print(f"A download could not be saved ({type(e).__name__})")
        for resp in pdf_responses:
            if resp.url in seen:
                continue
            i += 1
            try:
                body = resp.body()
            except Exception:
                body = b""
            if body[:4] != b"%PDF":
                try:                              # same logged-in session
                    body = ctx.request.get(resp.url).body()
                except Exception:
                    body = b""
            if body[:4] == b"%PDF":
                dest = target(i, total)
                dest.write_bytes(body); saved.append(dest); seen.add(resp.url)
                print(f"Saved: {dest}")
        if not saved:                             # tab opened, response not caught
            for p in new_pages:
                try:
                    body = ctx.request.get(p.url).body() if p.url.startswith("http") else b""
                except Exception:
                    body = b""
                if body[:4] == b"%PDF":
                    i += 1
                    dest = target(i, 1 + len(new_pages))
                    dest.write_bytes(body); saved.append(dest)
                    print(f"Saved: {dest} (from the new tab)")
        return saved
    finally:
        for p in new_pages:
            try:
                if not p.is_closed():
                    p.close()
            except Exception:
                pass
        for obj, ev, fn in ((ctx, "response", on_response), (ctx, "page", on_page),
                            (page, "download", on_download), (page, "dialog", on_dialog)):
            try:
                obj.remove_listener(ev, fn)
            except Exception:
                pass

def save_results_html(page, frame, folder: Path, tag):
    """Keep the results screen as it was BEFORE Generate, for diagnosis."""
    try:
        folder.mkdir(parents=True, exist_ok=True)
        (folder / f"_screen_{tag}.html").write_text(frame.content(), encoding="utf-8",
                                                    errors="ignore")
    except Exception:
        pass

def generate_certificates(page, save_folder: Path, section, folder: Path, stem, tag):
    """Read the result rows. Rows that already carry a TCC number were generated
    before and cannot be generated again - they are listed, not ticked. The rest
    are ticked and Generate is pressed once; the PDF is caught however it arrives."""
    frame, info = read_cert_rows(page)
    deadline = time.time() + 30
    while (info is None or not info["rows"]) and time.time() < deadline:
        if frame is not None and frame_shows_no_records(frame):
            break
        nap(page, 500)
        frame, info = read_cert_rows(page)
    if info is None or not info["rows"]:
        print(f"{section}: no result rows found.")
        record(section, tag, "no result rows")
        return "no rows"
    save_results_html(page, frame, folder.parent, tag)

    rows = info["rows"]
    if info["tccCol"] < 0:
        print(f"{section}: no TCC number column found (headings: {info['heads']}) - "
              f"cannot tell which rows were generated before; ticking every row.")
        todo, done = [r for r in rows if not r["disabled"]], []
        record(section, tag, "note", detail="TCC number column not found; all rows ticked")
    else:
        done = [r for r in rows if r["tcc"].strip().lower() not in EMPTY_TCC]
        todo = [r for r in rows if r not in done and not r["disabled"]]
    for r in done:
        record(section, r["text"][:80], "already generated", detail=f"TCC {r['tcc']}")
    print(f"{section}: {len(rows)} row(s) - {len(done)} already generated (TCC number shown), "
          f"{len(todo)} to generate.")
    if not todo:
        return "all generated before"

    stem = f"{stem}_{datetime.datetime.now():%Y%m%d_%H%M%S}"
    ticked = tick_cert_rows(page, frame, todo)
    print(f"Ticked {ticked} of {len(todo)} row(s).")
    button = find_generate_button(page, frame)
    if button is None:
        print(f"{section}: Generate button not found.")
        dump_controls(page, f"{tag}_no_generate")
        saved = []
    else:
        try:
            label = button.get_attribute("value") or button.inner_text()
        except Exception:
            label = "Generate"
        print(f"Clicking '{label}' (the portal shows a summary, not a PDF)...")
        def press():
            try:
                button.click(timeout=5000)
            except Exception:
                js_click(button)
        saved = click_and_capture_pdfs(page, press, folder, stem, first_wait=20)
    # Live portal (recording 3 Oct 2026): Generate TCC shows a summary screen and
    # NO PDF; the TCC number appears when the period is searched again.
    for r in todo:
        record(section, r["text"][:80], "generated" if saved else "generate pressed",
               detail="" if saved else "no PDF from the portal - search the period again to see the TCC No",
               file=", ".join(p.name for p in saved))
    return "saved" if saved else "pressed"

def download_all_links(page, frame, folder: Path, stem: str):
    """Save explicit download / PDF links only. It never clicks ordinary
    links - that is what once took the portal to the Change Password page."""
    saved = 0
    for sel in ("table a[href*='download' i]", "table a:has-text('Download')", "a[href$='.pdf']"):
        try:
            links = frame.locator(sel)
            n = links.count()
        except Exception:
            continue
        for i in range(n):
            if click_and_save(page, links.nth(i), folder, f"{Path(stem).stem}_{i + 1:03d}.pdf"):
                saved += 1
        if saved:
            return True
    return False

def generate_and_collect(page, frame, folder: Path, stem: str, wait_seconds=120):
    """Used by the (switched-off) return-copies section."""
    button = find_generate_button(page, frame)
    if button is None:
        return []
    return click_and_capture_pdfs(page, lambda: js_click(button), folder, stem,
                                  first_wait=wait_seconds)

def zip_folder(folder: Path):
    """Put everything in the folder into one zip next to it (SOP 9.4 step 5)."""
    if not folder.exists() or not any(folder.iterdir()):
        print(f"Nothing to zip in {folder}")
        return None
    zip_path = Path(str(folder) + ".zip")
    if zip_path.exists():
        zip_path.unlink()
    shutil.make_archive(str(folder), "zip", root_dir=folder)
    print(f"Zipped {len(list(folder.iterdir()))} file(s) into {zip_path}")
    return zip_path

def run_search(page, frame, extra_selectors):
    if not click_button(page, frame, ["Search"], "search"):
        for sel in extra_selectors:
            try:
                if frame.locator(sel).count() > 0:
                    safe_click(page, frame.locator(sel), "search")
                    break
            except Exception:
                continue
    results = wait_for_results(page)
    clear_overlays(page)
    return results or frame

def process_wht_certificates(page, period_from, period_to, save_folder: Path, tin, password):
    """SOP 9.1-9.2: Payee (Deductee), dates, Search, generate the rows that
    have no TCC number yet, save the PDF."""
    section = "WHT tax credit certificates"
    period_tag = f"{period_from}_to_{period_to}"
    folder = save_folder / "tax_credit_certificates" / "WHT"

    frame = open_cert_form(page, WHT_MENU_LINK, WHT_FORM, tin, password)
    try:
        frame.locator(WHT_GENERATE_FOR).select_option("D")        # Payee ( Deductee )
        print("Set 'Generate Certificate as' = Payee (Deductee).")
    except Exception as e:
        print(f"Could not set {WHT_GENERATE_FOR} ({type(e).__name__}) - set it yourself in the browser")

    clean_date_fields(page, frame, WHT_FORM)
    mode, date_from, date_to = ask_dates(
        "Withholding certificates: which period?",
        "Return Date From / Return Date To",
        defaults=period_dates(period_from, period_to))
    if mode in ("s", "q"):
        record(section, period_tag, "skipped by you")
        return
    if mode == "dates":
        set_form_date(frame, WHT_FORM, "Return Date From", date_from)
        set_form_date(frame, WHT_FORM, "Return Date To", date_to)

    frame = run_search(page, frame, WHT_SEARCH)
    if frame_shows_no_records(frame):
        print("WHT certificates: no records for those dates.")
        record(section, period_tag, "no records")
        return
    outcome = generate_certificates(page, save_folder, section, folder,
                                    f"WHT_TCC_{period_tag}", f"wht_{period_tag}")
    mark_progress(save_folder, "WHT Certificates", period_tag, outcome)

def process_paye_certificates(page, period_from, period_to, save_folder: Path, tin, password):
    """SOP 9.3-9.4: Headquarter, you pick the Period, Search, generate the rows
    that have no TCC number yet, save the PDF."""
    section = "PAYE tax credit certificates"
    period_tag = f"{period_from}_to_{period_to}"
    folder = save_folder / "tax_credit_certificates" / "PAYE"

    frame = open_cert_form(page, PAYE_MENU_LINK, PAYE_FORM, tin, password)
    try:
        frame.locator(PAYE_BRANCH).select_option(label="Headquarter")
        print("Set 'Return For' = Headquarter.")
    except Exception as e:
        print(f"Could not set {PAYE_BRANCH} ({type(e).__name__}) - set it yourself if needed")

    answer = hand_over("choose the Period for the PAYE certificates (the drop-down next to "
                       "'Period'), then press 'Done - continue'.")
    if answer in ("s", "q"):
        record(section, period_tag, "skipped by you")
        return

    frame = run_search(page, frame, PAYE_SEARCH)
    if frame_shows_no_records(frame):
        print("PAYE certificates: no records for that period.")
        record(section, period_tag, "no records")
        return
    outcome = generate_certificates(page, save_folder, section, folder,
                                    f"PAYE_TCC_{period_tag}", f"paye_{period_tag}")
    mark_progress(save_folder, "PAYE Certificates", period_tag, outcome)

# ---------- Section 10: PRN payment slips (SOP 10) ----------
# The ledger side uses the same screens as sections 7/8, so those selectors are
# known. The Re-Print screen is a guess until you record it - the lists below
# are the only things that should need changing.
PAYMENT_MENU = ["Payment", "Re-Print a Payment Slip"]
PRN_FIELD = ["PRN", "Payment Registration Number", "#prn", "#prnNo",
             "input[name*='prn' i]"]
PRN_SEARCH = ["Search", "Submit", "View"]
PRN_DOWNLOAD = ["Export To PDF", "Print", "Download", "Export To Excel"]

PRN_COLUMN_JS = r"""
(table) => {
  const norm = s => (s || '').replace(/\s+/g, ' ').trim();
  const rows = [...table.querySelectorAll('tr')];
  const head = rows.find(r => r.querySelectorAll('th').length > 2);
  if (!head) return [];
  const heads = [...head.children].map(c => norm(c.innerText).toLowerCase());
  let col = heads.findIndex(h => h.includes('payment registration'));
  if (col < 0) col = heads.findIndex(h => h.includes('prn'));
  if (col < 0) return [];
  const out = [];
  rows.forEach(r => {
    const cells = [...r.children];
    if (cells[col] && !cells[col].querySelector('th')) {
      const v = norm(cells[col].innerText);
      if (v && /^[A-Z0-9-]{6,}$/i.test(v) && !out.includes(v)) out.push(v);
    }
  });
  return out;
}"""

def prns_on_level3(frame):
    """Read the Payment Registration Numbers off a Level 3-2 detail screen."""
    tables = frame.locator("table")
    found = []
    for i in range(tables.count()):
        try:
            if not tables.nth(i).is_visible():
                continue
            for prn in tables.nth(i).evaluate(PRN_COLUMN_JS) or []:
                if prn not in found:
                    found.append(prn)
        except Exception:
            continue
    return found

def collect_prns(page, tax_type, month_label, save_folder, tin, password):
    """Open the month's Level 3 detail and return its PRNs."""
    frame = open_level2(page, tax_type, tin, password)
    marked = mark_payment_link(frame, month_label)
    if not marked.get("link"):
        print(f"{tax_type} {month_label}: no payment link - no PRNs.")
        return [], frame
    safe_click(page, frame.locator("[data-bot-target='1']"), f"payment for {month_label}")
    l3 = wait_for_level(page, "Level 3", timeout_seconds=25)
    if l3 is None:
        print(f"{tax_type} {month_label}: Level 3 did not open.")
        return [], frame
    prns = prns_on_level3(l3)
    print(f"{tax_type} {month_label}: PRNs {prns or 'none found'}")
    click_back(page, l3, container="#levelThreeTwoOfAR")
    return prns, wait_for_level(page, "Level 2", timeout_seconds=25) or frame

def open_payment_slip_form(page, tin, password):
    """e-Services > Payment > Re-Print a Payment Slip."""
    go_to_portal_home(page, tin, password)
    frame = find_frame_containing_text(page, "e-Services", timeout=60)
    safe_click(page, frame.locator("#mailbox header h3"), "e-Services menu")
    nap(page, 1200)
    for label in PAYMENT_MENU:
        clicked = False
        for loc in (frame.get_by_role("link", name=label, exact=False),
                    frame.get_by_text(label, exact=False)):
            try:
                if loc.count() > 0 and loc.first.is_visible():
                    safe_click(page, loc, f"menu '{label}'")
                    nap(page, 2500)
                    clicked = True
                    break
            except Exception:
                continue
        if not clicked:
            print(f"Menu item '{label}' not found - fix PAYMENT_MENU")
    clear_overlays(page)
    return frame

def download_payment_slip(page, prn, folder: Path, tax_type, month_label, tin, password):
    """Re-print one PRN's payment slip (SOP 10.2 steps 5-7)."""
    frame = open_payment_slip_form(page, tin, password)
    if not set_field(page, frame, PRN_FIELD, prn, "PRN"):
        dump_controls(page, "payment_slip_form")
        return False
    click_button(page, frame, PRN_SEARCH, "search")
    nap(page, 4000)
    clear_overlays(page)

    name = f"{tax_type}_PRN_{month_label}_{prn}.pdf".replace(" ", "_")
    for label in PRN_DOWNLOAD:
        for f in page.frames:
            for loc in (f.get_by_role("button", name=label, exact=False),
                        f.locator(f"input[value*='{label}' i]"),
                        f.get_by_text(label, exact=False)):
                try:
                    if loc.count() == 0 or not loc.first.is_visible():
                        continue
                except Exception:
                    continue
                if click_and_save(page, loc, folder, name):
                    return True
    print(f"PRN {prn}: nothing downloaded - the controls on screen:")
    dump_controls(page, "payment_slip_result")
    return False

def process_prn_slips(page, period_from, period_to, save_folder, tin, password,
                      tax_types=None):
    """SOP 10.2: PRN slips for the last month of the period, per tax type."""
    folder = save_folder / "prn_payment_slips"
    months = generate_month_labels(period_from, period_to)
    last_month = months[-1]                      # list runs oldest -> newest
    tax_types = tax_types or TAX_TYPES

    for tax_type in tax_types:
        key = f"{tax_type}|{last_month}"
        if already_done(save_folder, "PRN slips", key):
            print(f"PRN slips {tax_type} {last_month}: already done - skipping.")
            continue
        print(f"\n--- PRN slips: {tax_type} {last_month} ---")
        try:
            prns, _ = collect_prns(page, tax_type, last_month, save_folder, tin, password)
        except Exception as e:
            print(f"{tax_type}: could not read the PRNs ({e})")
            continue
        if not prns:
            mark_progress(save_folder, "PRN slips", key, "no PRNs")
            continue

        saved = 0
        for prn in prns:
            target = folder / f"{tax_type}_PRN_{last_month}_{prn}.pdf".replace(" ", "_")
            if target.exists():
                print(f"PRN {prn}: already downloaded - skipping.")
                saved += 1
                continue
            try:
                if download_payment_slip(page, prn, folder, tax_type, last_month, tin, password):
                    saved += 1
            except Exception as e:
                print(f"PRN {prn} failed: {e}")
        mark_progress(save_folder, "PRN slips", key,
                      f"{saved}/{len(prns)} saved" if saved else "download failed")

# ---------- Sections 4, 5, 6 and 10 ----------
# No recording exists for these screens yet, so the menu labels below are from
# the SOP. If one is not found the bot prints every menu/button on screen
# (dump_controls) and writes a debug_*.html so the label can be corrected here.

RETURN_HISTORY_MENU = ["Return", "Other Return Services", "View Return History"]
PAYMENT_SLIP_MENU = ["Payment", "Re-Print a Payment Slip"]

RETURN_TYPES = [
    "Income Tax Final", "Income Tax Provisional", "Rental Income Tax Final",
    "Pay As You Earn (PAYE)", "Withholding Tax", "Monthly Value-Added Tax",
    "Local Excise Duty",
]

def open_service(page, labels, top="e-Services", tin=None, password=None, wait_for=None):
    """Portal home -> top menu -> the labels, e.g.
    open_service(page, ["Return", "Other Return Services", "View Return History"])"""
    go_to_portal_home(page, tin, password)
    frame = find_frame_containing_text(page, top, timeout=60)
    if frame is None:
        raise RuntimeError(f"Could not find the '{top}' menu")
    safe_click(page, frame.locator("#mailbox header h3") if top == "e-Services"
               else frame.get_by_text(top, exact=True), f"{top} menu")
    nap(page, 1200)

    for label in labels:
        clicked = False
        for loc in (frame.get_by_role("link", name=label, exact=True),
                    frame.get_by_text(label, exact=True),
                    frame.get_by_text(label, exact=False)):
            try:
                if loc.count() > 0 and loc.first.is_visible():
                    safe_click(page, loc, f"menu '{label}'")
                    clicked = True
                    nap(page, 1500)
                    break
            except Exception:
                continue
        if not clicked:
            dump_controls(page, f"menu_{label.replace(' ', '_')[:30]}")
            raise RuntimeError(f"Menu item '{label}' not found - fix the menu list at the top "
                               f"of this section")
    nap(page, 3000)
    clear_overlays(page)

    if wait_for:
        deadline = time.time() + 60
        while time.time() < deadline:
            for f in page.frames:
                try:
                    if f.locator(wait_for).count() > 0:
                        return f
                except Exception:
                    continue
            nap(page, 1000)
        raise RuntimeError(f"The screen ({wait_for}) did not open")
    return frame

def save_page_evidence(page, folder: Path, stem: str):
    """Screenshot the current screen (SOP 5: 'save a screenshot or print to PDF')."""
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{stem}.png"
    page.screenshot(path=str(path), full_page=True)
    print(f"Saved: {path}")
    return path

# ---------- Section 4: return copies ----------
def process_return_copies(page, period_from, period_to, save_folder: Path, tin, password):
    """SOP 4: View Return History, one search per return type, download each
    return with its schedules."""
    folder = save_folder / "return_copies"
    tax_key = "Return copies"

    for return_type in RETURN_TYPES:
        if already_done(save_folder, tax_key, return_type):
            print(f"Returns - {return_type}: already done - skipping.")
            continue

        print(f"\n--- Return copies: {return_type} ---")
        try:
            frame = open_service(page, RETURN_HISTORY_MENU, tin=tin, password=password)
        except Exception as e:
            print(f"Could not open View Return History: {e}")
            return

        set_field(page, frame, ["Return Type", "#returnType", "select[name*='return' i]"],
                  return_type, "Return Type")
        set_field(page, frame, ["Return For", "#retBrnch", "select[name*='branch' i]"],
                  "Both", "Return For")

        mode, date_from, date_to = ask_dates(
            f"Return copies for {return_type}: which period?", "Period From / Period To")
        if mode == "q":
            print("Return copies stopped.")
            return
        if mode == "s":
            mark_progress(save_folder, tax_key, return_type, "skipped by user")
            continue
        if mode == "dates":
            set_form_date(frame, "form", "Period From", date_from)
            set_form_date(frame, "form", "Period To", date_to)

        click_button(page, frame, ["Search", "Submit"], "search")
        nap(page, 5000)
        clear_overlays(page)

        if frame_shows_no_records(frame):
            print(f"{return_type}: no returns filed for this period.")
            mark_progress(save_folder, tax_key, return_type, "no records")
            continue

        type_folder = folder / return_type.replace(" ", "_").replace("(", "").replace(")", "")
        files = generate_and_collect(page, frame, type_folder,
                                     f"{tin}_{return_type.replace(' ', '_')}")
        if not files:
            files = []
            if download_all_links(page, frame, type_folder,
                                  f"{tin}_{return_type.replace(' ', '_')}.pdf"):
                files = list(type_folder.iterdir())
        mark_progress(save_folder, tax_key, return_type,
                      f"{len(files)} saved" if files else "download failed")

# ---------- Section 5: taxpayer profile / registration details ----------
# ---------- Section 6: TIN / VAT certificate ----------
# ---------- Section 10: PRN payment slips ----------
PRN_COLUMN_HINTS = ("payment registration", "prn")

def prns_from_downloads(save_folder: Path):
    """Collect the PRNs out of the Level 3 files already downloaded, so the
    slips match the payments in the ledgers."""
    found = {}
    for path in sorted(save_folder.glob("*_L3_*.xls*")):
        # column-based read is more reliable than scanning every cell
        try:
            table = pd.read_excel(path, dtype=str)
        except Exception:
            continue
        for column in table.columns:
            if any(h in str(column).lower() for h in PRN_COLUMN_HINTS):
                for value in table[column].dropna():
                    value = str(value).strip()
                    if len(value) >= 8 and value.upper() == value:
                        found.setdefault(value, path.stem)
    print(f"{len(found)} PRN(s) found in the downloaded Level 3 files")
    return found

def process_payment_slips(page, save_folder: Path, tin, password):
    """SOP 10: re-print the payment slip for each PRN in the Level 3 ledgers."""
    tax_key = "PRN payment slips"
    prns = prns_from_downloads(save_folder)
    if not prns:
        print("No PRNs found - run the Level 3 downloads first, or check the column name "
              "in PRN_COLUMN_HINTS.")
        return

    folder = save_folder / "prn_payment_slips"
    for prn, source in prns.items():
        if already_done(save_folder, tax_key, prn):
            print(f"PRN {prn}: already done - skipping.")
            continue
        print(f"\n--- Payment slip for PRN {prn} (from {source}) ---")
        try:
            frame = open_service(page, PAYMENT_SLIP_MENU, tin=tin, password=password)
        except Exception as e:
            print(f"Could not open Re-Print a Payment Slip: {e}")
            return

        if not set_field(page, frame, ["PRN", "#prn", "input[name*='prn' i]"], prn, "PRN"):
            dump_controls(page, "payment_slip_form")
            return
        click_button(page, frame, ["Submit", "Search", "Print"], "submit")
        nap(page, 4000)
        clear_overlays(page)

        files = generate_and_collect(page, frame, folder, f"{tin}_PRN_{prn}")
        mark_progress(save_folder, tax_key, prn,
                      "saved" if files else "download failed")

# ---------- Sections 5 & 6: taxpayer profile and TIN/VAT certificate ----------
PROFILE_MENU = "#mailbox ul li:nth-of-type(13) a"                  # Taxpayer Profile
MY_PROFILE = "li:nth-of-type(13) > div > div > ul:nth-of-type(1) a"  # My Profile
TIN_CERT_LINK = "li:nth-of-type(13) ul:nth-of-type(3) a"           # Print TIN Certificate
VIEW_REG_DETAILS = "#submit2"                                      # View Registration Details

def open_mailbox_menu(page, tin=None, password=None, timeout=60):
    """Get the e-Services menu open and populated, and return its frame."""
    frame = go_to_portal_home(page, tin, password)
    deadline = time.time() + timeout
    while time.time() < deadline:
        if any_overlay(page):
            clear_overlays(page)
        frame = portal_menu_frame(page) or frame
        try:
            if frame.locator("#mailbox a").count() > 0:
                return frame
        except Exception:
            pass
        for loc in (frame.locator("#mailbox header h3"),
                    frame.get_by_text("e-Services", exact=True)):
            try:
                if loc.count() > 0:
                    js_click(loc)          # no hover, no actionability wait
                    break
            except Exception:
                continue
        # wait for the menu to actually fill in rather than sleeping
        wait_find(page, "#mailbox a", timeout=6000, visible=False)
    raise RuntimeError("The e-Services menu never opened (no links under #mailbox)")

def open_profile_item(page, item_selector, item_text, tin=None, password=None):
    """e-Services > hover 'Taxpayer Profile' > <My Profile | Print TIN Certificate>."""
    frame = open_mailbox_menu(page, tin, password)

    # Hover the Taxpayer Profile entry; its submenu only fills in on hover.
    parent = None
    for loc in (frame.locator(PROFILE_MENU),
                frame.get_by_role("link", name="Taxpayer Profile", exact=True),
                frame.locator("#mailbox a:has-text(\'Taxpayer Profile\')")):
        try:
            if loc.count() > 0:
                parent = loc.first
                break
        except Exception:
            continue
    if parent is None:
        raise RuntimeError("'Taxpayer Profile' is not in the e-Services menu")

    item = None
    for attempt in range(1, 4):
        try:
            parent.scroll_into_view_if_needed()
            parent.hover(force=True)      # the sub-menu opens on hover
            wait_find(page, item_selector, timeout=2500, visible=False)
        except Exception:
            pass

        for loc in (frame.locator(item_selector),
                    frame.locator(f"#mailbox a:has-text('{item_text}')"),
                    frame.get_by_role("link", name=item_text, exact=True)):
            try:
                if loc.count() > 0:
                    item = loc.first
                    break
            except Exception:
                continue
        if item is not None:
            print(f"Found '{item_text}' (hover attempt {attempt}).")
            return frame, item

        try:                       # hovering alone did not open it - click it
            safe_click(page, parent, "Taxpayer Profile menu")
            nap(page, 1500)
        except Exception:
            pass

    print(f"'{item_text}' not found - links currently in the menu:")
    for f in page.frames:
        try:
            links = f.evaluate(
                "() => [...document.querySelectorAll('#mailbox a')]"
                ".map(a => (a.innerText || '').trim()).filter(Boolean).slice(0, 60)")
        except Exception:
            continue
        if links:
            print(f"  {f.url[:70]}: {links}")
    raise RuntimeError(f"'{item_text}' not found in the Taxpayer Profile menu")

def page_to_pdf(target_page, dest: Path):
    """Print a page to PDF through Chrome DevTools - works with the browser
    window open, unlike page.pdf(), and avoids the native print dialog."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    session = target_page.context.new_cdp_session(target_page)
    result = session.send("Page.printToPDF", {"printBackground": True,
                                              "paperWidth": 8.27, "paperHeight": 11.69})
    dest.write_bytes(base64.b64decode(result["data"]))
    print(f"Saved: {dest}")
    return dest

def save_pdf_page(page, new_page, dest: Path):
    """Save a PDF that opened in a tab: fetch it with the logged-in session,
    falling back to printing the page."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    url = new_page.url
    if url and url.lower().startswith("http"):
        try:
            body = page.context.request.get(url).body()
            if body[:4] == b"%PDF":
                dest.write_bytes(body)
                print(f"Saved: {dest}")
                return dest
        except Exception as e:
            print(f"Could not fetch the PDF directly ({type(e).__name__})")
    try:
        return page_to_pdf(new_page, dest)
    except Exception as e:
        print(f"Could not save the certificate ({type(e).__name__})")
        return None

def process_taxpayer_profile(page, save_folder: Path, tin, password):
    """SOP 5: My Profile > View Registration Details > save the page as PDF."""
    key = "Taxpayer Profile"
    stamp = datetime.date.today().isoformat()
    dest = save_folder / "profile" / f"{tin}_Registration_{stamp}.pdf"
    if already_done(save_folder, key, stamp) and dest.exists():
        print("Taxpayer profile: already downloaded - skipping.")
        record("Taxpayer profile", "Registration details", "already downloaded", file=dest)
        return

    frame, item = open_profile_item(page, MY_PROFILE, "My Profile", tin, password)
    safe_click(page, item, "My Profile")
    nap(page, 4000)
    clear_overlays(page)

    detail_frame = find_frame_containing_selector(page, VIEW_REG_DETAILS, timeout=60)
    # the details open in a new tab (dealerProfileDetails)
    try:
        with page.context.expect_page(timeout=20000) as new_tab:
            safe_click(page, detail_frame.locator(VIEW_REG_DETAILS), "View Registration Details")
        target = new_tab.value
        target.wait_for_load_state("domcontentloaded")
    except Exception:
        target = page                       # opened in the same tab
        nap(page, 4000)

    clear_overlays(page)
    saved = None
    try:
        saved = page_to_pdf(target, dest)   # instead of #printid, which opens the print dialog
    except Exception as e:
        print(f"Could not print the profile ({type(e).__name__}) - saving a screenshot instead")
        try:
            target.screenshot(path=str(dest.with_suffix(".png")), full_page=True)
            saved = dest.with_suffix(".png")
            print(f"Saved: {saved}")
        except Exception:
            pass
    if target is not page:
        try:
            target.close()
        except Exception:
            pass
    mark_progress(save_folder, key, stamp, "saved" if saved else "failed")
    record("Taxpayer profile", "Registration details", "saved" if saved else "failed", file=saved or "")

def process_tin_certificate(page, save_folder: Path, tin, password):
    """SOP 6: Taxpayer Profile > Print TIN Certificate > save the PDF."""
    key = "TIN Certificate"
    dest = save_folder / "profile" / f"{tin}_TIN_Certificate.pdf"
    if dest.exists():
        print("TIN certificate: already downloaded - skipping.")
        record("TIN certificate", "TIN / VAT certificate", "already downloaded", file=dest)
        return

    frame, item = open_profile_item(page, TIN_CERT_LINK, "Print TIN Certificate", tin, password)

    saved = None
    try:                                     # it may arrive as a plain download
        with page.expect_download(timeout=15000) as dl:
            safe_click(page, item, "Print TIN Certificate")
        dl.value.save_as(str(dest))
        dest.parent.mkdir(parents=True, exist_ok=True)
        print(f"Saved: {dest}")
        saved = dest
    except Exception:
        nap(page, 4000)
        pages = page.context.pages
        if len(pages) > 1:                   # it opened in Chrome's PDF viewer
            saved = save_pdf_page(page, pages[-1], dest)
            try:
                pages[-1].close()
            except Exception:
                pass
        else:
            saved = save_pdf_page(page, page, dest)

    mark_progress(save_folder, key, "certificate", "saved" if saved else "failed")
    record("TIN certificate", "TIN / VAT certificate", "saved" if saved else "failed", file=saved or "")

# ============================================================================
# Runner - one function used by both the window and the old console mode
# ============================================================================

OTHER_JOBS = {
    "profile": "Taxpayer profile (registration details)",
    "tin_cert": "TIN / VAT certificate",
    "wht_tcc": "Withholding tax credit certificates",
    "paye_tcc": "PAYE tax credit certificates",
}

def write_summary(save_folder: Path):
    """Print the run summary and save it next to the downloads (plain Excel,
    no fills, no frozen panes; CSV if openpyxl is not installed)."""
    if not RESULTS:
        return None
    df = pd.DataFrame(RESULTS, columns=["Section", "Item", "Status", "Detail", "File"])
    print("\n=== Summary ===")
    for status, n in df["Status"].value_counts().items():
        print(f"  {status}: {n}")
    attention = df[df["Status"].isin(["failed", "generate failed", "no result rows", "note"])]
    for _, r in attention.iterrows():
        print(f"  CHECK: {r['Section']} - {r['Item']} - {r['Status']} {r['Detail']}")
    save_folder.mkdir(parents=True, exist_ok=True)
    stamp = f"{datetime.datetime.now():%Y%m%d_%H%M%S}"
    path = save_folder / f"_Run_Summary_{stamp}.xlsx"
    try:
        from engines.sheet_safety import neutralise_sheet
        with pd.ExcelWriter(path, engine="openpyxl") as xw:
            df.to_excel(xw, index=False, sheet_name="Summary")
            ws = xw.sheets["Summary"]
            neutralise_sheet(ws)
            for col, width in zip("ABCDE", (32, 44, 24, 40, 60)):
                ws.column_dimensions[col].width = width
    except Exception:
        from engines.sheet_safety import csv_text
        path = path.with_suffix(".csv")
        df.astype(str).map(csv_text).to_csv(path, index=False)
    print(f"Summary saved: {path}")
    return path

def launch_browser(p):
    launch_args = ["--disable-dev-shm-usage", "--disable-extensions",
                   "--disable-background-timer-throttling", "--start-maximized"]
    choices = [("chrome", "Google Chrome"), ("msedge", "Microsoft Edge"), (None, "bundled Chromium")]
    if BROWSER_CHANNEL:
        choices = [c for c in choices if c[0] == BROWSER_CHANNEL] + [(None, "bundled Chromium")]
    for channel, name in choices:
        try:
            kw = {"headless": False, "args": launch_args}   # visible: you type the security code
            if PROXY:
                kw["proxy"] = {"server": PROXY}
            if channel:
                kw["channel"] = channel
            browser = p.chromium.launch(**kw)
            print(f"Using {name}.")
            return browser
        except Exception as e:
            print(f"{name} not available ({type(e).__name__}).")
    raise RuntimeError("No browser found - install Google Chrome, or run: "
                       "playwright install chromium")

def _job_setup(settings):
    """Shared by run_job() and run_job_on_page(): reads the settings into the
    module globals and returns (tin, password, period_from, period_to,
    tax_types, other, save_folder)."""
    tin, password = settings["tin"], settings["password"]
    CREDS["tin"], CREDS["password"] = tin, password
    period_from, period_to = settings["period_from"], settings["period_to"]
    tax_types = settings["tax_types"]
    other = settings["other_jobs"]
    global SKIP_NIL, ANNUAL_L3, ANNUAL_ALL_YEARS
    SKIP_NIL = settings.get("skip_nil", False)
    ANNUAL_L3 = settings.get("annual_l3", False)
    ANNUAL_ALL_YEARS = settings.get("annual_all_years", False)
    RESULTS.clear()
    NOTES.clear()
    TIMINGS.clear()
    ROUTE_STATE.clear()
    global CURRENT_TIN, REFRESH_L2, KEEP_MONTH_FILES
    CURRENT_TIN = tin
    REFRESH_L2 = settings.get("refresh_l2", True)
    KEEP_MONTH_FILES = settings.get("keep_month_files", False)   # settings saved before 10.12: off
    global L3_ROUTE, DIRECT_SWITCH_AFTER_S
    L3_ROUTE = settings.get("l3_route", "auto")
    DIRECT_SWITCH_AFTER_S = int(settings.get("direct_switch_s", 0) or 0)
    DIRECT_STATE.clear()
    folder_name = safe_part(tin)
    save_folder = Path(OUTPUT_ROOT) / folder_name if OUTPUT_ROOT else app_dir() / "downloads" / folder_name
    # error_*.png / debug files go to WORK_DIR by absolute path (work_file) - no os.chdir
    from engines.excel_files import repair_folder
    fixed = repair_folder(save_folder)
    if fixed:
        print(f"Repaired {fixed} ledger file(s) saved by an earlier version (portal table -> real Excel).")
    print(f"Client TIN: {tin}")
    print(f"Period: {period_from} to {period_to}")
    print(f"Tax types: {', '.join(tax_types) or '(none)'}")
    print(f"Other documents: {', '.join(OTHER_JOBS[k] for k in other) or '(none)'}")
    print(f"Saving to: {save_folder}")
    return tin, password, period_from, period_to, tax_types, other, save_folder


def update_tax_head_files(save_folder: Path, tax_types):
    """User, 6 Oct 2026: one workbook per tax head, one table, updated in place (engines/ledger_taxhead.py).
    User, 8 Oct 2026: no Excel file per month - the workbook has Summary, Level 3-2 and a tab per month
    (like the WHT workbook) and the monthly files are removed once they are inside it.
    A failure is recorded as 'check' and never fails the step; the monthly files then stay on disk."""
    from engines import ledger_taxhead
    for tax in tax_types:
        if tax not in TAX_TYPES and tax not in ANNUAL_TAX_TYPES:
            continue
        statuses = {r["Item"]: (r["Status"], r["Detail"]) for r in RESULTS
                    if r["Section"] == f"{tax} ledger" and ledger_taxhead.is_period(r["Item"])}
        checks = []
        try:
            path = ledger_taxhead.update_tax_head(save_folder, tax, statuses=statuses,
                                                  keep_month_files=KEEP_MONTH_FILES, checks=checks)
            for period, why in checks:            # Level 2 payment vs Level 3-2 lines (10.13)
                print(f"{tax} {period}: {why}")
                record(f"{tax} ledger", period, "check", why[7:] if why.startswith("check: ") else why,
                       file=path or "")
            if path is not None:
                record("Ledgers", f"{tax} workbook", "saved", file=path)
        except Exception as e:
            msg = (str(e).splitlines() or [""])[0][:200]
            print(f"WARNING: {tax} workbook not updated: {msg}")
            record("Ledgers", f"{tax} workbook", "check", detail=msg)


def build_combined_workbook(save_folder: Path, tin, period_from, period_to, tax_types):
    """Ledgers_<TIN>_<from>_to_<to>.xlsx from the files on disk (engines/ledger_workbook.py).
    Local work only - no portal time.  A failure here never fails the ledger step: it is
    recorded as 'check' so the person sees it on the Review page."""
    from engines import ledger_workbook
    try:
        path = ledger_workbook.build(save_folder, tin, period_from, period_to, tax_types,
                                     generate_month_labels(period_from, period_to), ANNUAL_TAX_TYPES,
                                     results=list(RESULTS), progress=load_progress(save_folder))
    except Exception as e:
        print(f"Combined ledgers workbook not built ({type(e).__name__}: {e})")
        record("Ledgers workbook", "combined workbook", "check",
               f"not built ({type(e).__name__}: {e})"[:200])
        return None
    if path:
        print(f"Saved: {path}")
        record("Ledgers workbook", "combined workbook", "saved",
               "Contents + Level 2 and All Level 3 tabs per tax type", file=path)
    return path


def run_job_on_page(page, settings):
    """The ledger / certificate / profile work on a page that is ALREADY
    logged in.  Used by the suite so one login serves every module.
    Raises StopRequested when Stop is pressed; returns the failed list."""
    tin, password, period_from, period_to, tax_types, other, save_folder = _job_setup(settings)
    failed = []
    try:
        debug_dump(page, "after_login")
        dismiss_popup_if_present(page)

        if tax_types:
            print("\n=== Tax Statement (ledgers) ===")
            open_tax_statement(page)
            for tax_type in tax_types:
                print(f"\n=== Starting tax type: {tax_type} ===")
                for attempt in range(1, 4):
                    try:
                        if tax_type in ANNUAL_TAX_TYPES:
                            process_annual_tax_type(page, tax_type, save_folder,
                                                    tin=tin, password=password,
                                                    period_from=period_from,
                                                    period_to=period_to)
                        else:
                            process_tax_type(page, tax_type, period_from, period_to,
                                             save_folder, tin=tin, password=password)
                        break
                    except Exception as e:
                        from core.runner import BrowserClosed, browser_gone
                        if browser_gone(page, e):          # no retries, no further tax types
                            raise BrowserClosed(f"The browser window was closed during the {tax_type} ledger") from e
                        print(f"{tax_type} failed (attempt {attempt}/3): {e}")
                        safe_screenshot(page, f"error_{tax_type.replace(' ', '_')}_{attempt}.png")
                        if attempt == 3:
                            failed.append(tax_type)
                            record(f"{tax_type} ledger", "whole tax type", "failed", str(e)[:200])
                            break
                        if ONLINE_WAIT:
                            ONLINE_WAIT()
                        try:
                            recover_to_level1(page, tin, password)
                        except Exception as rec:
                            # 5 Oct 2026: this error escaped and ended the whole Ledgers step -
                            # attempts 2-3 and every remaining tax type were skipped.
                            print(f"Could not get back to Level 1 ({(str(rec).splitlines() or [''])[0][:150]}) "
                                  f"- attempt {attempt + 1} starts by opening the Tax Statement again.")

        profile_jobs = [(process_taxpayer_profile, "profile", "taxpayer profile"),
                        (process_tin_certificate, "tin_cert", "TIN/VAT certificate")]
        if any(k in other for _, k, _ in profile_jobs):
            print("\n=== Sections 5 & 6: taxpayer profile and TIN/VAT certificate ===")
            for job, key, label in profile_jobs:
                if key not in other:
                    continue
                try:
                    job(page, save_folder, tin, password)
                except Exception as e:
                    print(f"{label} failed: {e}")
                    failed.append(label)
                    record(label, label, "failed", str(e)[:200])
                    safe_screenshot(page, f"error_{label.split()[0]}.png")
                    clear_overlays(page)

        cert_jobs = [(process_wht_certificates, "wht_tcc", "withholding tax credit certificates"),
                     (process_paye_certificates, "paye_tcc", "PAYE tax credit certificates")]
        if any(k in other for _, k, _ in cert_jobs):
            print("\n=== Section 9: tax credit certificates ===")
            go_to_portal_home(page, tin, password)
            for job, key, label in cert_jobs:
                if key not in other:
                    continue
                for attempt in range(1, 3):
                    try:
                        job(page, period_from, period_to, save_folder, tin, password)
                        break
                    except Exception as e:
                        print(f"{label} failed (attempt {attempt}/2): {e}")
                        safe_screenshot(page, f"error_section9_{attempt}.png")
                        clear_overlays(page)
                        if attempt == 2:
                            failed.append(label)
                            record(label, label, "failed", str(e)[:200])
                        go_home_inside_portal(page)
                        if login_form_present(page):
                            login(page, tin, password)

        print("\n=== Finished ===")
        if failed:
            print("Not finished: " + ", ".join(failed) + " - run again to resume "
                  "(files already saved are skipped).")
        else:
            print("Everything selected was processed.")
        print(f"Files are in: {save_folder}")
    finally:
        try:
            summarise_level3_checks(save_folder)
        except Exception as e:
            print(f"Could not summarise the Level 3 times ({type(e).__name__})")
        if tax_types:
            update_tax_head_files(save_folder, tax_types)
        try:
            write_summary(save_folder)
        except Exception as e:
            print(f"Could not write the summary ({type(e).__name__}: {e})")
    return failed


def run_job(settings):
    """Stand-alone: start a browser, log in, run the job, close."""
    tin, password = settings["tin"], settings["password"]
    failed = []
    with sync_playwright() as p:
        browser = launch_browser(p)
        try:
            context = browser.new_context(accept_downloads=True,
                                          viewport={"width": 1400, "height": 900})
            speed_up(context)
            page = context.new_page()
            page.set_default_timeout(NAV_TIMEOUT_MS)
            print("\n=== Login ===")
            login(page, tin, password)
            dismiss_cookie_banner(page)
            dismiss_pre_login_popup(page)
            failed = run_job_on_page(page, settings)
        except StopRequested:
            print("\n=== Stopped by you ===  Run again to resume; saved files are skipped.")
        except Exception as e:
            print("Something went wrong:", e)
            traceback.print_exc()
            safe_screenshot(page, "error_screenshot.png")
        finally:
            try:
                browser.close()
            except Exception:
                pass
    return failed

def console_main():
    """The old behaviour: python ura_bot.py --console (uses .env, all tax types)."""
    tin, password = os.getenv("URA_TIN"), os.getenv("URA_PASSWORD")
    if not tin or not password:
        print("URA_TIN / URA_PASSWORD are missing from .env next to this file.")
        input("Press Enter to close...")
        return
    period_from, period_to = get_current_period()
    run_job({"tin": tin, "password": password, "period_from": period_from,
             "period_to": period_to, "tax_types": TAX_TYPES + ANNUAL_TAX_TYPES,
             "other_jobs": list(OTHER_JOBS)})
    input("Press Enter to close...")

# ============================================================================
# The window
# ============================================================================

class LogWriter:
    """Sends print() output to the window and to a log file."""
    def __init__(self, ui_queue, log_path):
        self.q = ui_queue
        self.buf = ""
        self.file = open(log_path, "a", encoding="utf-8")

    def write(self, text):
        self.file.write(text)
        self.file.flush()
        self.buf += text
        while "\n" in self.buf:
            line, self.buf = self.buf.split("\n", 1)
            self.q.put(("log", line))

    def flush(self):
        pass

def gui_main():
    import tkinter as tk
    from tkinter import ttk, messagebox
    from tkinter.scrolledtext import ScrolledText
    from dotenv import set_key

    ui_q = queue.Queue()
    reply_q = queue.Queue()
    state = {"worker": None, "log_path": None}

    root = tk.Tk()
    root.title("URA Tax Statement Downloader")
    root.geometry("920x840")
    root.minsize(820, 680)
    style = ttk.Style()
    try:
        style.theme_use("vista")
    except Exception:
        pass
    style.configure("Head.TLabel", font=("Arial", 10, "bold"))
    style.configure("Ask.TLabel", font=("Arial", 10, "bold"), foreground="#8a4b00")

    top = ttk.Frame(root, padding=8)
    top.pack(fill="x")

    # --- Login ---
    lf = ttk.LabelFrame(top, text=" Login ", padding=8)
    lf.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
    ttk.Label(lf, text="TIN / Login ID").grid(row=0, column=0, sticky="w")
    tin_var = tk.StringVar(value=os.getenv("URA_TIN", ""))
    ttk.Entry(lf, textvariable=tin_var, width=24).grid(row=0, column=1, pady=2)
    ttk.Label(lf, text="Password").grid(row=1, column=0, sticky="w")
    pw_var = tk.StringVar(value=os.getenv("URA_PASSWORD", ""))
    ttk.Entry(lf, textvariable=pw_var, width=24, show="*").grid(row=1, column=1, pady=2)
    remember_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(lf, text="Remember in .env (plain text)",
                    variable=remember_var).grid(row=2, column=0, columnspan=2, sticky="w")

    # --- Period ---
    pf = ttk.LabelFrame(top, text=" Period (YYYY-MM) ", padding=8)
    pf.grid(row=0, column=1, sticky="nsew", padx=(0, 6))
    d_from, d_to = get_current_period()
    from_var, to_var = tk.StringVar(value=d_from), tk.StringVar(value=d_to)
    ttk.Label(pf, text="From").grid(row=0, column=0, sticky="w")
    ttk.Entry(pf, textvariable=from_var, width=10).grid(row=0, column=1, pady=2)
    ttk.Label(pf, text="To").grid(row=1, column=0, sticky="w")
    ttk.Entry(pf, textvariable=to_var, width=10).grid(row=1, column=1, pady=2)
    def last18():
        a, b = get_current_period()
        from_var.set(a); to_var.set(b)
    ttk.Button(pf, text="Last 18 months", command=last18).grid(row=2, column=0, columnspan=2, pady=2)

    # --- Options ---
    of = ttk.LabelFrame(top, text=" Options ", padding=8)
    of.grid(row=0, column=2, sticky="nsew")
    speed_var = tk.BooleanVar(value=SPEED_UPS)
    ttk.Checkbutton(of, text="Speed-ups (block chat widget\nand third-party images)",
                    variable=speed_var).pack(anchor="w")
    debug_var = tk.BooleanVar(value=DEBUG_DUMPS)
    ttk.Checkbutton(of, text="Debug (slower: screenshots\nand blocked-request lines)",
                    variable=debug_var).pack(anchor="w")
    top.columnconfigure(0, weight=1)

    # --- Tax types ---
    mid = ttk.Frame(root, padding=(8, 0))
    mid.pack(fill="x")
    tf = ttk.LabelFrame(mid, text=" Tax ledgers to download (Reports > Tax Statement) ", padding=8)
    tf.pack(side="left", fill="both", expand=True, padx=(0, 6))
    tax_vars = {}
    for i, t in enumerate(TAX_TYPES + ANNUAL_TAX_TYPES):
        v = tk.BooleanVar(value=True)
        tax_vars[t] = v
        note = "  (Level 2 only, annual)" if t in ANNUAL_TAX_TYPES else "  (Level 2 + monthly Level 3)"
        ttk.Checkbutton(tf, text=t + note, variable=v).grid(row=i, column=0, sticky="w")
    skip_nil_var = tk.BooleanVar(value=False)
    annual_l3_var = tk.BooleanVar(value=False)
    annual_all_var = tk.BooleanVar(value=False)
    r0 = len(tax_vars)
    ttk.Separator(tf).grid(row=r0, column=0, sticky="ew", pady=4)
    ttk.Checkbutton(tf, text="Skip Level 3 for months with nil payment (faster)",
                    variable=skip_nil_var).grid(row=r0 + 1, column=0, sticky="w")
    ttk.Checkbutton(tf, text="Income Tax & Rental: also Level 3 (payment details)",
                    variable=annual_l3_var).grid(row=r0 + 2, column=0, sticky="w")
    ttk.Checkbutton(tf, text="     ...for all years, not only years in the period (slower)",
                    variable=annual_all_var).grid(row=r0 + 3, column=0, sticky="w")
    bf = ttk.Frame(tf)
    bf.grid(row=r0 + 4, column=0, sticky="w", pady=(4, 0))
    ttk.Button(bf, text="Select all", command=lambda: [v.set(True) for v in tax_vars.values()]).pack(side="left")
    ttk.Button(bf, text="Clear", command=lambda: [v.set(False) for v in tax_vars.values()]).pack(side="left", padx=4)

    xf = ttk.LabelFrame(mid, text=" Other documents ", padding=8)
    xf.pack(side="left", fill="both", expand=True)
    other_vars = {}
    for i, (k, label) in enumerate(OTHER_JOBS.items()):
        v = tk.BooleanVar(value=True)
        other_vars[k] = v
        ttk.Checkbutton(xf, text=label, variable=v).grid(row=i, column=0, sticky="w")
    for j, label in enumerate(["Return copies (screens not recorded yet)",
                               "PRN payment slips (screens not recorded yet)"]):
        cb = ttk.Checkbutton(xf, text=label)
        cb.state(["disabled", "!selected"])
        cb.grid(row=len(OTHER_JOBS) + j, column=0, sticky="w")

    # --- Buttons ---
    btns = ttk.Frame(root, padding=8)
    btns.pack(fill="x")
    start_btn = ttk.Button(btns, text="▶  Start")
    start_btn.pack(side="left")
    stop_btn = ttk.Button(btns, text="■  Stop")
    stop_btn.pack(side="left", padx=6)
    stop_btn.state(["disabled"])
    def open_folder():
        tin = tin_var.get().strip()
        folder = app_dir() / "downloads" / tin if tin else app_dir() / "downloads"
        folder.mkdir(parents=True, exist_ok=True)
        os.startfile(str(folder)) if hasattr(os, "startfile") else None
    ttk.Button(btns, text="Open downloads folder", command=open_folder).pack(side="left")
    def open_log():
        if state["log_path"] and hasattr(os, "startfile"):
            os.startfile(str(state["log_path"]))
    ttk.Button(btns, text="Open log file", command=open_log).pack(side="left", padx=6)
    status_var = tk.StringVar(value="Ready.")
    ttk.Label(btns, textvariable=status_var).pack(side="right")

    # --- Question panel (shown only when the bot needs you) ---
    ask = tk.Frame(root, bg="#fff4d6", bd=1, relief="solid", padx=8, pady=8)
    ask_msg = tk.Label(ask, bg="#fff4d6", font=("Arial", 10, "bold"), fg="#5a3a00",
                       wraplength=840, justify="left", anchor="w")
    ask_msg.pack(fill="x")
    ask_row = tk.Frame(ask, bg="#fff4d6")
    ask_row.pack(fill="x", pady=(6, 0))
    ask_from, ask_to = tk.StringVar(), tk.StringVar()

    # --- Log ---
    logbox = ScrolledText(root, height=18, font=("Consolas", 9), state="disabled", wrap="word")
    logbox.pack(fill="both", expand=True, padx=8, pady=(0, 8))

    def log(line):
        logbox.configure(state="normal")
        logbox.insert("end", line + "\n")
        logbox.see("end")
        logbox.configure(state="disabled")

    def answer(value):
        for w in ask_row.winfo_children():
            w.destroy()
        ask.pack_forget()
        root.unbind("<Return>")
        reply_q.put(value)
        status_var.set("Running...")

    def show_question(message, kind, defaults):
        for w in ask_row.winfo_children():
            w.destroy()
        ask_msg.configure(text=message)
        B = lambda text, val, **kw: tk.Button(ask_row, text=text, command=lambda: answer(val),
                                              padx=10, **kw)
        if kind == "login":
            B("Log In  (security code typed)", "", bg="#2e7d32", fg="white").pack(side="left")
            B("Cancel run", "q").pack(side="left", padx=6)
            root.bind("<Return>", lambda e: answer(""))
        elif kind == "dates":
            ask_from.set(defaults[0]); ask_to.set(defaults[1])
            tk.Label(ask_row, text="From (dd/mm/yyyy)", bg="#fff4d6").pack(side="left")
            ttk.Entry(ask_row, textvariable=ask_from, width=12).pack(side="left", padx=4)
            tk.Label(ask_row, text="To", bg="#fff4d6").pack(side="left")
            ttk.Entry(ask_row, textvariable=ask_to, width=12).pack(side="left", padx=4)
            def use_dates():
                try:
                    a = datetime.datetime.strptime(ask_from.get().strip(), "%d/%m/%Y")
                    b = datetime.datetime.strptime(ask_to.get().strip(), "%d/%m/%Y")
                except ValueError:
                    messagebox.showerror("Dates", "Please type both dates as dd/mm/yyyy.")
                    return
                if a > b:
                    messagebox.showerror("Dates", "'From' is after 'To'.")
                    return
                answer(("dates", ask_from.get().strip(), ask_to.get().strip()))
            tk.Button(ask_row, text="Use these dates", command=use_dates, padx=10,
                      bg="#2e7d32", fg="white").pack(side="left", padx=4)
            B("I set them in the browser", ("", None, None)).pack(side="left", padx=4)
            B("Skip", ("s", None, None)).pack(side="left", padx=4)
            B("Stop this section", ("q", None, None)).pack(side="left", padx=4)
        else:   # handover / continue
            B("Done - continue", "", bg="#2e7d32", fg="white").pack(side="left")
            B("Skip", "s").pack(side="left", padx=6)
            B("Stop this section", "q").pack(side="left")
            root.bind("<Return>", lambda e: answer(""))
        ask.pack(fill="x", padx=8, pady=(0, 8), before=logbox)
        status_var.set("Waiting for you - see the yellow box.")
        root.deiconify(); root.lift()
        root.attributes("-topmost", True)
        root.after(1200, lambda: root.attributes("-topmost", False))
        root.bell()

    def worker_prompt(message, kind, defaults, poll):
        """Runs in the worker thread: ask the window and wait for the answer,
        while still checking Stop and (for login) the browser itself."""
        while not reply_q.empty():
            reply_q.get_nowait()
        ui_q.put(("ask", (message, kind, defaults)))
        print(">>> " + message)
        while True:
            try:
                return reply_q.get(timeout=0.5)
            except queue.Empty:
                pass
            if STOP_EVENT.is_set():
                ui_q.put(("hide_ask", None))
                return ("q", None, None) if kind == "dates" else "q"
            if poll is not None:
                try:
                    if poll():
                        ui_q.put(("hide_ask", None))
                        return "done_in_browser"
                except Exception:
                    pass

    def pump():
        try:
            while True:
                kind, data = ui_q.get_nowait()
                if kind == "log":
                    log(data)
                elif kind == "ask":
                    show_question(*data)
                elif kind == "hide_ask":
                    for w in ask_row.winfo_children():
                        w.destroy()
                    ask.pack_forget()
                    root.unbind("<Return>")
                elif kind == "done":
                    start_btn.state(["!disabled"])
                    stop_btn.state(["disabled"])
                    status_var.set(data)
        except queue.Empty:
            pass
        root.after(100, pump)

    def start():
        global SPEED_UPS, DEBUG_DUMPS, DEBUG_BLOCKING, PROMPT_HANDLER
        tin, pw = tin_var.get().strip(), pw_var.get()
        pf_, pt_ = from_var.get().strip(), to_var.get().strip()
        if not tin or not pw:
            messagebox.showerror("Login", "Enter the TIN / Login ID and the password.")
            return
        try:
            a = datetime.datetime.strptime(pf_, "%Y-%m")
            b = datetime.datetime.strptime(pt_, "%Y-%m")
        except ValueError:
            messagebox.showerror("Period", "Type the period as YYYY-MM, e.g. 2025-07.")
            return
        if a > b:
            messagebox.showerror("Period", "'From' is after 'To'.")
            return
        taxes = [t for t, v in tax_vars.items() if v.get()]
        others = [k for k, v in other_vars.items() if v.get()]
        if not taxes and not others:
            messagebox.showerror("Nothing selected", "Tick at least one item to download.")
            return
        if remember_var.get():
            env = app_dir() / ".env"
            env.touch(exist_ok=True)
            set_key(str(env), "URA_TIN", tin)
            set_key(str(env), "URA_PASSWORD", pw)

        SPEED_UPS, DEBUG_DUMPS = speed_var.get(), debug_var.get()
        DEBUG_BLOCKING = DEBUG_DUMPS
        PROMPT_HANDLER = worker_prompt
        STOP_EVENT.clear()

        logs = app_dir() / "logs"
        logs.mkdir(exist_ok=True)
        state["log_path"] = logs / f"run_{tin}_{datetime.datetime.now():%Y%m%d_%H%M%S}.txt"
        writer = LogWriter(ui_q, state["log_path"])
        sys.stdout = sys.stderr = writer

        settings = {"tin": tin, "password": pw, "period_from": pf_, "period_to": pt_,
                    "tax_types": taxes, "other_jobs": others,
                    "skip_nil": skip_nil_var.get(), "annual_l3": annual_l3_var.get(),
                    "annual_all_years": annual_all_var.get()}

        def work():
            try:
                failed = run_job(settings)
                msg = "Stopped." if STOP_EVENT.is_set() else (
                    "Finished with items not done - see log." if failed else "Finished.")
            except BaseException as e:
                print(f"Could not run: {e}")
                traceback.print_exc()
                msg = "Stopped with an error - see log."
            writer.write("\n")
            ui_q.put(("done", msg))

        start_btn.state(["disabled"])
        stop_btn.state(["!disabled"])
        status_var.set("Running...")
        logbox.configure(state="normal"); logbox.delete("1.0", "end"); logbox.configure(state="disabled")
        state["worker"] = threading.Thread(target=work, daemon=True)
        state["worker"].start()

    def stop():
        STOP_EVENT.set()
        status_var.set("Stopping after the current step...")
        log("Stop requested - finishing the current step, then closing the browser.")

    start_btn.configure(command=start)
    stop_btn.configure(command=stop)

    def on_close():
        w = state["worker"]
        if w is not None and w.is_alive():
            if not messagebox.askyesno("Still running", "A download is running. Stop it and close?"):
                return
            STOP_EVENT.set()
        root.destroy()
    root.protocol("WM_DELETE_WINDOW", on_close)

    pump()
    root.mainloop()

if __name__ == "__main__":
    if "--console" in sys.argv:
        console_main()
    else:
        gui_main()
