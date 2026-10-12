"""core/client_import.py - add many clients at once from an Excel file (user, 6 Oct 2026).

Columns (found by their heading, in any order, on the first sheet): Client name | Login ID (TIN) |
Password.  Each row is checked exactly like 'Add client' (core.store.client_problem): name, Login ID
and password filled in, Login ID digits only, no Login ID twice.  A Login ID already in the client list
is offered as an update (new name / password).  Passwords never appear in any message; they are saved
encrypted with the Windows account like every other client password.
"""
from dataclasses import dataclass, field
from pathlib import Path

from core import store

HEAD = ("Client name", "Login ID (TIN)", "Password")
EXAMPLE = ("Example Traders Limited", "1000000000", "type the password here")


@dataclass
class Problem:
    row: int
    reason: str


@dataclass
class Result:
    new: list = field(default_factory=list)
    updates: list = field(default_factory=list)
    problems: list = field(default_factory=list)
    unchanged: int = 0


def make_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"
    ws.append(list(HEAD))
    ws.append(list(EXAMPLE))
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, width in zip("ABC", (40, 20, 28)):
        ws.column_dimensions[col].width = width
    for r in range(2, 1001):                                  # Login ID kept as text: every digit stays
        ws[f"B{r}"].number_format = "@"
    ws.freeze_panes = "A2"
    note = wb.create_sheet("How to fill")
    for line in ("One client per row on the 'Clients' sheet; keep the headings in row 1.",
                 "Login ID (TIN): digits only, as used to log in to the URA portal.",
                 "The example row is ignored; you can delete it.",
                 "After importing, delete this file or keep it somewhere safe: it holds passwords."):
        note.append([line])
    note.column_dimensions["A"].width = 90
    path = Path(path)
    wb.save(path)
    return path


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                                            # a TIN typed as a number
    return str(v).strip()


def _columns(rows):
    """(heading row index, {key: column}) - the first row naming all three columns."""
    for i, r in enumerate(rows[:20]):
        found = {}
        for j, v in enumerate(r):
            t = " ".join(_text(v).lower().replace("(", " ").replace(")", " ").split())
            words = set(t.split())
            if "password" in words:
                key = "password"
            elif words & {"tin", "login"}:
                key = "login"
            elif "name" in words or t in ("client", "company", "taxpayer"):
                key = "name"
            else:
                continue
            found.setdefault(key, j)
        if len(found) == 3:
            return i, found
    raise ValueError("No heading row with the columns " + ", ".join(HEAD) +
                     " was found in the first 20 rows of the first sheet.")


def read_clients(path, existing):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    head, col = _columns(rows)
    res = Result()
    known = {str(c.get("login", "")).strip(): c for c in existing}
    seen = {}
    for i, r in enumerate(rows[head + 1:], start=head + 2):            # Excel row numbers
        get = (lambda k: _text(r[col[k]]) if col[k] < len(r) else "")
        c = {"name": get("name"), "login": get("login"), "password": get("password")}
        if not any(c.values()):
            continue
        if (c["name"], c["login"], c["password"]) == EXAMPLE:
            continue
        if c["login"] in seen:
            res.problems.append(Problem(i, f"Login ID {c['login']} is twice in the file (rows "
                                           f"{seen[c['login']]} and {i}); only row {seen[c['login']]} is used."))
            continue
        problem = store.client_problem(c, [])                         # filled in, digits only
        if problem:
            res.problems.append(Problem(i, problem))
            continue
        seen[c["login"]] = i
        old = known.get(c["login"])
        if old is None:
            res.new.append(c)
        elif old.get("name") != c["name"] or old.get("password") != c["password"]:
            res.updates.append(c)
        else:
            res.unchanged += 1
    return res


def apply(existing, res, update_existing):
    """The new client list (existing order kept, new clients at the end)."""
    out = [dict(c) for c in existing]
    if update_existing:
        by_login = {c["login"]: c for c in res.updates}
        for c in out:
            u = by_login.get(c["login"])
            if u:
                c.update(name=u["name"], password=u["password"])
                c.pop("password_unreadable", None)
                c.pop("_blob", None)
    out.extend(dict(c) for c in res.new)
    return out


def summary(res):
    lines = [f"{len(res.new)} new client(s) to add.",
             f"{len(res.updates)} already in the list with a different name or password.",
             f"{res.unchanged} already in the list, unchanged."]
    if res.problems:
        lines.append(f"{len(res.problems)} row(s) not imported:")
        lines += [f"  row {p.row}: {p.reason}" for p in res.problems[:15]]
        if len(res.problems) > 15:
            lines.append(f"  ... and {len(res.problems) - 15} more")
    return "\n".join(lines)
