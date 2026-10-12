"""core/leakcheck.py - find and blank passwords that versions before 10.1 may have written
into log files, history.json and the summary workbooks (audit finding H6).

Clients > "Check old files for passwords".  Reports only file names and client names,
never the password itself.  Blanking replaces the password with *** in place.
"""
import json
from pathlib import Path

TEXT_TYPES = {".txt", ".json", ".csv", ".log", ".html"}


def _candidates(log_dir: Path, history_file: Path, output_root: Path):
    seen = set()
    roots = [Path(log_dir)] + ([Path(output_root)] if output_root else [])
    files = [Path(history_file)] if Path(history_file).exists() else []
    for root in roots:
        if root.exists():
            files += [p for p in root.rglob("*") if p.is_file()
                      and (p.suffix.lower() in TEXT_TYPES or
                           (p.suffix.lower() == ".xlsx" and ("Summary" in p.name)))]
    for p in files:
        if p.resolve() not in seen:
            seen.add(p.resolve())
            yield p


def _xlsx_hits(path: Path, secrets):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    hits = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and any(s in c.value for s in secrets):
                    hits += 1
    return wb, hits


def scan(clients, log_dir, history_file, output_root, redact=False):
    """[(file, [client names])] for every file containing a saved password.
    With redact=True each password is replaced by *** in those files."""
    secrets = {c["password"]: c.get("name") or c.get("login", "")
               for c in clients if c.get("password") and len(c["password"]) >= 4}
    found = []
    for path in _candidates(log_dir, history_file, output_root):
        try:
            if path.suffix.lower() == ".xlsx":
                wb, hits = _xlsx_hits(path, secrets)
                if not hits:
                    continue
                names = sorted({secrets[s] for s in secrets
                                for ws in wb.worksheets for row in ws.iter_rows() for c in row
                                if isinstance(c.value, str) and s in c.value})
                if redact:
                    for ws in wb.worksheets:
                        for row in ws.iter_rows():
                            for c in row:
                                if isinstance(c.value, str):
                                    for s in secrets:
                                        if s in c.value:
                                            c.value = c.value.replace(s, "***")
                    wb.save(path)
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
                names = sorted({n for s, n in secrets.items() if s in text})
                if not names:
                    continue
                if redact:
                    for s in secrets:
                        text = text.replace(s, "***")
                    if path.suffix.lower() == ".json":
                        json.loads(text)          # still valid JSON before writing it back
                    path.write_text(text, encoding="utf-8")
            found.append((str(path), names))
        except (OSError, ValueError, KeyError) as e:
            found.append((str(path), [f"(could not read: {type(e).__name__})"]))
    return found
