"""Finding H6 follow-up: old files written before 10.1 may hold passwords - find and blank them."""
import json

from openpyxl import Workbook, load_workbook

from core import leakcheck

CLIENTS = [{"name": "Client A", "login": "1000000001", "password": "Secret-A-123"},
           {"name": "Client B", "login": "1000000002", "password": "Secret-B-456"}]


def _setup(tmp_path):
    logs, out = tmp_path / "logs", tmp_path / "out"
    logs.mkdir()
    out.mkdir()
    (logs / "run_old.txt").write_text('Call log: fill("Secret-A-123")', encoding="utf-8")
    (logs / "run_clean.txt").write_text("nothing here", encoding="utf-8")
    hist = tmp_path / "history.json"
    hist.write_text(json.dumps([{"note": 'fill("Secret-B-456")'}]), encoding="utf-8")
    wb = Workbook()
    wb.active.append(["Reason", 'fill("Secret-A-123")'])
    wb.save(out / "Run_Summary_20261001.xlsx")
    return logs, hist, out


def test_finds_every_file_with_a_password_and_names_the_client_not_the_password(tmp_path):
    logs, hist, out = _setup(tmp_path)
    found = dict(leakcheck.scan(CLIENTS, logs, hist, out))
    assert set(found) == {str(logs / "run_old.txt"), str(hist), str(out / "Run_Summary_20261001.xlsx")}
    assert found[str(hist)] == ["Client B"]
    assert "Secret" not in repr(found)


def test_blanking_removes_the_passwords_and_keeps_files_readable(tmp_path):
    logs, hist, out = _setup(tmp_path)
    leakcheck.scan(CLIENTS, logs, hist, out, redact=True)
    assert "Secret" not in (logs / "run_old.txt").read_text(encoding="utf-8")
    assert json.loads(hist.read_text(encoding="utf-8"))[0]["note"] == 'fill("***")'
    cell = load_workbook(out / "Run_Summary_20261001.xlsx").active["B1"].value
    assert cell == 'fill("***")'
    assert leakcheck.scan(CLIENTS, logs, hist, out) == []
