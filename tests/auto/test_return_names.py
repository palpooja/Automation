"""10.13 (user, 8 Oct 2026, from the ura-returns-tax-statement-download skill): return files named for the
audit file - '<n>. <Month> <YYYY>.pdf' (n = month's place in the period), amended returns labelled, a branch
named, yearly returns 'Original Prov 2025-26' - and a 'No Records Found' screenshot for an empty return type."""
import json
import os

import pytest
from openpyxl import load_workbook

from engines import returns_common as rc

H = ["Sr No.", "Reference No.", "Return Period From", "Return Period To", "Return Applied Date",
     "Applied For Branch", "Application Type", "Status"]


def row(ref, a, b, filed, kind="Original", branch="MAIN BRANCH"):
    return ["1", ref, a, b, filed, branch, kind, "Submitted"]


def test_names_for_monthly_amended_branch_and_yearly_returns():
    fl = rc.friendly_label
    assert fl(row("CR0127007220631", "01/07/2025", "31/07/2025", "13/08/2025"), H, "Monthly VAT Return",
              "01/07/2025") == "1. July 2025"
    assert fl(row("CR0127007220632", "01/06/2026", "30/06/2026", "09/07/2026"), H, "PAYE Return",
              "01/07/2025") == "12. June 2026"
    assert fl(row("CR0126025885425", "01/02/2026", "28/02/2026", "10/06/2026", "Amended"), H,
              "Monthly VAT Return", "01/07/2025") == "8. February 2026 - Amended filed 10-06-2026"
    assert fl(row("CR0126000000009", "01/07/2025", "31/07/2025", "14/08/2025", branch="JINJA BRANCH"), H,
              "PAYE Return", "01/07/2025") == "1. July 2025 - JINJA BRANCH"
    assert fl(row("CR0126012456703", "01/07/2025", "30/06/2026", "23/12/2025"), H,
              "Provisional Return for Non-Individual", "01/07/2025") == "Original Prov 2025-26"
    assert fl(row("CR0126012456704", "01/07/2025", "30/06/2026", "20/06/2026", "Amended"), H,
              "Provisional Return for Non-Individual", "01/07/2025") == "Amended Prov 2025-26 filed 20-06-2026"
    # a period that is not July-June: month numbers follow the period asked for (Aone, 7 Oct: JUN-2025 first)
    assert fl(row("CR0126000452164", "01/06/2025", "30/06/2025", "07/07/2025"), H, "Monthly VAT Return",
              "01/06/2025") == "1. June 2025"
    assert fl(["R-VAT-01", "06/2026", "Submitted"], None, "Monthly VAT Return", "01/07/2025") == "", \
        "no period in the row: the old name is used"


def test_two_returns_never_share_a_name():
    mapping = {"3. September 2025": "CR0000000000001"}
    assert rc.unique_label("3. September 2025", "CR0000000000001", mapping) == "3. September 2025"
    assert rc.unique_label("3. September 2025", "CR0000000000002", mapping, "02/10/2025") == \
        "3. September 2025 (filed 02-10-2025)"


@pytest.mark.browser
def test_live_layout_names_resume_and_no_records_screenshot(page, mock_url, tmp_path, monkeypatch):
    from engines import returns_pw
    monkeypatch.setattr(rc, "TAB_WAIT_SECONDS", 5)
    monkeypatch.setattr(rc, "PAGE_WAIT_SECONDS", 5)

    def run():
        page.goto(f"{mock_url}/home.html")
        res, logs = {}, []
        returns_pw.run_on_page(page, "1000012345", "01/07/2025", "30/06/2026",
                               ["Withholding Tax Return", "PAYE Return"], logs.append, lambda: None,
                               str(tmp_path), patient=False, result=res, skip_existing=True)
        return res, logs

    res, logs = run()
    wht = tmp_path / "1000012345" / "Withholding Tax Return"
    assert sorted(p.name for p in wht.glob("*.pdf")) == [
        "1. July 2025.pdf", "2. August 2025 - Amended filed 20-09-2025.pdf",
        "2. August 2025 - JINJA BRANCH.pdf", "2. August 2025.pdf"]
    assert json.loads((wht / rc.SAVED_AS_FILE).read_text())["2. August 2025"] == "CR0127000000002"
    shot = tmp_path / "1000012345" / "PAYE Return" / "No Records Found 2025-26.png"
    assert shot.exists() and shot.stat().st_size > 1000
    assert any(i["status"] == "no records" and i.get("file") == str(shot) for i in res["items"])
    ws = load_workbook(tmp_path / "1000012345" / "Returns_Status.xlsx")["Returns"]
    head = [c.value for c in ws[1]]
    saved = {r[head.index("Saved as")] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "2. August 2025 - Amended filed 20-09-2025.pdf" in saved
    # second run: nothing downloaded again
    res2, _ = run()
    assert not [i for i in res2["items"] if i["status"] == "saved" and i["file"].endswith(".pdf")]
    assert sum(1 for i in res2["items"] if i["status"] == "skipped") == 4


@pytest.mark.browser
def test_returns_saved_by_10_12_are_renamed_not_downloaded_again(page, mock_url, tmp_path, monkeypatch):
    from engines import returns_pw
    monkeypatch.setattr(rc, "TAB_WAIT_SECONDS", 5)
    monkeypatch.setattr(rc, "PAGE_WAIT_SECONDS", 5)
    wht = tmp_path / "1000012345" / "Withholding Tax Return"
    wht.mkdir(parents=True)
    old = wht / "CR0127000000001 01-07-2025 to 31-07-2025 filed 14-08-2025.pdf"
    old.write_bytes(b"%PDF-1.4 " + b"x" * 5000)
    page.goto(f"{mock_url}/home.html")
    res = {}
    returns_pw.run_on_page(page, "1000012345", "01/07/2025", "30/06/2026", ["Withholding Tax Return"],
                           lambda *_: None, lambda: None, str(tmp_path), patient=False, result=res,
                           skip_existing=True)
    assert not old.exists() and (wht / "1. July 2025.pdf").read_bytes().startswith(b"%PDF-1.4 xxx")
    assert len([i for i in res["items"] if i["status"] == "saved" and i["file"].endswith(".pdf")]) == 3
