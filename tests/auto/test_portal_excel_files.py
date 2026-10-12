"""Live run 3 Oct 2026: every ledger file saved from the portal's 'Export To Excel' button
gave 'file format or extension is not valid' in Excel.  The portal sends an HTML table named
Level-2.xls; the program saved it as .xlsx.  Files must open in Excel, hold exactly the text the
portal sent, and files already saved that way must be repaired (data is inside them)."""
import os
import time

import pytest
from lxml import html as LH
from openpyxl import load_workbook

from engines import excel_files
from engines import ledger_engine as eng

PORTAL_TABLE = """<table id="tblData"><tbody>
<tr><th rowspan="2">S/N</th><th rowspan="2">Period</th><th colspan="2">Payment</th><th rowspan="2">Net Position</th></tr>
<tr><th>PRN</th><th>WHT Credit</th></tr>
<tr><td>1</td><td>AUG-2026</td><td>2260012345678</td><td>1,329,127</td><td>(46,475.50)</td></tr>
<tr><td>2</td><td>JUL-2026</td><td>0NAM27010142462</td><td>0</td><td>007</td></tr>
</tbody></table>"""


def _texts(path):
    ws = load_workbook(path).active
    return [[("" if c.value is None else c.value) for c in row] for row in ws.iter_rows()]


def test_html_table_becomes_a_real_workbook_with_the_same_text(tmp_path):
    dest = tmp_path / "VAT_L2.xlsx"
    excel_files.html_to_xlsx(PORTAL_TABLE, dest)
    assert dest.read_bytes()[:2] == b"PK"                       # a real .xlsx (zip) file
    rows = _texts(dest)
    assert rows[0][:2] == ["S/N", "Period"] and rows[0][2] == "Payment" and rows[1][2:4] == ["PRN", "WHT Credit"]
    assert rows[2] == [1, "AUG-2026", "2260012345678", 1329127, -46475.5]
    assert rows[3] == [2, "JUL-2026", "0NAM27010142462", 0, "007"]
    ws = load_workbook(dest).active
    assert "C1:D1" in {str(r) for r in ws.merged_cells.ranges}  # 'Payment' spans two columns, as on the portal


@pytest.mark.parametrize("text, value", [
    ("1,329,127", 1329127), ("46475", 46475), ("12.50", 12.5), ("(500)", -500), ("-1,000.25", -1000.25),
    ("0", 0), ("007", "007"), ("1000626970", "1000626970"), ("2260012345678", "2260012345678"),
    ("0NAM27010142462", "0NAM27010142462"), ("AUG-2026", "AUG-2026"), ("01/08/2026", "01/08/2026"), ("", None),
])
def test_only_amounts_become_numbers_identifiers_stay_text(text, value):
    assert excel_files.cell_value(text) == value


@pytest.mark.browser
def test_portal_export_is_saved_as_a_real_xlsx_and_the_original_is_kept(page, mock_url, tmp_path):
    page.goto(f"{mock_url}/Portal/ledger_bits.html")
    f = page.main_frame
    assert eng.export_table_to_excel(page, f.locator("body"), tmp_path, "A_L2.xlsx",
                                     table_locator=f.locator("#tblData"), wait_s=10) == "portal export"
    saved = tmp_path / "A_L2.xlsx"
    assert saved.read_bytes()[:2] == b"PK"
    original = tmp_path / "portal_original" / "A_L2.xls"
    assert original.read_bytes().lstrip().startswith(b"<table")
    portal_rows = [[" ".join(c.text_content().split()) for c in tr.xpath("./th|./td")]
                   for tr in LH.fromstring(original.read_text(encoding="utf-8")).xpath("//tr")]
    # amounts are stored as numbers (46,475 -> 46475, shown as 46,475 by the cell format); all else as text
    expected = [[("" if excel_files.cell_value(t) is None else excel_files.cell_value(t)) for t in row]
                for row in portal_rows]
    assert _texts(saved) == expected
    ws = load_workbook(saved).active
    shown_with_commas = [c.number_format for row in ws.iter_rows() for c in row
                         if isinstance(c.value, int) and c.value >= 1000]
    assert shown_with_commas and all(f == "#,##0" for f in shown_with_commas)


def test_files_saved_by_earlier_versions_are_repaired_and_keep_their_date(tmp_path):
    broken = tmp_path / "1000626970" / "PAYE_L3_AUG-2026.xlsx"
    broken.parent.mkdir()
    broken.write_text(PORTAL_TABLE, encoding="utf-8")
    old = time.time() - 3 * 86400
    os.utime(broken, (old, old))
    good = tmp_path / "1000626970" / "Real.xlsx"
    excel_files.html_to_xlsx(PORTAL_TABLE, good)
    good_bytes = good.read_bytes()
    assert excel_files.repair_folder(tmp_path) == 1
    assert broken.read_bytes()[:2] == b"PK" and _texts(broken)[2][1] == "AUG-2026"
    assert abs(broken.stat().st_mtime - old) < 2, "date kept, so 'saved on an earlier day' still works"
    assert (broken.parent / "portal_original" / "PAYE_L3_AUG-2026.xls").exists()
    assert good.read_bytes() == good_bytes                       # a real workbook is never touched
    assert excel_files.repair_folder(tmp_path) == 0               # nothing left to repair


def test_a_download_without_any_table_is_refused_not_saved_as_done(tmp_path):
    with pytest.raises(ValueError):
        excel_files.html_to_xlsx("<html><body>Session expired</body></html>", tmp_path / "x.xlsx")
    assert not (tmp_path / "x.xlsx").exists()
