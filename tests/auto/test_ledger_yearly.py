"""10.15 - Income Tax and Rental (yearly ledgers) on the test copy of the screens the user captured on 8 Oct 2026.

Live 8 Oct 2026, 22:33 and 22:35: 'Income Tax failed (attempt 1/3): Locator.click: Timeout 30000ms exceeded' - three
times each for Income Tax and Rental, so neither got a Level 3.  Cause found in the captures: the yearly Level 2
has two heading rows, so the program's column count put 'Payment' over the Interest cell; it opened Level 3-5
(Interest) and waited for an Export button that page does not have.
Now: the openLevelThreeTwo links only (PRN and Tax Credit open the same page - one request per year), a year with
both amounts 0 is not opened, one failing year never stops the others, everything ends in <Tax head>.xlsx.
Every amount and reference in the test copy is invented (tests/mock_portal/ledger_server.py)."""
import pytest
from openpyxl import load_workbook

from engines import ledger_engine as eng, ledger_taxhead as th

pytestmark = pytest.mark.browser
THIS_YEAR, LAST_YEAR = "JUL-2026 to JUN-2027", "JUL-2025 to JUN-2026"


@pytest.fixture
def yearly(page, mock_url, ledger_portal, monkeypatch, tmp_path):
    ledger_portal.reset(kind="yearly", level2_delay=0.3, level3_delay=0.2)
    level2 = f"{mock_url}{ledger_portal.BASE}loadLevelTwo"

    def open_level2(pg, tax, tin=None, password=None):
        pg.goto(level2)
        return pg.main_frame
    # the real code goes Level 1 -> Level 2; the test copy has Level 2 only
    monkeypatch.setattr(eng, "open_level2", open_level2)
    monkeypatch.setattr(eng, "recover_to_level1", lambda *a, **k: None)
    monkeypatch.setattr(eng, "_back_to_level1_after", lambda *a, **k: None)
    monkeypatch.setattr(eng, "ANNUAL_L3", True)
    monkeypatch.setattr(eng, "ANNUAL_ALL_YEARS", False)
    monkeypatch.setattr(eng, "WORK_DIR", str(tmp_path / "logs"))
    eng.RESULTS.clear()
    eng.NOTES.clear()
    return ledger_portal


def _results(tax="Income Tax"):
    return {r["Item"]: (r["Status"], r["Detail"]) for r in eng.RESULTS if r["Section"] == f"{tax} ledger"}


def test_each_year_opened_once_by_its_payment_link_and_saved_in_the_workbook(page, yearly, tmp_path):
    eng.process_annual_tax_type(page, "Income Tax", tmp_path, period_from="2025-07", period_to="2027-06")
    assert yearly.STATE["other_requests"] == [], "never Level 3-1 / 3-5 (the 10.14 fault)"
    assert set(yearly.STATE["level3_requests"]) == {THIS_YEAR, LAST_YEAR}
    got = _results()
    assert got[THIS_YEAR] == ("saved", "Payment 350,000; Tax credit 460,000")
    assert got[LAST_YEAR] == ("saved", "Payment 7,500,000; Tax credit 1,200,000")
    assert "JUL-2024 to JUN-2025" not in got and "JUL-2023 to JUN-2024" not in got, "outside the period"
    eng.update_tax_head_files(tmp_path, ["Income Tax"])
    wb = load_workbook(tmp_path / "Income Tax.xlsx")
    assert wb.sheetnames == ["Summary", "Level 3-2", LAST_YEAR, THIS_YEAR]
    rows = list(wb["Summary"].iter_rows(values_only=True))
    head = list(rows[0])
    summary = {r[0]: dict(zip(head, r)) for r in rows[1:]}
    assert summary[THIS_YEAR]["Check"] == "OK" and summary[LAST_YEAR]["Check"] == "OK"
    assert summary[THIS_YEAR]["Credits by type"] == "WHT Credit 60,000; ASYCUDA Credit 400,000"
    assert not list(th.months_folder(tmp_path, "Income Tax").glob("*.xlsx")), "the yearly files are inside the workbook"
    # a second run finds the years in the workbook and opens nothing
    yearly.STATE["level3_requests"].clear()
    eng.RESULTS.clear()
    eng.process_annual_tax_type(page, "Income Tax", tmp_path, period_from="2025-07", period_to="2027-06")
    assert yearly.STATE["level3_requests"] == []
    assert _results()[THIS_YEAR][0] == "already downloaded"


def test_year_with_nothing_paid_is_not_opened(page, yearly, tmp_path, monkeypatch):
    monkeypatch.setattr(eng, "ANNUAL_ALL_YEARS", True)
    eng.process_annual_tax_type(page, "Income Tax", tmp_path, period_from="2025-07", period_to="2026-06")
    assert "JUL-2024 to JUN-2025" not in yearly.STATE["level3_requests"]
    assert _results()["JUL-2024 to JUN-2025"] == ("no payment link", "Payment 0; Tax credit 0")
    assert _results()["JUL-2023 to JUN-2024"][0] == "saved", "all years: the old year with a payment too"


def test_one_failing_year_does_not_stop_the_others(page, yearly, tmp_path):
    yearly.STATE["fail_years"] = {LAST_YEAR}
    eng.process_annual_tax_type(page, "Income Tax", tmp_path, period_from="2025-07", period_to="2027-06")
    got = _results()
    assert got[LAST_YEAR][0] == "failed"
    assert got[THIS_YEAR][0] == "saved", "the next year still done"
    copies = list((tmp_path / "logs").glob("*error_Income_Tax_JUL-2025*"))
    assert any(p.suffix == ".png" for p in copies) and any(p.suffix == ".html" for p in copies), copies
    # the failed year is fetched by the next run (or the second try at the end of the client)
    yearly.STATE["fail_years"] = set()
    yearly.STATE["level3_requests"].clear()
    eng.RESULTS.clear()
    eng.process_annual_tax_type(page, "Income Tax", tmp_path, period_from="2025-07", period_to="2027-06")
    assert set(yearly.STATE["level3_requests"]) == {LAST_YEAR}
    assert _results()[LAST_YEAR][0] == "saved"


def test_year_in_period_means_the_year_overlaps_the_period():
    assert eng.year_in_period(LAST_YEAR, "2025-07", "2026-06")
    assert not eng.year_in_period("JUL-2024 to JUN-2025", "2025-07", "2026-06"), "10.14 took this one too"
    assert not eng.year_in_period(THIS_YEAR, "2025-07", "2026-06")
    assert eng.year_in_period(THIS_YEAR, "2025-05", "2026-10"), "the default 18 months reach into it"
