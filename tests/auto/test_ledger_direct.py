"""10.16 - the direct way for Level 3 (user, 9 Oct 2026: 'build direct but also keep current way').

The payment link only asks URA for one page (openLevelThreeTwo -> POST loadLevelThreeTwo?period=...).  The direct
way sends that request from the Level 2 page itself and saves the ledger table of URA's answer - nothing drawn,
no Back.  Live check by the user, 9 Oct 2026: URA answered it (status 200, the full Level 3-2 page); MAR-2026 equal
to the portal's Export cell for cell; FEB-2026 the same 4 lines with two lines of one payment in the other order.
The current (screen) way stays: it proves the first month of each tax type, and takes over on any clear failure.
Test portal: tests/mock_portal/ledger_server.py (invented values); 'cors' = a direct request, 'navigate' = a page."""
import json

import pytest

from engines import excel_files, ledger_engine as eng

pytestmark = pytest.mark.browser
MONTHS = [("JUN-2026", "1,329,127"), ("MAY-2026", "8,068,063"), ("APR-2026", "9,704,197")]


@pytest.fixture
def ledger(page, mock_url, ledger_portal, monkeypatch):
    level2 = f"{mock_url}{ledger_portal.BASE}loadLevelTwo"
    monkeypatch.setattr(eng, "refresh_level2", lambda pg, fr, tax: (pg.goto(level2), pg.main_frame)[1])
    monkeypatch.setattr(eng, "L3_ROUTE", "auto")

    def start(**state):
        ledger_portal.reset(**state)
        page.goto(level2)
        for d in (eng.TIMINGS, eng.ROUTE_STATE, eng.DIRECT_STATE):
            d.clear()
        eng.NOTES.clear()
        return page.main_frame
    return start


def _months(page, frame, folder):
    out = []
    for m, amt in MONTHS:
        frame, outcome = eng.level3_month(page, frame, folder, "PAYE", m, amt)
        out.append(outcome)
    return out


def _check(folder):
    return json.loads((folder / eng.ROUTE_CHECK_FILE).read_text(encoding="utf-8"))


def test_first_month_proved_then_the_others_asked_for_directly(page, ledger, ledger_portal, tmp_path, capsys):
    frame = ledger()
    loads = ledger_portal.STATE["level2_loads"]
    assert _months(page, frame, tmp_path / "auto") == ["saved"] * 3
    modes = ledger_portal.STATE["level3_modes"]
    assert modes.count("cors") == 3, f"the proof copy + 2 months directly: {modes}"
    assert ledger_portal.STATE["level2_loads"] == loads + 1, "only the first month's Back rebuilt Level 2"
    assert _check(tmp_path / "auto")["PAYE - direct"]["result"] == "identical"
    out = capsys.readouterr().out
    assert "direct way identical" in out
    assert out.count("asking URA for the Level 3-2 directly") == 2, "MAY and APR"
    # the same files as the current way alone
    eng.L3_ROUTE = "screen"
    _months(page, ledger(), tmp_path / "screen")
    for m, _ in MONTHS:
        name = f"PAYE_L3_{m}.xlsx"
        assert eng.compare_level3(tmp_path / "auto" / name, tmp_path / "screen" / name)[0] == "identical"
        assert (tmp_path / "auto" / "portal_original" / f"PAYE_L3_{m}.xls").exists(), "URA's table kept as sent"


def test_current_way_only_never_sends_a_direct_request(page, ledger, ledger_portal, tmp_path, monkeypatch):
    monkeypatch.setattr(eng, "L3_ROUTE", "screen")
    assert _months(page, ledger(), tmp_path) == ["saved"] * 3
    assert "cors" not in ledger_portal.STATE["level3_modes"]


def test_portal_refusing_direct_requests_keeps_the_current_way(page, ledger, ledger_portal, tmp_path):
    assert _months(page, ledger(refuse_fetch=True), tmp_path) == ["saved"] * 3
    assert _check(tmp_path)["PAYE - direct"]["result"] == "not possible"
    assert ledger_portal.STATE["level3_modes"].count("cors") == 1, "tried once (the proof), then never again"


def test_direct_failure_switches_to_the_current_way_at_once(page, ledger, ledger_portal, tmp_path, monkeypatch,
                                                           capsys):
    monkeypatch.setattr(eng, "L3_ROUTE", "direct")                 # no proof: direct from the first month
    assert _months(page, ledger(refuse_fetch=True), tmp_path) == ["saved"] * 3
    out = capsys.readouterr().out
    assert "the direct way did not work (URA answered 403) - the screen way for the rest of PAYE" in out
    assert _check(tmp_path)["PAYE - direct"]["result"] == "switched off"


def test_a_slow_answer_is_waited_for_and_shown_in_the_log(page, ledger, ledger_portal, tmp_path, monkeypatch,
                                                          capsys):
    monkeypatch.setattr(eng, "L3_ROUTE", "direct")
    monkeypatch.setattr(eng, "DIRECT_HEARTBEAT_S", 2)
    frame = ledger(level3_delay=4)
    frame, outcome = eng.level3_month(page, frame, tmp_path, "PAYE", "JUN-2026", "1,329,127")
    assert outcome == "saved"
    out = capsys.readouterr().out
    assert "JUN-2026: URA still working on it (" in out
    assert ledger_portal.STATE["level3_modes"] == ["cors"], "one request, no retry, no click"


def test_optional_switch_time_hands_a_slow_month_to_the_current_way(page, ledger, ledger_portal, tmp_path,
                                                                    monkeypatch, capsys):
    monkeypatch.setattr(eng, "L3_ROUTE", "direct")
    monkeypatch.setattr(eng, "DIRECT_SWITCH_AFTER_S", 2)
    frame = ledger(level3_delay=4)
    frame, outcome = eng.level3_month(page, frame, tmp_path, "PAYE", "JUN-2026", "1,329,127")
    assert outcome == "saved"
    assert "no answer within 2 s (switch time set on the Ledgers page)" in capsys.readouterr().out


def test_same_lines_in_another_order_pass_the_proof(page, ledger, ledger_portal, tmp_path):
    """Live VAT FEB-2026: a payment split over two lines with the same date came in either order."""
    assert _months(page, ledger(flip_order=True), tmp_path) == ["saved"] * 3
    check = _check(tmp_path)
    # click (1st request), background proof (2nd, other order), direct proof (3rd, same order as the 1st)
    assert check["PAYE"]["result"] == eng.SAME_LINES, "the background way's proof passes too"
    assert check["PAYE - direct"]["result"] in ("identical", eng.SAME_LINES)
    assert ledger_portal.STATE["level3_modes"].count("cors") == 3, "the other months still asked for directly"


def test_compare_level3_tells_order_from_a_real_difference(tmp_path):
    def l3(name, rows, stamp):
        body = "".join("<tr>" + "".join(f"<td>{v}</td>" for v in r) + "</tr>" for r in rows)
        excel_files.html_to_xlsx(f'<table><tr><td colspan="4">Tax Position as at {stamp}</td></tr>'
                                 '<tr><th colspan="4">Period: FEB-2026</th></tr><tr><th>Sr No.</th>'
                                 f'<th>Payment Basis</th><th>Liability Basis</th><th>Principal</th></tr>{body}</table>',
                                 tmp_path / name)
    a = [(1, "Payment (PRN)", "Original Return", "1,000"), (2, "Payment (PRN)", "Penalty Assessment", "0")]
    l3("a.xlsx", a, "08/10/2026 10:32:54")
    l3("b.xlsx", a, "09/10/2026 02:29:25")
    l3("c.xlsx", [(1, *a[1][1:]), (2, *a[0][1:])], "09/10/2026 02:29:25")
    l3("d.xlsx", [a[0], (2, "Payment (PRN)", "Penalty Assessment", "27")], "09/10/2026 02:29:25")
    assert eng.compare_level3(tmp_path / "a.xlsx", tmp_path / "b.xlsx") == ("identical", "")
    assert eng.compare_level3(tmp_path / "a.xlsx", tmp_path / "c.xlsx")[0] == eng.SAME_LINES
    assert eng.compare_level3(tmp_path / "a.xlsx", tmp_path / "d.xlsx")[0] == "different"


def test_yearly_ledger_uses_the_direct_way_too(page, mock_url, ledger_portal, tmp_path, monkeypatch):
    ledger_portal.reset(kind="yearly", level2_delay=0.3, level3_delay=0.2)
    level2 = f"{mock_url}{ledger_portal.BASE}loadLevelTwo"
    monkeypatch.setattr(eng, "open_level2", lambda pg, tax, tin=None, password=None: (pg.goto(level2),
                                                                                       pg.main_frame)[1])
    monkeypatch.setattr(eng, "recover_to_level1", lambda *a, **k: None)
    monkeypatch.setattr(eng, "_back_to_level1_after", lambda *a, **k: None)
    monkeypatch.setattr(eng, "ANNUAL_L3", True)
    monkeypatch.setattr(eng, "ANNUAL_ALL_YEARS", True)
    monkeypatch.setattr(eng, "L3_ROUTE", "auto")
    monkeypatch.setattr(eng, "WORK_DIR", str(tmp_path / "logs"))
    for d in (eng.RESULTS, eng.NOTES, eng.DIRECT_STATE, eng.ROUTE_STATE):
        d.clear()
    eng.process_annual_tax_type(page, "Income Tax", tmp_path, period_from="2025-07", period_to="2027-06")
    got = {r["Item"]: r["Status"] for r in eng.RESULTS if r["Section"] == "Income Tax ledger"}
    assert [got[y] for y in ("JUL-2026 to JUN-2027", "JUL-2025 to JUN-2026", "JUL-2023 to JUN-2024")] == ["saved"] * 3
    assert ledger_portal.STATE["level3_modes"].count("cors") == 3, "proof of the first year + 2 years directly"
    assert ledger_portal.STATE["other_requests"] == []
    eng.update_tax_head_files(tmp_path, ["Income Tax"])
    from openpyxl import load_workbook
    rows = list(load_workbook(tmp_path / "Income Tax.xlsx")["Summary"].iter_rows(values_only=True))
    head = list(rows[0])
    checks = {r[0]: r[head.index("Check")] for r in rows[1:] if r[head.index("Level 3 rows")]}
    assert set(checks.values()) == {"OK"}, checks
