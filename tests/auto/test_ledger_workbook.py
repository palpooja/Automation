"""Combined Ledgers workbook: one per client - Contents, then per tax type an L2 tab and an All L3 tab."""
import time

from openpyxl import load_workbook

from engines import excel_files, ledger_engine as eng, ledger_workbook as lw

L2 = """<table><tr><th rowspan="2">S/N</th><th rowspan="2">Period</th><th colspan="2">Payment</th></tr>
<tr><th>PRN</th><th>WHT Credit</th></tr>
<tr><td>1</td><td>AUG-2026</td><td>2260012345678</td><td>1,329,127</td></tr>
<tr><td>2</td><td>JUL-2026</td><td>2260012345679</td><td>0</td></tr></table>"""


def l3(rows, extra_col=False):
    head = "<tr><th>Sr No.</th><th>Payment Basis</th><th>PRN</th><th>Amount</th>" + \
           ("<th>Refund</th>" if extra_col else "") + "</tr>"
    body = "".join(f"<tr><td>{i}</td><td>{b}</td><td>{p}</td><td>{a}</td>" + ("<td>0</td>" if extra_col else "")
                   + "</tr>" for i, (b, p, a) in enumerate(rows, 1))
    return f"<table>{head}{body}</table>"


def make(folder):
    excel_files.html_to_xlsx(L2, folder / "Value Added Tax_L2.xlsx")
    excel_files.html_to_xlsx(l3([("Return", "2260011111111", "1,000,000"), ("=1+1", "2260022222222", "329,127")]),
                             folder / "Value Added Tax_L3_AUG-2026.xlsx")
    excel_files.html_to_xlsx(l3([("Return", "2260033333333", "50,000")], extra_col=True),
                             folder / "Value Added Tax_L3_SEP-2026.xlsx")
    excel_files.html_to_xlsx(L2, folder / "PAYE_L2.xlsx")
    excel_files.html_to_xlsx(L2, folder / "Income Tax_L2.xlsx")
    excel_files.html_to_xlsx(l3([("Assessment", "2260044444444", "10")]), folder / "Income Tax_L3_Payment_2025_10.xlsx")
    excel_files.html_to_xlsx(l3([("Assessment", "2260055555555", "20")]), folder / "Income Tax_L3_Payment_2026_20.xlsx")


def values(ws):
    """Cell values per row, without the empty cells openpyxl pads up to the widest row."""
    out = []
    for row in ws.iter_rows():
        v = [c.value for c in row]
        while v and v[-1] is None:
            v.pop()
        out.append(v)
    return out


def test_combined_workbook_layout_contents_and_exact_values(tmp_path):
    make(tmp_path)
    results = [{"Section": "Value Added Tax ledger", "Item": "JUL-2026", "Status": "nil payment - skipped"},
               {"Section": "PAYE ledger", "Item": "AUG-2026", "Status": "failed"}]
    progress = {"Value Added Tax": {"OCT-2026": "not in ledger @2026-10-04"}}
    path = lw.build(tmp_path, "1000626970", "2026-07", "2026-10", ["Value Added Tax", "PAYE", "Income Tax"],
                    ["JUL-2026", "AUG-2026", "SEP-2026", "OCT-2026"], ["Income Tax"], results, progress)
    assert path.name == "Ledgers_1000626970_2026-07_to_2026-10.xlsx" and path.read_bytes()[:2] == b"PK"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Contents", "Value Added Tax - L2", "Value Added Tax - All L3",
                             "PAYE - L2", "Income Tax - L2", "Income Tax - All L3"]
    lines = [(r + [None] * 5)[:4] for r in values(wb["Contents"])[4:]]
    assert lines == [
        ["Value Added Tax", "Level 2", "included", None],
        ["Value Added Tax", "JUL-2026", "nil payment - skipped", None],
        ["Value Added Tax", "AUG-2026", "included", 2],
        ["Value Added Tax", "SEP-2026", "included", 1],
        ["Value Added Tax", "OCT-2026", "not in ledger", None],
        ["PAYE", "Level 2", "included", None],
        ["PAYE", "JUL-2026", "not downloaded", None],
        ["PAYE", "AUG-2026", "failed", None],
        ["PAYE", "SEP-2026", "not downloaded", None],
        ["PAYE", "OCT-2026", "not downloaded", None],
        ["Income Tax", "Level 2", "included", None],
        ["Income Tax", "Payment_2025_10", "included", 1],
        ["Income Tax", "Payment_2026_20", "included", 1]]
    # Level 2 copied exactly, merged 'Payment' header kept
    src = load_workbook(tmp_path / "Value Added Tax_L2.xlsx").active
    assert values(wb["Value Added Tax - L2"]) == values(src)
    assert "C1:D1" in {str(m) for m in wb["Value Added Tax - L2"].merged_cells.ranges}
    # All Level 3: Month + Source file, AUG rows, then SEP under its own (different) header
    l3v = values(wb["Value Added Tax - All L3"])
    assert l3v[0] == ["Month", "Source file", "Sr No.", "Payment Basis", "PRN", "Amount"]
    assert l3v[1] == ["AUG-2026", "Value Added Tax_L3_AUG-2026.xlsx", 1, "Return", "2260011111111", 1000000]
    assert l3v[2][:4] == ["AUG-2026", "Value Added Tax_L3_AUG-2026.xlsx", 2, "=1+1"]
    assert wb["Value Added Tax - All L3"].cell(row=3, column=4).data_type == "s"      # text, not a formula
    assert l3v[3] == []                                                              # blank line
    assert l3v[4] == ["Month", "Source file", "Sr No.", "Payment Basis", "PRN", "Amount", "Refund"]
    assert l3v[5] == ["SEP-2026", "Value Added Tax_L3_SEP-2026.xlsx", 1, "Return", "2260033333333", 50000, 0]
    assert wb["Value Added Tax - All L3"].cell(row=2, column=6).number_format == "#,##0"


def test_nothing_to_combine_writes_no_file(tmp_path):
    assert lw.build(tmp_path, "1000626970", "2026-07", "2026-07", ["PAYE"], ["JUL-2026"], []) is None
    assert not list(tmp_path.glob("Ledgers_*.xlsx"))


def test_ledger_step_records_the_combined_workbook_as_a_saved_file(tmp_path, monkeypatch):
    make(tmp_path)
    monkeypatch.setattr(eng, "RESULTS", [])
    eng.build_combined_workbook(tmp_path, "1000626970", "2026-07", "2026-10", ["Value Added Tax"])
    rec = eng.RESULTS[-1]
    assert rec["Status"] == "saved" and rec["File"].endswith("Ledgers_1000626970_2026-07_to_2026-10.xlsx")


def test_building_a_full_year_for_four_tax_types_takes_seconds_not_minutes(tmp_path):
    months = eng.generate_month_labels("2025-07", "2026-06")
    rows = [("Return", f"22600{i:08d}", f"{i * 1000:,}") for i in range(1, 21)]
    taxes = ["Value Added Tax", "IT-WHT", "PAYE", "Local Excise Duty"]
    for t in taxes:
        excel_files.html_to_xlsx(L2, tmp_path / f"{t}_L2.xlsx")
        for m in months:
            excel_files.html_to_xlsx(l3(rows), tmp_path / f"{t}_L3_{m}.xlsx")
    t0 = time.monotonic()
    path = lw.build(tmp_path, "1000000001", "2025-07", "2026-06", taxes, months, [])
    took = time.monotonic() - t0
    ws = load_workbook(path)["PAYE - All L3"]
    assert ws.max_row == 1 + 12 * 20
    print(f"\n4 tax types x 12 months x 20 rows: built in {took:.2f} s")
    assert took < 15


def test_a_ledger_job_ends_with_one_updated_workbook_per_tax_head(tmp_path, monkeypatch):
    """The real run_job_on_page; only the portal steps are replaced, writing what the live run wrote.
    Since 10.9 (user, 6 Oct 2026): one workbook per tax head, one table, updated in place - the
    period-named Ledgers_<TIN>_<from>_to_<to>.xlsx is no longer created."""
    from engines import ledger_taxhead as th
    for name in ("debug_dump", "dismiss_popup_if_present", "open_tax_statement", "recover_to_level1"):
        monkeypatch.setattr(eng, name, lambda *a, **k: None)

    def fake_tax_type(page, tax, period_from, period_to, save_folder, tin=None, password=None):
        save_folder.mkdir(parents=True, exist_ok=True)
        excel_files.html_to_xlsx(L2, save_folder / f"{tax}_L2.xlsx")
        eng.record(f"{tax} ledger", "Level 2", "saved", "portal export", save_folder / f"{tax}_L2.xlsx")
        months = th.months_folder(save_folder, tax)
        months.mkdir(parents=True, exist_ok=True)
        excel_files.html_to_xlsx(l3([("Return", "2260011111111", "1,329,127")]), months / f"{tax}_L3_AUG-2026.xlsx")
        eng.record(f"{tax} ledger", "AUG-2026", "saved", "Payment 1329127", months / f"{tax}_L3_AUG-2026.xlsx")
        eng.record(f"{tax} ledger", "JUL-2026", "nil payment - skipped")

    monkeypatch.setattr(eng, "process_tax_type", fake_tax_type)
    monkeypatch.setattr(eng, "OUTPUT_ROOT", str(tmp_path))
    monkeypatch.setattr(eng, "WORK_DIR", str(tmp_path))
    failed = eng.run_job_on_page(None, {"tin": "1000000001", "password": "pw", "period_from": "2026-07",
                                        "period_to": "2026-08", "tax_types": ["Value Added Tax", "PAYE"],
                                        "other_jobs": []})
    assert failed == []
    folder = tmp_path / "1000000001"
    assert not list(folder.glob("Ledgers_*.xlsx")), "no period-named combined workbook any more"
    for tax in ("Value Added Tax", "PAYE"):
        path = folder / f"{tax}.xlsx"
        assert any(r["Status"] == "saved" and r["File"] == str(path) for r in eng.RESULTS)
        wb = load_workbook(path)
        assert wb.sheetnames == ["Summary", "Level 3-2", "AUG-2026"] and len(wb["Level 3-2"].tables) == 1
        rows = list(wb["Level 3-2"].iter_rows(values_only=True))
        assert rows[1][0] == "AUG-2026" and "2260011111111" in rows[1]
        summary = {r[0]: r[1] for r in list(wb["Summary"].iter_rows(values_only=True))[1:]}
        assert summary == {"JUL-2026": "nil payment - skipped", "AUG-2026": "Level 3 saved"}
        assert not th.months_folder(folder, tax).exists(), "no Excel file per month (user, 8 Oct 2026)"
