"""PAYE certificates on a test copy of the LIVE page structure (user's capture of 8 Oct 2026: Recorder JSON +
network record).  The older test page (Portal/paye.html) drew results by script and showed month-wise details
in the same window; the live portal does neither:
* Search posts the form and the whole PAYE page is replaced;
* 'Get Month-Wise Details' opens a separate pop-up window (#divPostContent there, with Close);
* GenerateTCC takes at most 5 employees per press (var count = 5; a 6th tick is refused with an alert).
Every name, TIN and amount in the test copy is invented (tests/mock_portal/paye_server.py)."""
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engines import paye_engine as pe

pytestmark = pytest.mark.browser


def _run(page, mock_url, tmp_path, outer="paye_outer.html", period=("2025-07", "2026-06"), **kw):
    page.goto(f"{mock_url}/Portal/{outer}")
    page.wait_for_timeout(800)
    logs, recs = [], []
    files = pe.run(page, "1000000001", period[0], period[1], tmp_path, logs.append, lambda: None,
                   lambda *a: recs.append(a), debug_dir=tmp_path, **kw)
    return files, recs, logs


def test_slow_answer_with_a_busy_page_still_finds_the_employees(page, mock_url, paye_portal, tmp_path):
    """6 Oct 2026, live: 'No employees listed' 10 s after Search for a client filing PAYE monthly.
    Reproduced: while the portal takes a few seconds, ANY other frame that reloads (a chat or tracking widget)
    was taken as the portal's answer; 6 s later 'no employees'."""
    paye_portal.reset(search_delay=3)
    files, recs, logs = _run(page, mock_url, tmp_path, outer="paye_outer_busy.html", generate=False, monthwise=False)
    assert not [r for r in recs if r[1] == "no records"], logs
    assert len(files) == 1
    assert load_workbook(files[0])["Employees"].max_row >= 4


def test_no_employees_says_why_and_keeps_a_copy_of_the_page(page, mock_url, paye_portal, tmp_path):
    files, recs, logs = _run(page, mock_url, tmp_path, period=("2023-07", "2024-06"), generate=False, monthwise=False)
    assert files == []
    (sec, status, detail), = [r for r in recs if r[1] == "no records"]
    assert "no employees list" in detail.lower(), detail
    assert (tmp_path / "paye_no_employees_2023-2024.html").exists()
    assert (tmp_path / "paye_no_employees_2023-2024.png").exists()


def test_month_wise_details_are_read_from_the_pop_up_window(page, mock_url, paye_portal, tmp_path):
    t0 = time.monotonic()
    files, recs, logs = _run(page, mock_url, tmp_path, generate=False, monthwise=True)
    took = time.monotonic() - t0
    assert not [r for r in recs if r[1] == "failed"], recs
    rows = list(load_workbook(files[0])["Month-wise details"].iter_rows(values_only=True))
    tins = [r for r in rows if r and any(str(v) in ("1000000011", "1000000022", "1000000033") for v in r)]
    assert len(tins) == 3 * 12, "12 months for each of the 3 employees"
    assert paye_portal.STATE["monthwise"] == ["1000000011", "1000000022", "1000000033"]
    assert len(page.context.pages) == 1, "every pop-up closed"
    assert took < 25, f"{took:.0f} s (was 30 s per employee, then failed)"


def test_generate_works_in_batches_of_the_portals_limit(page, mock_url, paye_portal, tmp_path):
    files, recs, logs = _run(page, mock_url, tmp_path, period=("2024-07", "2025-06"), generate=True, monthwise=False)
    assert len(paye_portal.STATE["generated"]) == 7, "7 employees without a TCC No.: 5 + 2"
    assert not [r for r in recs if r[1] in ("incomplete", "failed")], recs
    assert any("still without a TCC No." in m and "0 still" in m for m in logs), logs


# ---------- 10.15: live 8 Oct 2026, 22:40 - 'no results within 60 s after Search' failed the whole Certificates step


def test_search_press_that_never_reaches_the_portal_is_pressed_again(page, mock_url, paye_portal, tmp_path,
                                                                     monkeypatch):
    monkeypatch.setattr(pe, "PRESS_AGAIN_S", 2)
    paye_portal.reset(drop_presses=1)
    files, recs, logs = _run(page, mock_url, tmp_path, generate=False, monthwise=False)
    assert len(files) == 1, logs
    assert any("pressing it again" in m for m in logs), logs
    assert len(paye_portal.STATE["searches"]) == 1, "one search reached the portal"


def test_unanswered_search_is_tried_again_then_only_that_year_fails(page, mock_url, paye_portal, tmp_path,
                                                                    monkeypatch):
    monkeypatch.setattr(pe, "WAIT_S", 4)
    monkeypatch.setattr(pe, "PRESS_AGAIN_S", 1.5)       # a press lost while the page reloads is pressed again
    paye_portal.reset(hang={"2025-2026": 2})
    failed = []
    page.goto(f"{mock_url}/Portal/paye_outer.html")
    page.wait_for_timeout(800)
    logs, recs, reopened = [], [], []
    files = pe.run(page, "1000000001", "2025-07", "2027-06", tmp_path, logs.append, lambda: None,
                   lambda *a: recs.append(a), debug_dir=tmp_path, generate=False, monthwise=False,
                   reopen=lambda: (reopened.append(1), page.goto(f"{mock_url}/Portal/paye_outer.html", wait_until="load"),
                                   page.wait_for_timeout(800)), failed=failed)
    assert failed == ["2025-2026"]
    (sec, status, detail), = [r for r in recs if r[1] == "failed"]
    assert sec == "PAYE certificates July 2025 to June 2026"
    assert "the portal received the search" in detail, detail
    assert reopened == [1], "tried once more from the portal home page"
    assert (tmp_path / "paye_search_timeout_2025-2026.html").exists()
    assert (tmp_path / "paye_search_timeout_2025-2026.png").exists()
    assert [Path(f).name for f in files] == ["PAYE_TCC_1000000001_2026-2027.xlsx"], "the next year still done"
    # the second try at the end of the client: only the failed year
    paye_portal.reset()
    files = pe.run(page, "1000000001", "2025-07", "2027-06", tmp_path, logs.append, lambda: None,
                   lambda *a: recs.append(a), debug_dir=tmp_path, generate=False, monthwise=False,
                   years=failed)
    assert [Path(f).name for f in files] == ["PAYE_TCC_1000000001_2025-2026.xlsx"]


def test_one_unanswered_search_is_saved_by_the_second_search(page, mock_url, paye_portal, tmp_path, monkeypatch):
    monkeypatch.setattr(pe, "WAIT_S", 4)
    paye_portal.reset(hang={"2025-2026": 1})
    files, recs, logs = _run(page, mock_url, tmp_path, generate=False, monthwise=False)
    assert len(files) == 1 and not [r for r in recs if r[1] == "failed"], logs
    assert any("Trying once more" in m for m in logs)


# ---------- 10.17: live 9 Oct 2026 (Uganda connection) - the Search press never reached the portal, four times


def test_search_is_the_paye_forms_own_button_not_the_outer_pages_site_search(page, mock_url, paye_portal, tmp_path):
    """The outer ura.go.ug page has a 'Search' of its own.  The old selector list tried a bare 'Search' in every
    frame, the outer frame first: that press sent nothing to the portal (live 9 Oct 2026, 4 min 34 s lost)."""
    files, recs, logs = _run(page, mock_url, tmp_path, outer="paye_outer_sitesearch.html", generate=False,
                             monthwise=False)
    assert len(files) == 1, logs
    assert len(paye_portal.STATE["searches"]) == 1, "the PAYE search reached the portal at the first press"
    assert not any("pressing it again" in m for m in logs), logs
    assert page.main_frame.locator("#site-search-note").inner_text() == "", "the site's search was never pressed"


def test_a_press_that_never_sends_is_given_up_quickly_described_and_not_tried_a_third_time(page, mock_url, paye_portal,
                                                                                         tmp_path, monkeypatch):
    monkeypatch.setattr(pe, "PRESS_AGAIN_S", 1.5)
    monkeypatch.setattr(pe, "UNSENT_GIVE_UP_S", 1.5)
    paye_portal.reset(drop_presses=50)                 # every press on this page is swallowed
    failed, reopened = [], []
    page.goto(f"{mock_url}/Portal/paye_outer.html")
    page.wait_for_timeout(800)
    logs, recs = [], []

    def reopen():
        reopened.append(1)
        paye_portal.STATE["drop_presses"] = 50         # the fresh page swallows its presses too
        page.goto(f"{mock_url}/Portal/paye_outer.html", wait_until="load")
        page.wait_for_timeout(800)
    t0 = time.monotonic()
    files = pe.run(page, "1000000001", "2025-07", "2026-06", tmp_path, logs.append, lambda: None,
                   lambda *a: recs.append(a), debug_dir=tmp_path, generate=False, monthwise=False,
                   reopen=reopen, failed=failed)
    took = time.monotonic() - t0
    assert files == []
    (sec, status, detail), = [r for r in recs if r[1] == "failed"]
    assert "never reached the portal" in detail and "pressed twice" in detail, detail
    assert paye_portal.STATE["searches"] == [], "nothing was ever sent"
    assert reopened == [1], "one retry from the portal home page, no more"
    assert failed == [], "a press that never sends is not put off to the end of the client again"
    assert any("not tried again at the end of the client" in m for m in logs), logs
    assert took < 30, f"{took:.0f} s - the old code waited 60 s per press (2 x 68 s)"
    # the log says what the button is, and both attempts keep their own page copies
    assert any("button: <button" in m and "Period '2025-2026'" in m for m in logs), logs
    assert any("pressed again by" in m for m in logs), logs
    for name in ("paye_search_timeout_2025-2026.html", "paye_search_timeout_2025-2026_2.html",
                 "paye_search_timeout_2025-2026.png", "paye_search_timeout_2025-2026_2.png"):
        assert (tmp_path / name).exists(), name


def test_a_message_box_after_the_press_ends_the_wait_at_once_with_the_message(page, mock_url, paye_portal, tmp_path,
                                                                             monkeypatch):
    monkeypatch.setattr(pe, "WAIT_S", 30)
    paye_portal.reset(alert_presses=5)
    page.goto(f"{mock_url}/Portal/paye_outer.html")
    page.wait_for_timeout(800)
    logs, recs = [], []
    t0 = time.monotonic()
    pe.run(page, "1000000001", "2025-07", "2026-06", tmp_path, logs.append, lambda: None,
           lambda *a: recs.append(a), debug_dir=tmp_path, generate=False, monthwise=False)
    took = time.monotonic() - t0
    (sec, status, detail), = [r for r in recs if r[1] == "failed"]
    assert "Please select Period" in detail, detail
    assert took < 15, f"{took:.0f} s - a refused search is not waited for"
    assert any("answered the Search press with a message" in m for m in logs), logs


def test_a_page_that_comes_back_for_another_period_is_a_check_item_not_no_records(page, mock_url, paye_portal,
                                                                                 tmp_path):
    """Live 9 Oct 2026 (Tronics): 'No employees listed' for a year with 12.6M of PAYE paid - probably the chosen
    year was not applied.  A page showing another period than the one asked for needs a person."""
    paye_portal.reset(answer_period="2023-2024")       # 2023-2024 has no employees: the page shows it instead
    files, recs, logs = _run(page, mock_url, tmp_path, generate=False, monthwise=False)
    assert files == []
    (sec, status, detail), = [r for r in recs if r[1] == "check"]
    assert "2023-2024" in detail and "2025-2026" in detail, detail
    assert not [r for r in recs if r[1] == "no records"]
    assert any("Check this year by hand" in m for m in logs), logs
