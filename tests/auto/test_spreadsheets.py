"""Findings L2 (text that becomes a live formula) and L6 (amounts left out of totals)."""
import pandas as pd
import pytest
from openpyxl import load_workbook

from engines import wht_outputs as wo

EVIL = '=HYPERLINK("http://example.invalid","click")'


def test_run_summary_never_turns_portal_text_into_a_formula(tmp_path):
    import app
    rows = [{"when": "03/10/2026 10:00", "client": EVIL, "login": "1000000001",
             "module": "Returns", "status": "Completed", "files": 1, "seconds": 5}]
    items = [{"client": EVIL, "login": "1000000001", "module": "Returns", "item": "+SUM(1,2)",
              "status": "failed", "detail": "@cmd", "file": ""}]
    path = app.write_run_summary(rows, items, tmp_path / "s.xlsx", master=tmp_path / "m.xlsx")
    for wb in (load_workbook(path), load_workbook(tmp_path / "m.xlsx")):
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    assert c.data_type != "f", f"{ws.title}!{c.coordinate} became a formula"


def test_wht_workbook_keeps_portal_text_as_text_but_keeps_its_own_totals(tmp_path):
    r = wo.MonthResult(2026, 6)
    r.df = pd.DataFrame({"Sl. No.": [1], "Agent TIN": ["1000000001"], "Agent Name": [EVIL],
                         "Period": ["06/2026"], "Gross Amount": [1000.0], "Tax Withheld": [60.0]})
    r.status = "ok"
    path = wo.save_excel([r], str(tmp_path / "w.xlsx"), tin="1000000002", gen_label="Payee")
    ws = load_workbook(path)["2026-06"]
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    names = [c for c in cells if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert names and all(c.data_type == "s" for c in names)
    assert any(c.data_type == "f" and str(c.value).startswith("=SUM(") for c in cells), \
        "the program's own total formulas must stay"


@pytest.mark.parametrize("text, value", [
    ("1,000,000", 1000000.0), ("1,000,000.50", 1000000.5), ("-500,000", -500000.0),
    ("(500,000)", -500000.0), ("UGX 60,000", 60000.0), ("60,000/=", 60000.0),
    ("-", None), ("", None), ("abc", None),
])
def test_amounts_in_brackets_or_with_currency_marks_are_counted(text, value):
    assert wo.as_number(text) == value


def test_month_page_copies_merge_into_one_pdf_with_a_bookmark_each(tmp_path):
    """merge_pdfs uses pypdf; this guards the version upgrade (5.9 -> 6.19, finding M4)."""
    import io
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    def one_page(text):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(100, 750, text)
        c.save()
        return buf.getvalue()

    months = []
    for m in (6, 7):
        r = wo.MonthResult(2026, m)
        r.status, r.pdf_bytes = "ok", one_page(f"month {m}")
        months.append(r)
    path = wo.merge_pdfs(months, str(tmp_path / "pages.pdf"))
    reader = PdfReader(path)
    assert len(reader.pages) == 2
    assert [o.title for o in reader.outline] == ["June 2026", "July 2026"]
