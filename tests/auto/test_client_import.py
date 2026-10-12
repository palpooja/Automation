"""Bulk client upload from Excel (user, 6 Oct 2026).  Columns: Client name | Login ID (TIN) | Password.
Rules are the same as 'Add client' (core.store.client_problem); passwords never appear in messages."""
from openpyxl import Workbook, load_workbook

from core import client_import as ci

EXISTING = [{"name": "Client A", "login": "1000000001", "password": "old-pass"}]


def sheet(path, rows, header=("Client name", "Login ID (TIN)", "Password"), blank_rows=0):
    wb = Workbook()
    ws = wb.active
    for _ in range(blank_rows):
        ws.append([])
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def test_template_has_the_three_columns_and_reads_back(tmp_path):
    path = ci.make_template(tmp_path / "Clients_template.xlsx")
    ws = load_workbook(path).worksheets[0]
    assert [c.value for c in ws[1]][:3] == ["Client name", "Login ID (TIN)", "Password"]
    assert ws["B2"].number_format == "@", "the TIN column is text, so Excel keeps every digit"
    res = ci.read_clients(path, [])
    assert res.new == [] and res.problems == [], "the template's example row is not imported"


def test_new_updated_and_problem_rows_are_sorted_out(tmp_path):
    path = sheet(tmp_path / "c.xlsx", [
        ("Client B", 1000000002, "pw-b"),              # TIN typed as a number in Excel
        ("Client A renamed", "1000000001", "new-pass"),  # already in the list -> update
        ("Client C", "10000 00003", "pw-c"),            # space inside the TIN
        ("", "1000000004", "pw-d"),                     # no name
        ("Client B again", "1000000002", "pw-x"),       # twice in the file
        ("Client E", "1000000005", ""),                 # no password
        (None, None, None),                             # empty row: ignored
    ], blank_rows=2)
    res = ci.read_clients(path, EXISTING)
    assert [(c["name"], c["login"], c["password"]) for c in res.new] == [("Client B", "1000000002", "pw-b")]
    assert [(c["name"], c["login"]) for c in res.updates] == [("Client A renamed", "1000000001")]
    rows = {p.row: p.reason for p in res.problems}
    assert sorted(rows) == [6, 7, 8, 9]                    # Excel row numbers (2 blank + heading)
    assert "digits only" in rows[6] and "client name" in rows[7]
    assert "twice" in rows[8] and "password" in rows[9]
    text = " ".join(p.reason for p in res.problems) + ci.summary(res)
    for pw in ("pw-b", "new-pass", "pw-c", "pw-d", "pw-x"):
        assert pw not in text, "a password must never appear in a message"


def test_apply_adds_and_optionally_updates(tmp_path):
    res = ci.read_clients(sheet(tmp_path / "c.xlsx", [("Client B", "1000000002", "pw-b"),
                                                      ("Client A renamed", "1000000001", "new")]), EXISTING)
    kept = ci.apply(EXISTING, res, update_existing=False)
    assert [(c["name"], c["password"]) for c in kept] == [("Client A", "old-pass"), ("Client B", "pw-b")]
    both = ci.apply(EXISTING, res, update_existing=True)
    assert [(c["name"], c["password"]) for c in both] == [("Client A renamed", "new"), ("Client B", "pw-b")]


def test_columns_are_found_by_heading_in_any_order(tmp_path):
    path = sheet(tmp_path / "c.xlsx", [("pw-z", "Client Z", "1000000026")],
                 header=("PASSWORD", "Name of client", "TIN"))
    res = ci.read_clients(path, [])
    assert [(c["name"], c["login"], c["password"]) for c in res.new] == [("Client Z", "1000000026", "pw-z")]
