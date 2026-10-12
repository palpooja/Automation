"""End-to-end behaviour against the mock portal (tests/mock_portal).

These pin what the program does today, so that any speed change that alters a
result turns a test red.  They replace the print-only scripts now kept in
tests/scripts/ for reference."""
import os

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.browser


def _statuses(months):
    return [(r.label, r.status, r.rows) for r in months]


def test_wht_months_ok_empty_ok_and_workbook(page, mock_url, tmp_path, monkeypatch):
    from engines import wht_engine as we, wht_outputs as wo
    monkeypatch.setattr(we, "RESULT_WAIT_MS", 30000)
    page.goto(f"{mock_url}/home.html")
    bot = we.WhtBot(page, lambda *_: None, lambda: None, screenshot_dir=str(tmp_path))
    bot.dismiss_popups()
    bot.open_wht_page()
    months = [wo.MonthResult(y, m) for y, m in wo.month_sequence(2026, 6, 2026, 8)]
    we.run_months(bot, months, want_pdf=False, want_html=True, want_screenshot=False, generate_for="D")
    assert _statuses(months) == [("2026-06", "ok", 2), ("2026-07", "empty", 0), ("2026-08", "ok", 1)]
    path = wo.save_excel(months, str(tmp_path / "W.xlsx"), tin="1000012345")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Combined", "2026-06", "2026-08"]
    summary = [r for r in wb["Summary"].iter_rows(values_only=True) if r and r[0] in
               ("June 2026", "July 2026", "August 2026")]
    assert [(r[0], r[2], r[3]) for r in summary] == [
        ("June 2026", 2, "OK"), ("July 2026", 0, "No records"), ("August 2026", 1, "OK")]


def test_wht_live_page_payee_payer_and_unknown_choice(page, mock_url, monkeypatch):
    from engines import wht_engine as we, wht_outputs as wo
    monkeypatch.setattr(we, "RESULT_WAIT_MS", 30000)
    page.goto(f"{mock_url}/Portal/wht_live.html")
    bot = we.WhtBot(page, lambda *_: None, lambda: None)
    got = {}
    for code in ("D", "W", "X"):
        months = [wo.MonthResult(2026, m) for m in (1, 2, 3)]
        we.run_months(bot, months, want_pdf=False, want_html=False, generate_for=code)
        got[code] = _statuses(months)
    assert got["D"] == [("2026-01", "empty", 0), ("2026-02", "ok", 2), ("2026-03", "ok", 1)]
    assert got["W"] == [("2026-01", "ok", 3), ("2026-02", "empty", 0), ("2026-03", "ok", 1)]
    assert [s for _, s, _ in got["X"]] == ["error"] * 3


def test_wht_generate_tcc_fills_the_missing_numbers(page, mock_url, monkeypatch):
    from engines import wht_engine as we, wht_outputs as wo
    monkeypatch.setattr(we, "RESULT_WAIT_MS", 20000)
    page.goto(f"{mock_url}/Portal/wht_tcc.html")
    bot = we.WhtBot(page, lambda *_: None, lambda: None)
    months = [wo.MonthResult(2026, m) for m in (1, 2, 3)]
    we.run_months(bot, months, want_pdf=False, want_html=False, generate_for="D", generate_tcc=True)
    feb = next(r for r in months if r.label == "2026-02")
    assert feb.status == "ok"
    tcc = [str(v) for v in feb.df["TCC No"]]
    assert tcc and all(v.strip() for v in tcc), tcc


def test_returns_first_run_then_resume_skips_saved(page, mock_url, tmp_path, monkeypatch):
    from engines import returns_common, returns_pw
    monkeypatch.setattr(returns_common, "TAB_WAIT_SECONDS", 5)
    monkeypatch.setattr(returns_common, "PAGE_WAIT_SECONDS", 5)
    wanted = ["Monthly VAT Return", "PAYE Return", "Local Excise Duty Return"]

    def run():
        page.goto(f"{mock_url}/home.html")
        res = {}
        returns_pw.run_on_page(page, "1000012345", "01/07/2025", "03/10/2026", wanted,
                               lambda *_: None, lambda: None, str(tmp_path), patient=False,
                               result=res, skip_existing=True)
        return res

    def returns_saved(res):      # the status workbook is also recorded as 'saved'; count returns only
        return [i for i in res["items"] if i["status"] == "saved"
                and not i["file"].endswith("Returns_Status.xlsx")]

    first = run()
    saved = returns_saved(first)
    assert len(saved) == 3                                    # 2 VAT PDFs + 1 excise web page
    assert all(os.path.getsize(i["file"]) > 1000 for i in saved)
    second = run()
    assert sum(1 for i in second["items"] if i["status"] == "skipped") == 3
    assert returns_saved(second) == []


def test_paye_years_records_and_workbooks(page, mock_url, tmp_path):
    from engines import paye_engine as pe
    assert pe.financial_years("2025-07", "2026-06") == [pe.financial_years("2025-07", "2026-06")[0]]
    recs = []
    page.goto(f"{mock_url}/home.html")
    page.wait_for_timeout(800)
    files = pe.run(page, "1000626970", "2024-05", "2026-06", tmp_path, lambda *_: None,
                   lambda: None, lambda *a: recs.append(a), generate=False, debug_dir=tmp_path)
    kinds = [(r[0], r[1]) for r in recs if not r[0].endswith("TCC numbers")]
    assert kinds == [("PAYE certificates July 2023 to June 2024", "no records"),
                     ("PAYE certificates July 2024 to June 2025 - workbook", "saved"),
                     ("PAYE certificates July 2025 to June 2026 - workbook", "saved")]
    assert len(files) == 2
    for f in files:
        wb = load_workbook(f)
        assert wb.sheetnames == ["Employees", "Month-wise details"]
        assert wb["Employees"].max_row > 4 and wb["Month-wise details"].max_row > 4


def test_ledger_export_portal_file_and_screen_fallback(page, mock_url, tmp_path):
    import pandas as pd
    from engines import ledger_engine as eng
    page.goto(f"{mock_url}/Portal/ledger_bits.html")
    f = page.main_frame
    assert eng.export_table_to_excel(page, f.locator("body"), tmp_path, "A_L2.xlsx",
                                     table_locator=f.locator("#tblData"), wait_s=10) == "portal export"
    assert eng.export_table_to_excel(page, f.locator("#dead"), tmp_path, "B_L2.xlsx",
                                     table_locator=f.locator("#tblData"), wait_s=3) == "built from the screen"
    assert pd.read_excel(tmp_path / "B_L2.xlsx").values.tolist() == [[1, "JUN-2026", 100, 46475],
                                                                     [2, "MAY-2026", 90, 215136]]
    assert eng.tick_cert_rows(page, f, [{"n": 0}, {"n": 2}]) == 2
    assert f.evaluate("()=>[...document.querySelectorAll('[data-bot-cert]')].map(e=>e.checked)") == \
        [True, False, True]
    eng.safe_click(page, f.locator("#far"), "far link")
    assert f.evaluate("()=>window.farClicks") == 1
