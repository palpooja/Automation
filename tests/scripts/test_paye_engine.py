import sys, os, shutil, glob; sys.path.insert(0,'/home/claude/ura_suite')
from playwright.sync_api import sync_playwright
from engines import paye_engine as pe
print("FY for 2025-05..2026-06:", pe.financial_years("2025-05","2026-06"), "| 2025-07..2026-06:", pe.financial_years("2025-07","2026-06"), "| 2023-01..2025-12:", pe.financial_years("2023-01","2025-12"))
out="/tmp/out_paye"; shutil.rmtree(out, ignore_errors=True); recs=[]
with sync_playwright() as p:
    b=p.chromium.launch(executable_path="/tmp/chr/chromium", headless=True, args=["--no-sandbox","--disable-gpu","--disable-dev-shm-usage"])
    pg=b.new_context().new_page(); pg.goto("http://localhost:8765/home.html"); pg.wait_for_timeout(800)
    files=pe.run(pg, "1000626970", "2024-05", "2026-06", out, print, lambda: None, lambda *a: recs.append(a), debug_dir=out)
    b.close()
print("RECORDS:", [r[:2] for r in recs])
from openpyxl import load_workbook
for f in files:
    wb=load_workbook(f); print("==", os.path.basename(f), wb.sheetnames)
    for name in wb.sheetnames:
        rows=[r for r in wb[name].iter_rows(values_only=True) if any(v is not None for v in r)]
        print("  ", name, "-", len(rows), "rows; header:", [v for v in rows[3] if v] if len(rows)>3 else None)
        print("     first data:", [v for v in rows[4] if v is not None][:7] if len(rows)>4 else None)
