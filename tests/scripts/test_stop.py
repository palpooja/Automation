import sys, os, time, glob, subprocess, faulthandler; faulthandler.dump_traceback_later(90, repeat=True, file=open('/tmp/stacks_full.txt','w')); sys.path.insert(0,'/home/claude/ura_suite'); os.chdir('/home/claude/ura_suite')
from engines import ledger_engine as eng, returns_common
eng.LOGIN_URL="http://localhost:8765/login.html"
eng.launch_browser=lambda p: p.chromium.launch(executable_path="/tmp/chr/chromium", headless=True, args=["--no-sandbox","--disable-gpu","--disable-dev-shm-usage"])
returns_common.TAB_WAIT_SECONDS = returns_common.PAGE_WAIT_SECONDS = 5
import app
from core import store
a=app.App(); a.geometry("1240x820+0+0")
a.clients=[{"name":"ABC Traders Ltd","login":"1000012345","password":"pw1"},{"name":"Kampala Hardware Co","login":"1000067890","password":"pw2"}]
a.ticked_logins={c["login"] for c in a.clients}; a._refresh_clients()
[v.set(k!='ledger') for k,v in a.menu_vars.items()]
# Returns page: three types; WHT: Jun-Aug 2026; Ledgers: defaults
rm=app.MODULES[0]; [v.set(c in ("Monthly VAT Return","PAYE Return","Local Excise Duty Return")) for c,v in rm.cat_vars.items()]
rm.patient.set(True); wm=app.MODULES[2]; wm.from_var.set("2026-06"); wm.to_var.set("2026-08")
a.settings["suite"]={"open_folder_when_done":False,"skip_done_today":False,"nav_timeout_s":15}
answered=[]; shots=[]
def person():                     # answers the yellow bar like a human, 1 s later
    if a.ask.winfo_ismapped():
        txt=a.ask_msg.cget("text"); answered.append(txt.splitlines()[0])
        if len(answered)==1: shots.append(1); subprocess.run(["import","-display",":99","-window","root","/tmp/app_yellowbar.png"])
        a.after(800, lambda: a._answer(""))
    a.after(700, person)
a.start(all_ticked=True); a.after(1500, person)
T0=time.time()
a.after(25000, lambda: (print('PRESSING STOP at', round(time.time()-T0),'s', flush=True), a.stop()))
def watch():
    if a.worker and not a.worker.is_alive(): print('WORKER ENDED at', round(time.time()-T0),'s', flush=True); a.after(300,a.destroy); return
    a.after(500, watch)
a.after(26000, watch)
a.mainloop()
print("YELLOW BAR ASKED FOR:", answered)
log=open(sorted(store.LOG_DIR.glob("run_*.txt"))[-1]).read()
print("---- key log lines"); [print(l[:150]) for l in log.splitlines() if any(k in l for k in ("=====","-----","Logged in.","Logged out","failed for","saved ->","web page saved","No returns","Captured","No records","Run summary","All done","WARNING","could not be downloaded","Saved:"))]
from openpyxl import load_workbook
p=sorted(glob.glob("/root/Downloads/URA_Suite/Run_Summary_*.xlsx"))[-1]; wb=load_workbook(p)
for ws in wb: print(f"== {ws.title}: {ws.max_row-1} rows"); [print("   ",[str(v)[:28] for v in r][:6]) for r in list(ws.iter_rows(values_only=True))[:8]]
