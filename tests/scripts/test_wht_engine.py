import sys, shutil; sys.path.insert(0,'/home/claude/ura_suite')
from playwright.sync_api import sync_playwright
from engines import wht_engine as we, wht_outputs as wo
import os; we.RESULT_WAIT_MS=int(os.environ.get("RW","60000"))
with sync_playwright() as p:
    b=p.chromium.launch(executable_path="/tmp/chr/chromium", headless=True, args=["--no-sandbox","--disable-gpu","--disable-dev-shm-usage"])
    pg=b.new_context().new_page(); pg.goto("http://localhost:8765/home.html")
    bot=we.WhtBot(pg, print, lambda: None, screenshot_dir="/tmp")
    bot.dismiss_popups(); bot.open_wht_page()
    months=[wo.MonthResult(y,m) for y,m in wo.month_sequence(2026,6,2026,8)]
    we.run_months(bot, months, want_pdf=True, want_html=True, want_screenshot=False, generate_for="D")
    b.close()
print("\n==== RESULTS")
for r in months: print(f"  {r.title}: {r.status}, rows={r.rows}, pdf={len(r.pdf_bytes or b'')}B, warnings={r.warnings}")
ok=[r for r in months if r.df is not None]
if ok: print(ok[0].df.to_string()); print("columns:", list(ok[0].df.columns))
shutil.rmtree("/tmp/out_wht", ignore_errors=True); import os; os.makedirs("/tmp/out_wht")
print(wo.save_excel(months,"/tmp/out_wht/W.xlsx",tin="1000012345"), wo.save_table_pdf(months,"/tmp/out_wht/W.pdf"), wo.merge_pdfs(months,"/tmp/out_wht/P.pdf"))
from openpyxl import load_workbook; wb=load_workbook("/tmp/out_wht/W.xlsx"); print("sheets:", wb.sheetnames)
for r in list(wb["Summary"].iter_rows(values_only=True))[5:10]: print("  ", [v for v in r if v is not None])
