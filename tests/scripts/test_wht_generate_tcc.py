import sys, os, shutil; sys.path.insert(0,'/home/claude/ura_suite')
from playwright.sync_api import sync_playwright
from engines import wht_engine as we, wht_outputs as wo
we.RESULT_WAIT_MS=20000
with sync_playwright() as p:
    b=p.chromium.launch(executable_path="/tmp/chr/chromium", headless=True, args=["--no-sandbox","--disable-gpu","--disable-dev-shm-usage"])
    pg=b.new_context().new_page(); pg.goto("http://localhost:8765/Portal/wht_tcc.html")
    logs=[]; bot=we.WhtBot(pg, logs.append, lambda: None)
    for code in ("D","W"):
        months=[wo.MonthResult(2026,m) for m in (1,2,3)]
        we.run_months(bot, months, want_pdf=False, want_html=False, generate_for=code, generate_tcc=True)
        for r in months:
            tcc = list(r.df["TCC No"]) if r.df is not None and "TCC No" in r.df.columns else None
            print(code, r.label, r.status, r.rows, "| TCC:", tcc, "|", r.message)
        if code=="D":
            os.makedirs("/tmp/out_tcc", exist_ok=True); wo.save_excel(months, "/tmp/out_tcc/D.xlsx")
    b.close()
print("\n".join(l for l in logs if any(k in l for k in ("ticked","Summary","Generated","Portal message","still"))))
from openpyxl import load_workbook
ws=load_workbook("/tmp/out_tcc/D.xlsx")["2026-02"]
print("Excel 2026-02 TCC column:", [row[-2] for row in ws.iter_rows(values_only=True) if row and row[0] and str(row[0]).isdigit()])
